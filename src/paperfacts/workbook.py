"""The Excel export: a consolidated dataset written to an atomic workbook.

Kept apart from :mod:`paperfacts.dataset` on purpose: that module's source is hashed into ``comparison_key``,
because the rows it assembles are verdicts. How those rows are laid out on a sheet (column order, widths,
number formats) decides nothing, so editing it must never rename a stored comparison. This module is in no
cache key.
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from paperfacts.columns import FieldColumn, field_columns
from paperfacts.dataset import DocumentDataset, Row
from paperfacts.kinds import CellValue, interval_text
from paperfacts.profile import IMPLICIT_ENTITY, DomainProfile, EntitySpec
from paperfacts.storage import write_atomic


def data_columns(profile: DomainProfile, entity: EntitySpec | None = None) -> tuple[tuple[str, str], ...]:
    """``(key, header)`` of a paper or sample row, in order: who the row is, then one column per field -- the
    paper-level fields and ``entity``'s own (the primary entity's by default, which in a profile without entity
    types is every field). Here rather than in :mod:`paperfacts.dataset`: a header is display text, and that
    module is hashed."""
    entity = entity or profile.primary
    return (
        ("document_id", "文档ID"),
        ("filename", "文件名"),
        ("sample_id", "样品ID"),
        ("sample_label", "样品标签"),
        ("conditions", "样品及测量条件"),
        ("available_fields", "可用字段数"),
        ("agree_fields", "双路一致字段数"),
        *(
            (key, key)
            for spec in profile.fields
            if not spec.is_sample_level or profile.entity_of(spec).name == entity.name
            for key in (_interval_keys(spec.name) if spec.kind == "interval" else (spec.name,))
        ),
    )


def _interval_keys(name: str) -> tuple[str, str]:
    """The two numeric columns an interval field fills on a data sheet: its low end and its high end."""
    return f"{name} 下限", f"{name} 上限"


def format_cell(value: CellValue, column: FieldColumn) -> CellValue:
    """A field's cell as a sheet holds it, decided by the column rather than by the value's shape: the values of
    a ``many`` column joined with "; ", an interval as one text ("2.8–4.3", "≥ 80", "≤ 5"; a data sheet gives it
    two numeric columns instead), every other value as it is (a boolean is written TRUE/FALSE)."""
    if column.cardinality == "many" and isinstance(value, list):
        return "; ".join("" if item is None else str(item) for item in value)
    if column.kind == "interval" and isinstance(value, list):
        return interval_text(value)  # type: ignore[arg-type]
    return value


def _formatted(rows: Sequence[Row], columns: dict[str, FieldColumn], *, quality: bool = False) -> list[Row]:
    """``rows`` with every field cell through :func:`format_cell`: a data row's field columns, or a quality row's
    ``value`` under the column its ``field`` names."""
    if quality:
        return [
            {**row, "value": format_cell(row.get("value"), columns[str(row["field"])])}
            if row.get("field") in columns
            else row
            for row in rows
        ]
    return [_data_row(row, columns) for row in rows]


def _data_row(row: Row, columns: dict[str, FieldColumn]) -> Row:
    """A data row's field cells through :func:`format_cell`, an interval's split into its two end columns."""
    formatted: dict[str, CellValue] = {}
    for key, value in row.items():
        column = columns.get(key)
        if column is not None and column.kind == "interval":
            ends = value if isinstance(value, list) else [None, None]
            formatted.update(zip(_interval_keys(key), ends, strict=True))
        else:
            formatted[key] = value if column is None else format_cell(value, column)
    return formatted


# What the 字段说明 sheet says of a field with no unit.
_KIND_ZH = {"boolean": "是/否", "date": "日期（ISO）"}


# The 字段说明 rule of an interval column, and of a list column (cardinality "many").
_INTERVAL_RULE = "区间：下限、上限各占一列，开口一端留空；冲突、多条件或无引用定位时两列都留空。"
_REFERENCE_RULE = (
    "引用：所引样品的样品ID，与该实体数据表的样品ID一列一致；"
    "两路所引样品未被匹配为同一样品或引用未定位到列表中的样品时留空。"
)
_LIST_RULE = "多值：两路已定位证据的并集，以“; ”分隔，每个元素的来源通道见数据质量说明；有分类时按分类顺序排列。"
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
# What Excel refuses in a sheet title, and its length limit.
_SHEET_TITLE_REFUSED = re.compile(r"[\[\]:*?/\\]")
_SHEET_TITLE_MAX = 31
_QUALITY_COLUMNS = (
    ("document_id", "文档ID"),
    ("filename", "文件名"),
    ("sample_id", "样品ID"),
    ("field", "字段"),
    ("decision", "最终决策"),
    ("value", "输出值"),
    ("unit", "标准单位"),
    ("conditions", "条件"),
    ("source_ids", "合并证据来源"),
    ("lanes", "证据来源通道"),
    ("series", "系列级"),
    ("detail", "说明"),
)
# Values a vision model read off charts. The rows come from paperfacts.readings (which this module does not
# import: the readings are no part of any verdict here); this is only the sheet's layout.
_FIGURE_COLUMNS = (
    ("document_id", "文档ID"),
    ("filename", "文件名"),
    ("figure", "图"),
    ("page", "页码"),
    ("source_id", "图块来源"),
    ("panel", "子图"),
    ("field", "字段"),
    ("series", "系列"),
    ("x", "横轴（仅供参考，不用于对应样品）"),
    ("value", "读数（近似值）"),
    ("unit", "标准单位"),
    ("precision", "精度"),
    ("value_raw", "图中原始读数"),
    ("scale", "纵轴刻度"),
    ("caption", "图注"),
    ("detail", "说明"),
)


def _worksheet(
    workbook: Workbook,
    title: str,
    columns: Sequence[tuple[str, str]],
    rows: Sequence[Row],
    table_id: str,
    *,
    scientific: frozenset[str] = frozenset(),
) -> Worksheet:
    """``scientific`` names the columns whose numbers the profile displays as ``display_format: scientific``."""
    sheet = workbook.create_sheet(title)
    # openpyxl refuses a control character as it appends, so a stray \x01 in a quote would fail the whole export.
    sheet.append([_clean(label) for _, label in columns])
    for row in rows:
        sheet.append([_clean(row.get(key)) for key, _ in columns])
    # A sheet of sample rows keeps who each row is in view: every column up to the sample id.
    keys = [key for key, _ in columns]
    sheet.freeze_panes = f"{get_column_letter(keys.index('sample_id') + 2)}2" if "sample_id" in keys else "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column, (key, _) in enumerate(columns, start=1):
        width = {
            "document_id": 20,
            "filename": 48,
            "conditions": 48,
            "source_ids": 48,
            "detail": 80,
            "sample_id": 28,
            "sample_label": 35,
            "caption": 60,
            "x": 36,
            "description": 68,
            "rule": 70,
        }.get(key, 23)
        sheet.column_dimensions[get_column_letter(column)].width = width
        for cells in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            cell = cells[0]
            if isinstance(cell.value, bool):
                # openpyxl writes a bool as Excel TRUE/FALSE; a number format would be misleading.
                pass
            elif isinstance(cell.value, str):
                # PDF-derived strings are data even when their first character is '='.
                cell.data_type = "s"
            elif isinstance(cell.value, (int, float)):
                cell.number_format = "0.0000E+00" if key in scientific else "0.############"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if rows:
        table = Table(displayName=table_id, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    return sheet


def _clean(value: object) -> object:
    return _CONTROL.sub("", value) if isinstance(value, str) else value


def _entity_sheets(profile: DomainProfile) -> list[tuple[EntitySpec, str]]:
    """Each entity type's data sheet and its title. A profile without entity types keeps the one 样品数据 sheet;
    with several, each sheet is named after its entity ("催化剂数据"), cleaned of what Excel refuses in a title and
    kept distinct from every other sheet."""
    if len(profile.entities) == 1:
        return [(profile.primary, "样品数据")]
    taken = {"论文数据", "字段说明", "数据质量", "图中读数", "运行记录"}
    sheets: list[tuple[EntitySpec, str]] = []
    for entity in profile.entities:
        base = _SHEET_TITLE_REFUSED.sub("", _CONTROL.sub("", f"{entity.label_zh or entity.name}数据")).strip("'")
        # Excel refuses a title ending in an apostrophe, which the cut may leave.
        title, n = base[:_SHEET_TITLE_MAX].rstrip("'"), 1
        while title.casefold() in {name.casefold() for name in taken}:
            n += 1
            title = f"{base[: _SHEET_TITLE_MAX - len(str(n)) - 1]} {n}"
        taken.add(title)
        sheets.append((entity, title))
    return sheets


def write_dataset(
    documents: Sequence[DocumentDataset],
    output: Path,
    profile: DomainProfile,
    *,
    failures: Sequence[dict[str, str]] = (),
    figure_rows: Sequence[Row] = (),
) -> None:
    """Replace a workbook atomically; repeated PDF hashes produce exactly one paper row.

    ``figure_rows`` (from :func:`paperfacts.readings.figure_rows`) only fill the 图中读数 sheet: chart readings
    are approximate and never compared, so they never reach a sample or paper row.
    """
    unique = sorted(
        {document.document_id: document for document in documents}.values(), key=lambda document: document.document_id
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    several = len(profile.entities) > 1
    by_name = {column.name: column for column in field_columns(profile)}
    scientific = frozenset(spec.name for spec in profile.fields if spec.display_format == "scientific")
    papers = _formatted([doc.paper_row for doc in unique], by_name)
    _worksheet(workbook, "论文数据", data_columns(profile), papers, "Papers", scientific=scientific)
    for entity, title in _entity_sheets(profile):
        rows = [
            row
            for doc in unique
            for row in doc.sample_rows
            if not several or row.get("entity", IMPLICIT_ENTITY) == entity.name
        ]
        _worksheet(
            workbook,
            title,
            data_columns(profile, entity),
            _formatted(rows, by_name),
            f"Samples_{entity.name}" if several else "Samples",
            scientific=scientific,
        )
    entity_labels = {entity.name: entity.label_zh or entity.name for entity in profile.entities}
    descriptions: list[Row] = [
        column.model_dump()
        | {
            # The sheet says the same things in Chinese, for a reader who opens the workbook alone.
            "scope": f"{profile.ui.entity_label_zh}级" if column.scope == "sample" else profile.ui.paper_level_label_zh,
            "unit": column.unit or _KIND_ZH.get(column.kind, "文本"),
            "rule": "冲突、多条件、多值、范围、上下界或无引用定位时留空；近似值和 ± 不确定度保留中心值并备注。",
        }
        | (
            # A list column says so, and that its cell is the union of both lanes, not one agreed value.
            {"unit": "文本（多值）", "rule": _LIST_RULE} if column.cardinality == "many" else {}
        )
        | ({"rule": _INTERVAL_RULE} if column.kind == "interval" else {})
        | (
            # A reference cell is the sample_id of a row on the referenced entity's sheet.
            {"unit": f"{entity_labels.get(column.references, column.references)}样品ID", "rule": _REFERENCE_RULE}
            if column.references is not None
            else {}
        )
        | ({"entity": entity_labels.get(column.entity or "", "")} if several else {})
        for column in by_name.values()
    ]
    _worksheet(
        workbook,
        "字段说明",
        (
            ("name", "字段"),
            ("label", "中文名"),
            ("scope", "层级"),
            *((("entity", "实体"),) if several else ()),
            ("unit", "标准单位"),
            ("description", "中文说明"),
            ("rule", "单值与缺失规则"),
        ),
        descriptions,
        "Fields",
    )
    quality = _formatted([row for doc in unique for row in doc.quality_rows], by_name, quality=True)
    # With several entity types a quality row says whose sample it is about; a paper-level one names none.
    quality_columns = (
        (*_QUALITY_COLUMNS[:2], ("entity", "实体"), *_QUALITY_COLUMNS[2:]) if several else _QUALITY_COLUMNS
    )
    _worksheet(
        workbook,
        "数据质量",
        quality_columns,
        [
            {**row, "entity": entity_labels[str(row["entity"])]} if row.get("entity") in entity_labels else row
            for row in quality
        ],
        "Quality",
    )
    _worksheet(workbook, "图中读数", _FIGURE_COLUMNS, figure_rows, "Figures")
    runs: list[Row] = [
        {
            "document_id": doc.document_id,
            "filename": doc.filename,
            "status": "incomplete" if doc.incomplete else "success",
            "samples": len(doc.sample_rows),
            "extractor_key": doc.extractor_key,
            "comparison_key": doc.comparison_key,
            "detail": f"未完成，下次运行重问：{doc.incomplete}"
            if doc.incomplete
            else "论文行采用一个完整样品；空白为缺失或未通过唯一值质量规则。",
        }
        for doc in unique
    ]
    runs.extend(
        {
            "document_id": failure.get("document_id", ""),
            "filename": failure.get("filename") or failure.get("pdf") or failure.get("path", ""),
            "status": "failed",
            "detail": failure.get("error") or failure.get("detail", str(failure)),
        }
        for failure in failures
    )
    _worksheet(
        workbook,
        "运行记录",
        (
            ("document_id", "文档ID"),
            ("filename", "文件名"),
            ("status", "状态"),
            ("samples", "合并后样品数"),
            ("extractor_key", "抽取版本"),
            ("comparison_key", "比较版本"),
            ("detail", "说明"),
        ),
        runs,
        "Runs",
    )
    write_atomic(output, workbook.save)
