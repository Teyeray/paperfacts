"""Conservative, one-value-per-field datasets and an atomic Excel export.

The paper table selects a complete sample row. It must never manufacture a sample by combining the best
measurement of each field from different experimental conditions. Which value a cell holds -- and whether it
holds one at all -- is decided per cell by :mod:`paperfacts.decide`; this module gathers each cell's evidence,
assembles the rows and writes the workbook.
"""

from __future__ import annotations

import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from types import MappingProxyType

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ConfigDict

from paperfacts.compare import ComparisonReport, FieldComparison
from paperfacts.decide import CellValue, Decision, decide, joined
from paperfacts.fields import (
    AMBIGUOUS_MATCH_CONFIDENCE,
    FIELD_SPECS,
    SAMPLE_FIELDS,
    TARGET_FIELDS,
    FieldSpec,
)
from paperfacts.models import Backend, DocumentInput
from paperfacts.normalize import normalize_lane
from paperfacts.records import LaneExtraction, SampleRecord
from paperfacts.storage import write_atomic

Row = Mapping[str, CellValue]

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_DATA_COLUMNS = (
    ("document_id", "文档ID"),
    ("filename", "文件名"),
    ("sample_id", "样品ID"),
    ("sample_label", "样品标签"),
    ("conditions", "样品及测量条件"),
    ("available_fields", "可用字段数"),
    ("agree_fields", "双路一致字段数"),
    *((spec.name, spec.name) for spec in FIELD_SPECS),
)
_QUALITY_COLUMNS = (
    ("document_id", "文档ID"),
    ("filename", "文件名"),
    ("sample_id", "样品ID"),
    ("field", "字段"),
    ("decision", "最终决策"),
    ("value", "输出值"),
    ("unit", "标准单位"),
    ("conditions", "条件"),
    ("source_ids", "合并证据来源"),
    ("lanes", "证据来源通道"),
    ("series", "系列级"),
    ("detail", "说明"),
)
# Values a vision model read off charts. The rows come from paperfacts.figures (which this module does not
# import: the readings are no part of any verdict here); this is only the sheet's layout.
_FIGURE_COLUMNS = (
    ("document_id", "文档ID"),
    ("filename", "文件名"),
    ("figure", "图"),
    ("page", "页码"),
    ("source_id", "图块来源"),
    ("panel", "子图"),
    ("field", "字段"),
    ("series", "系列"),
    ("x", "横轴（仅供参考，不用于对应样品）"),
    ("value", "读数（近似值）"),
    ("unit", "标准单位"),
    ("precision", "精度"),
    ("value_raw", "图中原始读数"),
    ("scale", "纵轴刻度"),
    ("caption", "图注"),
    ("detail", "说明"),
)


class FieldColumn(BaseModel):
    """What a reader needs to know about one column, built once for both the web UI and the Excel sheet.

    ``label`` and ``description`` are display only and may be empty when ``config.json`` declares neither;
    ``unit`` is absent for a text field.
    """

    model_config = ConfigDict(frozen=True)

    name: str
    label: str = ""
    unit: str | None = None
    scope: str
    description: str = ""


class DatasetPayload(BaseModel):
    """One document's consolidated dataset as it crosses the disk and HTTP boundaries.

    The same model is written to ``dataset.json``, parsed back from it and returned by the endpoint, so
    the browser's contract is declared once and FastAPI can publish a schema for it. The field list
    travels with the data because the rows carry values only: the browser needs the canonical unit and
    the paper/sample scope to build a header it can trust.
    """

    model_config = ConfigDict(frozen=True)

    document_id: str = ""
    filename: str = ""
    extractor_key: str = ""
    comparison_key: str = ""
    # Per lane, the content hash of the parse the table was built from (see ParsedArtifact.content_hash): its
    # cells cite blocks by position. A lane absent here is unknown (a file from before it was recorded).
    artifact_sha256: dict[Backend, str] = {}
    fields: tuple[FieldColumn, ...] = ()
    paper_row: dict[str, CellValue] = {}
    sample_rows: tuple[dict[str, CellValue], ...] = ()
    quality_rows: tuple[dict[str, CellValue], ...] = ()


def field_columns() -> tuple[FieldColumn, ...]:
    """The field table as columns, in the order the dataset writes them."""
    return tuple(
        FieldColumn(
            name=spec.name,
            label=spec.label,
            unit=spec.canonical_unit,
            scope="sample" if spec.is_sample_level else "target",
            description=spec.description_zh,
        )
        for spec in FIELD_SPECS
    )


@dataclass(frozen=True)
class DocumentDataset:
    document_id: str
    filename: str
    paper_row: Row
    sample_rows: tuple[Row, ...]
    quality_rows: tuple[Row, ...]
    extractor_key: str = ""
    comparison_key: str = ""
    artifact_sha256: Mapping[Backend, str] = field(default_factory=dict)

    def to_payload(self) -> DatasetPayload:
        """The serialisable view the web UI and ``dataset.json`` share.

        Rows are ``MappingProxyType`` so nothing downstream can mutate a consolidated row; pydantic copies
        each into a plain dict here rather than weakening the model.
        """
        return DatasetPayload(
            document_id=self.document_id,
            filename=self.filename,
            extractor_key=self.extractor_key,
            comparison_key=self.comparison_key,
            artifact_sha256=dict(self.artifact_sha256),
            fields=field_columns(),
            paper_row=dict(self.paper_row),
            sample_rows=tuple(dict(row) for row in self.sample_rows),
            quality_rows=tuple(dict(row) for row in self.quality_rows),
        )

    @classmethod
    def from_payload(cls, payload: DatasetPayload) -> DocumentDataset:
        """The exact inverse of :meth:`to_payload`, so a dataset read back from disk can be exported again
        without re-running the pipeline. The field list is not restored: it is derived from FIELD_SPECS on
        the way out, and a payload written under a different field table lives under a different
        extractor_key and is never read next to this one.

        The payload arrives already validated -- a malformed file fails at the disk boundary, where the
        caller can decide whether to skip that document or raise.
        """
        return cls(
            document_id=payload.document_id,
            filename=payload.filename,
            paper_row=MappingProxyType(dict(payload.paper_row)),
            sample_rows=tuple(MappingProxyType(dict(row)) for row in payload.sample_rows),
            quality_rows=tuple(MappingProxyType(dict(row)) for row in payload.quality_rows),
            extractor_key=payload.extractor_key,
            comparison_key=payload.comparison_key,
            artifact_sha256=dict(payload.artifact_sha256),
        )


def write_dataset_json(dataset: DocumentDataset, path: Path) -> None:
    """Write one document's consolidated dataset for the web UI, atomically like every other artifact."""
    payload = dataset.to_payload().model_dump_json(indent=2)
    write_atomic(path, lambda tmp: tmp.write_text(payload, encoding="utf-8"))


@dataclass(frozen=True)
class _Scope:
    sample_id: str
    report_scope: str
    a: SampleRecord | None
    b: SampleRecord | None
    confidence: float | None = None
    matching_failed: bool = False


def _matching_blocked(scope: _Scope | None) -> str | None:
    """Why nothing measured on this scope may be committed, or None if it may.

    Scope-wide rather than per-field: if the two lanes' samples were not confidently identified as the
    same sample, no value on them can be trusted, whatever the per-field comparison says. The paper-level
    target row has no scope and so is never blocked this way.
    """
    if scope is None:
        return None
    if scope.matching_failed:
        return "样品匹配失败，无法确认跨通道身份"
    if scope.confidence is not None and scope.confidence < AMBIGUOUS_MATCH_CONFIDENCE:
        return "样品匹配置信度低于阈值"
    return None


def _scopes(lanes: Mapping[Backend, LaneExtraction], report: ComparisonReport) -> tuple[_Scope, ...]:
    lane_a, lane_b = lanes[report.backend_a], lanes[report.backend_b]
    scopes: list[_Scope] = []
    for pair in report.matching.pairs:
        a, b = lane_a.sample(pair.a_id), lane_b.sample(pair.b_id)
        if a is not None and b is not None:
            sample_id = pair.a_id if pair.a_id == pair.b_id else f"{pair.a_id} | {pair.b_id}"
            scopes.append(_Scope(sample_id, f"sample:{pair.a_id}|{pair.b_id}", a, b, pair.confidence))
    for backend, ids, side in (
        (report.backend_a, report.matching.unmatched_a, "a"),
        (report.backend_b, report.matching.unmatched_b, "b"),
    ):
        for sample_id in ids:
            sample = lanes[backend].sample(sample_id)
            if sample is not None:
                scopes.append(
                    _Scope(
                        f"{backend}:{sample_id}",
                        f"sample:{sample_id}",
                        sample if side == "a" else None,
                        sample if side == "b" else None,
                        matching_failed=report.matching.failed,
                    )
                )
    return tuple(sorted(scopes, key=lambda scope: (scope.sample_id, scope.report_scope)))


def _scope_comparisons(scope: _Scope, report: ComparisonReport) -> tuple[FieldComparison, ...]:
    # Unmatched IDs can be identical across lanes; the comparison's populated side
    # disambiguates them without parsing the report's deliberately lossy scope string.
    return tuple(
        comparison
        for comparison in report.comparisons
        if comparison.scope == scope.report_scope
        and (scope.a is not None or comparison.a is None)
        and (scope.b is not None or comparison.b is None)
    )


def consolidate_document(
    document: DocumentInput, lanes: Mapping[Backend, LaneExtraction], report: ComparisonReport
) -> DocumentDataset:
    """Collapse source evidence, then select the most complete trustworthy sample row."""
    if report.document_id != document.document_id or any(
        lane.document_id != document.document_id for lane in lanes.values()
    ):
        raise ValueError("document, extraction lanes and comparison report must refer to the same PDF")
    if any(lane.extractor_key != report.extractor_key for lane in lanes.values()):
        raise ValueError("extraction lanes and comparison report have different extractor keys")
    lanes = {backend: normalize_lane(lane) for backend, lane in lanes.items()}
    metadata: dict[str, CellValue] = {"document_id": document.document_id, "filename": document.display_filename}
    quality: list[Row] = []

    def record(sample_id: str, spec: FieldSpec, decision: Decision) -> None:
        quality.append(
            MappingProxyType(
                {
                    **metadata,
                    "sample_id": sample_id,
                    "field": spec.name,
                    "decision": decision.status,
                    "value": decision.value,
                    "unit": spec.canonical_unit,
                    "conditions": decision.conditions,
                    "source_ids": decision.sources,
                    "lanes": "; ".join(decision.lanes),
                    "series": decision.series,
                    "detail": decision.detail,
                }
            )
        )

    target: dict[str, Decision] = {}
    for spec in TARGET_FIELDS:
        evidence = [
            (backend, field)
            for backend, lane in lanes.items()
            if lane.target
            for field in lane.target.fields
            if field.field == spec.name
        ]
        target[spec.name] = decide(
            spec, evidence, [c for c in report.comparisons if c.scope == "target" and c.field == spec.name]
        )
    for spec in TARGET_FIELDS:
        record("target", spec, target[spec.name])

    sample_rows: list[Row] = []
    for scope in _scopes(lanes, report):
        scope_comparisons = _scope_comparisons(scope, report)
        decisions = dict(target)
        for spec in SAMPLE_FIELDS:
            evidence = [
                (backend, field)
                for backend, sample in ((report.backend_a, scope.a), (report.backend_b, scope.b))
                if sample is not None
                for field in sample.fields
                if field.field == spec.name
            ]
            row_sources = frozenset(
                source
                for sample in (scope.a, scope.b)
                if sample is not None
                for field in sample.fields
                if field.field != spec.name and field.grounded
                for source in field.source_ids
            )
            decision = decide(
                spec,
                evidence,
                [c for c in scope_comparisons if c.field == spec.name],
                blocked=_matching_blocked(scope),
                row_sources=row_sources,
            )
            decisions[spec.name] = decision
            record(scope.sample_id, spec, decision)
        samples = [sample for sample in (scope.a, scope.b) if sample is not None]
        conditions = joined(
            [f"{key}={value}" for sample in samples for key, value in sorted(sample.conditions.items())]
            + [
                f"{spec.name}: {decisions[spec.name].conditions}"
                for spec in SAMPLE_FIELDS
                if decisions[spec.name].conditions
            ]
        )
        sample_rows.append(
            MappingProxyType(
                {
                    **metadata,
                    "sample_id": scope.sample_id,
                    "sample_label": joined([sample.label for sample in samples]),
                    "conditions": conditions,
                    "available_fields": sum(d.value is not None for d in decisions.values()),
                    "agree_fields": sum(d.status == "agree" for d in decisions.values()),
                    **{name: decision.value for name, decision in decisions.items()},
                }
            )
        )
    if sample_rows:
        paper_row = min(
            sample_rows,
            key=lambda row: (-int(row["available_fields"] or 0), -int(row["agree_fields"] or 0), str(row["sample_id"])),
        )
        selection = "按可用字段数最多、双路一致字段数最多、样品ID稳定排序，选择整行；未跨样品拼接字段"
    else:
        paper_row = MappingProxyType(
            {
                **metadata,
                "sample_id": "",
                "sample_label": "",
                "conditions": "",
                "available_fields": sum(d.value is not None for d in target.values()),
                "agree_fields": sum(d.status == "agree" for d in target.values()),
                **{spec.name: target[spec.name].value if spec.name in target else None for spec in FIELD_SPECS},
            }
        )
        selection = "未提取到可匹配样品；论文行仅保留唯一的靶材字段"
    quality.append(
        MappingProxyType(
            {
                **metadata,
                "sample_id": paper_row["sample_id"],
                "field": "__selection__",
                "decision": "selected_sample",
                "detail": selection,
            }
        )
    )
    return DocumentDataset(
        document.document_id,
        document.display_filename,
        paper_row,
        tuple(sample_rows),
        tuple(quality),
        report.extractor_key,
        report.comparison_key,
        {
            backend: sha
            for backend, sha in (
                (report.backend_a, report.artifact_sha256_a),
                (report.backend_b, report.artifact_sha256_b),
            )
            if sha is not None
        },
    )


def _worksheet(
    workbook: Workbook, title: str, columns: Sequence[tuple[str, str]], rows: Sequence[Row], table_id: str
) -> Worksheet:
    sheet = workbook.create_sheet(title)
    sheet.append([label for _, label in columns])
    for row in rows:
        sheet.append([row.get(key) for key, _ in columns])
    sheet.freeze_panes = "D2" if title in {"论文数据", "样品数据", "数据质量"} else "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column, (key, _) in enumerate(columns, start=1):
        width = {
            "document_id": 20,
            "filename": 48,
            "conditions": 48,
            "source_ids": 48,
            "detail": 80,
            "sample_id": 28,
            "sample_label": 35,
            "caption": 60,
            "x": 36,
            "description": 68,
            "rule": 70,
        }.get(key, 23)
        sheet.column_dimensions[get_column_letter(column)].width = width
        for cells in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            cell = cells[0]
            if isinstance(cell.value, bool):
                # openpyxl writes a bool as Excel TRUE/FALSE; a number format would be misleading.
                pass
            elif isinstance(cell.value, str):
                cell.value = _CONTROL.sub("", cell.value)
                # PDF-derived strings are data even when their first character is '='.
                cell.data_type = "s"
            elif isinstance(cell.value, (int, float)):
                cell.number_format = "0.0000E+00" if key in {"resistance", "resistivity"} else "0.############"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if rows:
        table = Table(displayName=table_id, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    return sheet


def write_dataset(
    documents: Sequence[DocumentDataset],
    output: Path,
    *,
    failures: Sequence[dict[str, str]] = (),
    figure_rows: Sequence[Row] = (),
) -> None:
    """Replace a workbook atomically; repeated PDF hashes produce exactly one paper row.

    ``figure_rows`` (from :func:`paperfacts.figures.figure_rows`) only fill the 图中读数 sheet: chart readings
    are approximate and never compared, so they never reach a sample or paper row.
    """
    unique = sorted(
        {document.document_id: document for document in documents}.values(), key=lambda document: document.document_id
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    _worksheet(workbook, "论文数据", _DATA_COLUMNS, [doc.paper_row for doc in unique], "Papers")
    _worksheet(workbook, "样品数据", _DATA_COLUMNS, [row for doc in unique for row in doc.sample_rows], "Samples")
    descriptions: list[Row] = [
        column.model_dump()
        | {
            # The sheet says the same things in Chinese, for a reader who opens the workbook alone.
            "scope": "样品级" if column.scope == "sample" else "靶材（论文级）",
            "unit": column.unit or "文本",
            "rule": "冲突、多条件、多值、范围、上下界或无引用定位时留空；近似值和 ± 不确定度保留中心值并备注。",
        }
        for column in field_columns()
    ]
    _worksheet(
        workbook,
        "字段说明",
        (
            ("name", "字段"),
            ("label", "中文名"),
            ("scope", "层级"),
            ("unit", "标准单位"),
            ("description", "中文说明"),
            ("rule", "单值与缺失规则"),
        ),
        descriptions,
        "Fields",
    )
    _worksheet(workbook, "数据质量", _QUALITY_COLUMNS, [row for doc in unique for row in doc.quality_rows], "Quality")
    _worksheet(workbook, "图中读数", _FIGURE_COLUMNS, figure_rows, "Figures")
    runs: list[Row] = [
        {
            "document_id": doc.document_id,
            "filename": doc.filename,
            "status": "success",
            "samples": len(doc.sample_rows),
            "extractor_key": doc.extractor_key,
            "comparison_key": doc.comparison_key,
            "detail": "论文行采用一个完整样品；空白为缺失或未通过唯一值质量规则。",
        }
        for doc in unique
    ]
    runs.extend(
        {
            "document_id": failure.get("document_id", ""),
            "filename": failure.get("filename") or failure.get("pdf") or failure.get("path", ""),
            "status": "failed",
            "detail": failure.get("error") or failure.get("detail", str(failure)),
        }
        for failure in failures
    )
    _worksheet(
        workbook,
        "运行记录",
        (
            ("document_id", "文档ID"),
            ("filename", "文件名"),
            ("status", "状态"),
            ("samples", "合并后样品数"),
            ("extractor_key", "抽取版本"),
            ("comparison_key", "比较版本"),
            ("detail", "说明"),
        ),
        runs,
        "Runs",
    )
    write_atomic(output, workbook.save)
