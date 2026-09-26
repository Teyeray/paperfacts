"""The ``catalysis`` example profile (round 2 spec §6), end to end: catalysts and the reaction tests run on them, two
entity types in one paper linked by a reference field, through the real pipeline with no code edit.

Nothing is parsed and no model is called. The two parser outputs are hand-built into the data_root, which the parse
stage takes as a cache hit, and a scripted :class:`FakeLlmClient` answers each question from its text: which entity an
inventory asks about, which field a question asks for, and which lane's excerpts it was shown. The two lanes name
their samples differently, so every matching is a question to the model. No gold set of real catalysis papers exists,
so none of this measures extraction quality; it proves every structural feature the profile uses runs.
"""

from __future__ import annotations

import dataclasses
import importlib.util
import json
import re
import shutil
import sys
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from typer.testing import CliRunner

from paperfacts.batch import export_document
from paperfacts.cli import app
from paperfacts.config import DEFAULT_REPO_ROOT, Settings
from paperfacts.errors import ConfigError
from paperfacts.keys import extractor_key_for
from paperfacts.kinds import NO_CONTEXT
from paperfacts.models import BACKENDS, DocumentGeometry, DocumentInput
from paperfacts.normalize import normalize_field
from paperfacts.profile import DomainProfile
from paperfacts.profile_loader import load_profile
from paperfacts.prompts import inventory_system_prompt, matching_system_prompt
from paperfacts.records import ExtractedRecords, FieldValue, SampleRecord
from paperfacts.storage import DataLayout
from paperfacts.voting import merge_passes
from paperfacts.web.app import create_app
from paperfacts.web.jobs import JobManager
from paperfacts.workflow import PipelineResult, load_run_profile, run_document, stage_names
from support.factories import RawOutputFactory, paddle_page_entry
from support.llm import FakeLlmClient
from support.profiles import SHIPPED_PROFILE_PATH

CATALYSIS_PATH = SHIPPED_PROFILE_PATH.with_name("catalysis.json")
GOLD_FIXTURES = Path(__file__).parent / "fixtures" / "catalysis_gold"

TITLE = "Cu/ZnO catalysts for CO2 hydrogenation to methanol"
PARAGRAPHS = (
    "Received 12 March 2021; accepted in revised form in May.",
    "The precursors were Cu(NO3)2·3H2O, Zn(NO3)2·6H2O and ZrO(NO3)2. Two Cu/ZnO catalysts, CZ-1 with 10 % Cu and"
    " CZ-2 with 20 % Cu, were prepared by co-precipitation and calcined at 450-500 °C.",
    "The catalysts were characterized by XRD and XPS, by BET and by UV-vis; the BET surface area of CZ-1 is"
    " 85 m2 g-1 and that of CZ-2 is 64 m2 g-1.",
    "Both catalysts were pre-reduced in H2 before the tests. Test T1 ran CZ-1 at 250 °C and 3 MPa over the window"
    " 200–300 °C, GHSV 6000 mL g-1 h-1, H2/CO2 = 3:1, giving a CO2 conversion of 12.5 % and a methanol selectivity"
    " of 68 %.",
    "Test T2 ran CZ-2 above 240 °C at 30 bar with the same feed.",
)
# Each lane names its samples its own way; the reference answers below name them as that lane listed them.
INVENTORY = {
    ("mineru", "catalysts"): (("CZ-1", "10 % Cu"), ("CZ-2", "20 % Cu")),
    ("paddleocr_vl", "catalysts"): (("Cat 1", "10 % Cu catalyst"), ("Cat 2", "20 % Cu catalyst")),
    ("mineru", "reaction tests"): (("T1", "CZ-1 at 250 °C"), ("T2", "CZ-2 above 240 °C")),
    ("paddleocr_vl", "reaction tests"): (("run-1", "first test"), ("run-2", "second test")),
}


def _answer(sample: int | None, value_raw: str, unit_raw: str | None = None, **extra: Any) -> dict[str, Any]:
    """One answer: the listed sample it is for (None for a paper-level field), the quote, and anything else the
    answer carries (``holds``; ``cite``, the text of the block to cite when the quote alone would find another)."""
    return {"sample": sample, "value_raw": value_raw, "unit_raw": unit_raw, **extra}


BOTH: dict[str, list[dict[str, Any]]] = {
    "received_date": [_answer(None, "12 March 2021")],
    "composition": [_answer(0, "Cu/ZnO"), _answer(1, "Cu/ZnO")],
    "metal_loading": [_answer(0, "10", "%", cite="10 % Cu and"), _answer(1, "20", "%", cite="20 % Cu,")],
    "preparation_method": [_answer(0, "co-precipitation"), _answer(1, "co-precipitation")],
    "calcination_temperature": [_answer(0, "450-500", "°C"), _answer(1, "450-500", "°C")],
    "bet_surface_area": [_answer(0, "85", "m2 g-1"), _answer(1, "64", "m2 g-1")],
    "reaction_temperature": [_answer(0, "250", "°C")],
    "temperature_window": [_answer(0, "200–300", "°C"), _answer(1, "240", "°C")],
    "reaction_pressure": [_answer(0, "3", "MPa", cite="3 MPa"), _answer(1, "30", "bar", cite="30 bar")],
    "ghsv": [_answer(0, "6000", "mL g-1 h-1")],
    "h2_co2_ratio": [_answer(0, "3:1")],
    "co2_conversion": [_answer(0, "12.5", "%")],
    "methanol_selectivity": [_answer(0, "68", "%")],
}
ANSWERS: dict[str, dict[str, list[dict[str, Any]]]] = {
    "mineru": BOTH
    | {
        # "UV-vis" names no category and is refused as an element; only this lane reads BET.
        "characterization_techniques": [_answer(None, v) for v in ("XPS", "XRD", "BET", "UV-vis")],
        "precursors": [_answer(None, v) for v in ("Cu(NO3)2·3H2O", "Zn(NO3)2·6H2O")],
        "pre_reduced": [_answer(0, "pre-reduced", holds=True), _answer(1, "pre-reduced", holds=True)],
        # T2's catalyst is an id the catalyst list does not hold.
        "catalyst": [_answer(0, "CZ-1"), _answer(1, "CZ-9")],
    },
    "paddleocr_vl": BOTH
    | {
        # "XRD and XPS" in one quote names two categories and is refused as an element.
        "characterization_techniques": [_answer(None, v) for v in ("XRD", "XPS", "XRD and XPS")],
        "precursors": [_answer(None, v) for v in ("Cu(NO3)2·3H2O", "ZrO(NO3)2")],
        # The second answer has no holds: a yes/no field without it is dropped at cleaning.
        "pre_reduced": [_answer(0, "pre-reduced", holds=True), _answer(1, "pre-reduced")],
        # "cat-1" is the listed "Cat 1" by sample_key.
        "catalyst": [_answer(0, "cat-1")],
    },
}

_SOURCE = re.compile(r"<!-- source: (\S+) -->")
_LISTED_ID = re.compile(r"^- id: (.+?) \| label:", re.MULTILINE)
_FIELD = re.compile(r"\AField to extract:\n- `([^`]+)`")
_PLURAL = re.compile(r"list the (.+?) the paper reports")


def _lane(user: str) -> str:
    return "paddleocr_vl" if "<!-- source: paddleocr_vl_" in user else "mineru"


def _citing(excerpts: str, text: str) -> list[str]:
    for block in re.split(r"(?=<!-- source: )", excerpts):
        found = _SOURCE.match(block)
        if found and text in block:
            return [found.group(1)]
    return _SOURCE.findall(excerpts)[:1]


def scripted_model(profile: DomainProfile):
    def respond(system: str, user: str) -> str:
        lane = _lane(user)
        if user.startswith("Paper excerpts"):
            plural = _PLURAL.search(system)
            assert plural is not None, system[:200]
            cited = _SOURCE.findall(user)[:1]
            samples = [
                {"sample_id": sid, "label": label, "conditions": {}, "source_ids": cited}
                for sid, label in INVENTORY[(lane, plural.group(1))]
            ]
            return json.dumps({"samples": samples, profile.prompt.no_samples_key: False})
        if user.startswith("List A (parser:"):
            list_a, list_b = user.split("\n\nList B (parser:", 1)
            pairs = [
                {"a": a, "b": b, "confidence": 0.9, "justification": "same position"}
                for a, b in zip(_LISTED_ID.findall(list_a), _LISTED_ID.findall(list_b), strict=True)
            ]
            return json.dumps({"pairs": pairs, "unmatched_a": [], "unmatched_b": []})
        field = _FIELD.match(user)
        assert field is not None, user[:200]
        question, excerpts = user.split("Excerpts (Markdown with provenance markers):", 1)
        # A reference question lists the referenced catalysts after the field's own tests: the own list comes first.
        listed = _LISTED_ID.findall(
            question.split("Catalysts this paper reports:")[0] if field.group(1) == "catalyst" else question
        )
        values = []
        for answer in ANSWERS[lane].get(field.group(1), []):
            index = answer["sample"]
            value = {
                "sample_id": None if index is None else listed[index],
                "value_raw": answer["value_raw"],
                "unit_raw": answer["unit_raw"],
                "condition": None,
                "source_ids": _citing(excerpts, answer.get("cite", answer["value_raw"])),
            }
            if "holds" in answer:
                value["holds"] = answer["holds"]
            values.append(value)
        return json.dumps({"values": values})

    return respond


def seed_parses(document: DocumentInput, layout: DataLayout, geometry: DocumentGeometry) -> None:
    blocks = (TITLE, *PARAGRAPHS)
    for backend in BACKENDS:
        raw_dir = layout.raw_dir(document.document_id, backend)
        factory = RawOutputFactory(raw_dir.parent, document, geometry)
        if backend == "mineru":
            content = [
                {"type": "text", "page_idx": 0, "bbox": [100, 80 + 120 * i, 900, 190 + 120 * i], "text": text}
                | ({"text_level": 1} if i == 0 else {})
                for i, text in enumerate(blocks)
            ]
            factory.mineru(content, dir_name=raw_dir.name)
        else:
            parsing = [
                {
                    "block_id": i,
                    "block_order": i,
                    "block_label": "doc_title" if i == 0 else "text",
                    "block_bbox": [200, 150 + 300 * i, 1450, 420 + 300 * i],
                    "block_content": text,
                }
                for i, text in enumerate(blocks)
            ]
            page = paddle_page_entry(0, {"res": {"parsing_res_list": parsing}}, (1654, 2339))
            factory.paddle([page], dir_name=raw_dir.name)


@pytest.fixture(scope="module")
def profile() -> DomainProfile:
    return load_profile(CATALYSIS_PATH)


@pytest.fixture
def run(
    tmp_path: Path, document: DocumentInput, geometry: DocumentGeometry, monkeypatch
) -> tuple[PipelineResult, FakeLlmClient, Settings, DocumentInput]:
    settings = Settings(
        data_root=tmp_path / "data",
        repo_root=tmp_path,
        llm_api_key="sk-test",
        llm_model="fake-model",
        profile=str(CATALYSIS_PATH),
    )
    seed_parses(document, DataLayout(settings.data_root), geometry)
    run_profile = load_run_profile(settings)
    client = FakeLlmClient(scripted_model(run_profile))
    monkeypatch.setattr("paperfacts.workflow.build_llm_client", lambda _settings: client)
    return run_document(document, settings, run_profile), client, settings, document


def _rows(result: PipelineResult) -> dict[tuple[str, str], dict[str, Any]]:
    return {(row["entity"], row["sample_id"]): dict(row) for row in result.dataset.sample_rows}


def _quality(result: PipelineResult) -> dict[tuple[str | None, str, str], dict[str, Any]]:
    return {(q.get("entity"), q["sample_id"], q["field"]): dict(q) for q in result.dataset.quality_rows}


# ---- 1, 2: the profile loads, checks clean, and runs in passage mode only ----------------------------------------


def test_the_catalysis_profile_is_an_example_of_two_entity_types_linked_by_a_reference(profile: DomainProfile):
    assert (profile.name, profile.maturity) == ("catalysis", "example")
    assert [entity.name for entity in profile.entities] == ["catalyst", "test"]
    assert profile.by_name["catalyst"].references == "catalyst"
    kinds = {spec.kind for spec in profile.fields}
    assert kinds == {"numeric", "composition", "text", "boolean", "date", "interval", "reference"}
    assert {spec.name for spec in profile.fields if spec.cardinality == "many"} == {
        "characterization_techniques",
        "precursors",
    }
    assert profile.by_name["calcination_temperature"].range_policy == "upper"
    assert profile.units.convert("MPa", "bar") == (0.1, 0.0)
    assert profile.units.convert("mL/(g·h)", "mL g-1 h-1") == (1.0, 0.0)
    assert profile.units.convert("m2/g", "m2 g-1") == (1.0, 0.0)


def test_profiles_check_passes_and_notes_the_mode():
    result = CliRunner().invoke(app, ["profiles", "--check", str(CATALYSIS_PATH)])

    assert result.exit_code == 0, result.output
    assert result.output.splitlines()[-1] == "ok"
    assert "runs only in passage mode" in result.output


def test_document_mode_is_refused_when_it_is_loaded_to_run(tmp_path: Path):
    settings = Settings(repo_root=tmp_path, profile=str(CATALYSIS_PATH), extraction_mode="document")

    with pytest.raises(ConfigError, match=r"extraction\.mode"):
        load_run_profile(settings)


# ---- 2-10: one run through the real pipeline ----------------------------------------------------------------------


def test_each_lane_asks_one_inventory_per_entity_and_each_entity_is_matched(run, profile: DomainProfile):
    result, client, _, _ = run

    inventories = [call.system for call in client.calls if call.user.startswith("Paper excerpts")]
    assert sorted(inventories) == sorted([inventory_system_prompt(profile, e) for e in profile.entities] * 2)
    matchings = [call.system for call in client.calls if call.user.startswith("List A")]
    assert sorted(matchings) == sorted(matching_system_prompt(profile, e) for e in profile.entities)
    assert list(result.report.matchings) == ["catalyst", "test"]
    assert result.report.counts.samples_matched == 4 and result.report.counts.samples_unmatched == 0
    assert not result.dataset.incomplete


def test_the_paper_level_lists_and_date(run):
    result, _, _, _ = run
    paper = result.dataset.paper_row
    quality = _quality(result)

    # 3. Category order, "UV-vis" and "XRD and XPS" refused as elements, single_source because one lane lacks BET.
    assert paper["characterization_techniques"] == ["XRD", "XPS", "BET"]
    techniques = quality[(None, "paper", "characterization_techniques")]
    assert techniques["decision"] == "single_source"
    assert "UV-vis" in techniques["detail"] and "XRD and XPS" in techniques["detail"]
    # 4. The union, each element with the lanes that read it.
    assert paper["precursors"] == ["Cu(NO3)2·3H2O", "Zn(NO3)2·6H2O", "ZrO(NO3)2"]
    detail = quality[(None, "paper", "precursors")]["detail"]
    assert "Cu(NO3)2·3H2O（mineru, paddleocr_vl）" in detail and "ZrO(NO3)2（paddleocr_vl）" in detail
    # 5. The date at the precision written.
    assert paper["received_date"] == "2021-03-12"
    assert quality[(None, "paper", "received_date")]["decision"] == "agree"


def test_a_date_that_could_be_read_two_ways_is_refused(profile: DomainProfile):
    spec = profile.by_name["received_date"]

    read = normalize_field(FieldValue(field="received_date", value_raw="03/04/2021"), spec, profile.units, NO_CONTEXT)

    assert read.iso_date is None and read.normalization_note


def test_the_catalyst_rows(run):
    result, _, _, _ = run
    rows = _rows(result)

    assert set(rows) == {
        ("catalyst", "CZ-1 | Cat 1"),
        ("catalyst", "CZ-2 | Cat 2"),
        ("test", "T1 | run-1"),
        ("test", "T2 | run-2"),
    }
    one = rows[("catalyst", "CZ-1 | Cat 1")]
    # 6. "450-500 °C" under range_policy upper is its upper end, with the note saying so.
    assert one["calcination_temperature"] == pytest.approx(500.0)
    assert "按字段配置取上限" in _quality(result)[("catalyst", "CZ-1 | Cat 1", "calcination_temperature")]["detail"]
    # Declared units: "85 m2 g-1" is 85 m2/g.
    assert one["bet_surface_area"] == pytest.approx(85.0)
    assert one["metal_loading"] == pytest.approx(10.0) and one["preparation_method"] == "co-precipitation"
    assert "reaction_temperature" not in one
    # 10. The paper row is one of the catalysts' rows.
    assert result.dataset.paper_row["entity"] == "catalyst"


def test_the_reaction_test_rows(run):
    result, _, _, _ = run
    rows = _rows(result)
    quality = _quality(result)
    first, second = rows[("test", "T1 | run-1")], rows[("test", "T2 | run-2")]

    # 7. A window gives both ends; "240" quoted out of "above 240 °C" is a bound, its upper end open.
    assert first["temperature_window"] == [200.0, 300.0]
    assert second["temperature_window"] == [240.0, None]
    # "30 bar" is 3 MPa: one declared unit, two spellings.
    assert first["reaction_pressure"] == pytest.approx(3.0) and second["reaction_pressure"] == pytest.approx(3.0)
    assert first["ghsv"] == pytest.approx(6000.0) and first["co2_conversion"] == pytest.approx(12.5)
    # 8. A yes/no field: agreed, and an answer without holds dropped with the reason in the lane's audit.
    assert first["pre_reduced"] is True and quality[("test", "T1 | run-1", "pre_reduced")]["decision"] == "agree"
    assert second["pre_reduced"] is True
    assert quality[("test", "T2 | run-2", "pre_reduced")]["decision"] == "single_source"
    assert any("boolean field without holds" in line for line in result.lanes["paddleocr_vl"].dropped)
    # 9. The reference cell is the catalyst's row id; an id the catalyst list does not hold is refused.
    assert first["catalyst"] == "CZ-1 | Cat 1"
    assert quality[("test", "T1 | run-1", "catalyst")]["decision"] == "agree"
    assert second["catalyst"] is None
    assert quality[("test", "T2 | run-2", "catalyst")]["decision"] == "ungrounded"


def test_a_reference_still_fills_its_cell_after_the_lanes_are_read_back(run, profile: DomainProfile):
    first, client, settings, document = run
    for derived in ("comparisons", "datasets"):
        shutil.rmtree(DataLayout(settings.data_root).doc_dir(document.document_id) / derived)
    calls = client.call_count

    second = run_document(document, settings, profile)

    # Both lanes come from their files through read_lane, which re-grounds every value; only matching is asked again.
    assert all(call.user.startswith("List A") for call in client.calls[calls:])
    assert _rows(second)[("test", "T1 | run-1")]["catalyst"] == "CZ-1 | Cat 1"
    assert [dict(row) for row in second.dataset.sample_rows] == [dict(row) for row in first.dataset.sample_rows]


def test_the_workbook_has_a_sheet_per_entity_and_re_export_is_identical(run, profile: DomainProfile):
    result, _, settings, document = run
    book = load_workbook(result.excel_path)

    assert book.sheetnames[:3] == ["论文数据", "催化剂数据", "反应测试数据"]
    tests = list(book["反应测试数据"].values)
    header = tests[0]
    assert "temperature_window 下限" in header and "temperature_window 上限" in header
    assert "calcination_temperature" not in header
    catalysts = list(book["催化剂数据"].values)
    assert "bet_surface_area" in catalysts[0] and "ghsv" not in catalysts[0]

    exported = export_document(document, settings, profile)
    assert [dict(row) for row in exported.sample_rows] == [dict(row) for row in result.dataset.sample_rows]
    assert [dict(row) for row in exported.quality_rows] == [dict(row) for row in result.dataset.quality_rows]
    assert dict(exported.paper_row) == dict(result.dataset.paper_row)


# ---- 8: the vote keeps a yes and a no apart --------------------------------------------------------------------


def test_with_three_passes_a_true_false_split_never_votes_as_one_value():
    def one_pass(holds: bool) -> ExtractedRecords:
        value = FieldValue(field="pre_reduced", value_raw="pre-reduced", source_ids=("b1",), holds=holds)
        sample = SampleRecord(sample_id="T1", entity="test", fields=(value,))
        return ExtractedRecords(paper=None, samples=(sample,), invalid_source_ids=(), dropped=())

    split = merge_passes([one_pass(True), one_pass(False)], reference_fields={"catalyst"})
    majority = merge_passes([one_pass(True), one_pass(False), one_pass(True)], reference_fields={"catalyst"})

    assert split.samples[0].fields == ()
    assert [(v.holds, v.agreement) for v in majority.samples[0].fields] == [(True, pytest.approx(2 / 3))]


# ---- 11: the web page learns the entities ---------------------------------------------------------------------


def test_the_profile_endpoint_lists_the_entities(tmp_path: Path, profile: DomainProfile):
    settings = Settings(data_root=tmp_path / "data", repo_root=tmp_path, llm_api_key="sk-test")
    jobs = JobManager(lambda *args, **kwargs: None, stage_names())

    with TestClient(create_app(settings, profile=profile, jobs=jobs)) as client:
        body = client.get("/api/profile").json()

    assert [(e["name"], e["label_zh"]) for e in body["entities"]] == [("catalyst", "催化剂"), ("test", "反应测试")]
    catalyst = next(field for field in body["fields"] if field["name"] == "catalyst")
    assert (catalyst["entity"], catalyst["references"]) == ("test", "catalyst")


# ---- 12: TCO is untouched -------------------------------------------------------------------------------------


def test_tco_keys_are_those_pinned_and_differ_from_the_catalysis_run(run):
    from test_tco_fingerprints_pinned import B1

    result, _, settings, _ = run
    tco = load_profile(SHIPPED_PROFILE_PATH)
    assert {function: function(tco) for function in B1} == B1
    tco_settings = dataclasses.replace(settings, profile=str(SHIPPED_PROFILE_PATH))
    assert extractor_key_for(tco_settings, tco) != result.dataset.extractor_key
    assert result.excel_path.name == "catalysis.xlsx"


# ---- 13: the scorer on a synthetic gold set ---------------------------------------------------------------------


def _score_module():
    spec = importlib.util.spec_from_file_location("paperfacts_score_catalysis", DEFAULT_REPO_ROOT / "eval" / "score.py")
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def test_the_scorer_scores_the_synthetic_gold_set(tmp_path: Path):
    score = _score_module()
    datasets = [f"{path.stem}={path}" for path in sorted((GOLD_FIXTURES / "datasets").glob("*.json"))]
    cells = tmp_path / "cells.json"

    assert score.main([
        "--profile", str(CATALYSIS_PATH),
        "--gold", str(GOLD_FIXTURES / "gold"),
        *(item for dataset in datasets for item in ("--dataset", dataset)),
        "--out", str(tmp_path / "report.md"),
        "--json", str(cells),
    ]) == 0  # fmt: skip

    # One cell per list element: the extra precursor, the missing TEM, the wrong support and yes/no, the reference
    # left empty where the lanes named no listed catalyst.
    outcomes = [[c["doc"], c["sample"], c["field"], c["outcome"]] for c in json.loads(cells.read_text("utf-8"))]
    assert outcomes == json.loads((GOLD_FIXTURES / "expected.json").read_text(encoding="utf-8"))
