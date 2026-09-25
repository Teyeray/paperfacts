"""A batch must survive one bad paper, resume caches, and never duplicate a PDF's row."""

from __future__ import annotations

import re
import shutil

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from paperfacts.batch import discover_pdfs, run_batch
from paperfacts.cli import app
from paperfacts.config import Settings
from paperfacts.errors import ConfigError, ParserError
from paperfacts.matching import SampleMatch, SampleMatching
from paperfacts.models import DocumentInput
from paperfacts.storage import DataLayout
from paperfacts.workflow import run_document
from support.factories import make_blank_pdf
from test_workflow_run import install_fake_pipeline


def test_discovery_is_recursive_case_insensitive_and_ignores_non_pdfs(tmp_path):
    lower = make_blank_pdf(tmp_path / "a.pdf")
    upper = make_blank_pdf(tmp_path / "nested" / "b.PDF")
    (tmp_path / "not_pdf.txt").write_text("ignored")
    assert discover_pdfs(tmp_path) == (lower, upper)
    assert discover_pdfs(upper) == (upper,)


def test_empty_input_fails_without_creating_a_success_workbook(tmp_path):
    settings = Settings(data_root=tmp_path / "data")
    with pytest.raises(ConfigError, match="no PDF"):
        run_batch(tmp_path, settings)
    assert not DataLayout(settings.data_root).batch_dataset_path().exists()


def test_duplicate_pdf_content_gets_one_pipeline_run_and_one_row(monkeypatch, tmp_path):
    source = tmp_path / "papers"
    original = make_blank_pdf(source / "original.pdf")
    shutil.copyfile(original, source / "copy.PDF")
    spy = install_fake_pipeline(monkeypatch)

    result = run_batch(source, Settings(data_root=tmp_path / "data"))

    assert len(result.documents) == 1
    assert result.duplicate_count == 1
    assert len(spy.parse) == 2
    workbook = load_workbook(result.excel_path)
    assert workbook["论文数据"].max_row == 2
    workbook.close()


def test_a_failed_paper_is_recorded_and_remaining_papers_are_exported(monkeypatch, tmp_path):
    source = tmp_path / "papers"
    make_blank_pdf(source / "a_bad.pdf", sizes=[(200, 200)])
    make_blank_pdf(source / "b_good.pdf", sizes=[(300, 300)])
    install_fake_pipeline(monkeypatch)
    output = tmp_path / "output.xlsx"
    snapshots = []
    from paperfacts.workbook import write_dataset

    def capture_write(documents, path, profile, *, failures=(), figure_rows=()):
        write_dataset(documents, path, profile, failures=failures)
        if path == output:
            snapshots.append((len(documents), len(failures), path.is_file()))

    def fake_run(document, settings, **kwargs):
        if document.pdf_path.name.startswith("a_bad"):
            raise ParserError("mineru", "run", "bad PDF")
        return run_document(document, settings, **kwargs)

    monkeypatch.setattr("paperfacts.batch.run_document", fake_run)
    monkeypatch.setattr("paperfacts.batch.write_dataset", capture_write)
    result = run_batch(source, Settings(data_root=tmp_path / "data"), output=output)

    assert len(result.documents) == len(result.failures) == 1
    assert result.failures[0]["filename"] == "a_bad.pdf"
    assert "bad PDF" in result.failures[0]["error"]
    assert snapshots == [(0, 1, True), (1, 1, True)]
    workbook = load_workbook(output)
    assert "a_bad.pdf" in str(list(workbook["运行记录"].values))
    workbook.close()


def test_a_paper_with_an_unanswered_question_is_listed_as_incomplete(monkeypatch, tmp_path):
    # Its rows are written (the unanswered field refused in both lanes), but it is not a success: the next
    # run asks that question again.
    source = make_blank_pdf(tmp_path / "paper.pdf")
    install_fake_pipeline(monkeypatch, unanswered="thickness")
    output = tmp_path / "result.xlsx"

    cli = CliRunner().invoke(
        app, ["batch", str(source), "--output", str(output), "--data-root", str(tmp_path / "data")]
    )

    assert cli.exit_code == 0, cli.output
    assert "1 incomplete" in cli.output
    workbook = load_workbook(output)
    [run] = [row for row in workbook["运行记录"].iter_rows(min_row=2, values_only=True)]
    workbook.close()
    assert run[2] == "incomplete"
    assert "no valid answer to mineru:thickness" in run[-1]


def test_cli_run_says_when_its_result_is_not_kept(monkeypatch, tmp_path):
    source = make_blank_pdf(tmp_path / "paper.pdf")
    paired = SampleMatching(pairs=(SampleMatch(a_id="S0", b_id="S0", confidence=1.0, method="llm", justification="t"),))
    install_fake_pipeline(monkeypatch, unanswered="thickness", matching=paired)

    cli = CliRunner().invoke(app, ["run", str(source), "--data-root", str(tmp_path / "data")])

    assert cli.exit_code == 0, cli.output
    assert "Incomplete, not kept as finished: no valid answer to mineru:thickness" in cli.output
    # The comparison counts above print such a field as missing; the cells say what it really is.
    assert re.search(r"[1-9]\d* cells unanswered \(thickness\)", cli.output), cli.output


def test_offline_export_never_runs_parser_or_llm(monkeypatch, tmp_path):
    source = make_blank_pdf(tmp_path / "paper.pdf")
    settings = Settings(data_root=tmp_path / "data")
    install_fake_pipeline(monkeypatch)
    document = DocumentInput.from_path(source)
    dataset = run_document(document, settings).dataset
    monkeypatch.setattr("paperfacts.batch.export_document", lambda doc, cfg: dataset)

    def unexpected(*args, **kwargs):
        pytest.fail("offline export tried to call the pipeline")

    monkeypatch.setattr("paperfacts.batch.run_document", unexpected)
    result = run_batch(source, settings, export_only=True)
    assert result.documents == (dataset,)
    assert result.excel_path.is_file()


def test_write_failure_propagates_instead_of_claiming_batch_success(monkeypatch, tmp_path):
    source = make_blank_pdf(tmp_path / "paper.pdf")
    install_fake_pipeline(monkeypatch)
    settings = Settings(data_root=tmp_path / "data")
    dataset = run_document(DocumentInput.from_path(source), settings).dataset
    monkeypatch.setattr("paperfacts.batch.export_document", lambda doc, cfg: dataset)

    def fail_write(*args, **kwargs):
        raise PermissionError("workbook is locked")

    monkeypatch.setattr("paperfacts.batch.write_dataset", fail_write)
    with pytest.raises(PermissionError, match="locked"):
        run_batch(source, settings, export_only=True)


def test_bad_output_suffix_fails_before_parser_runs(monkeypatch, tmp_path):
    source = make_blank_pdf(tmp_path / "paper.pdf")
    spy = install_fake_pipeline(monkeypatch)
    with pytest.raises(ConfigError, match="xlsx"):
        run_batch(source, Settings(), output=tmp_path / "output.csv")
    assert spy.parse == []


def test_cli_batch_produces_workbook_and_reports_path(monkeypatch, tmp_path):
    source = make_blank_pdf(tmp_path / "paper.pdf")
    install_fake_pipeline(monkeypatch)
    output = tmp_path / "result.xlsx"
    result = CliRunner().invoke(
        app, ["batch", str(source), "--output", str(output), "--data-root", str(tmp_path / "data")]
    )
    assert result.exit_code == 0, result.output
    assert "Completed: 1 papers" in result.output
    assert str(output) in result.output
    assert output.is_file()


def test_cli_offline_export_reports_missing_cache_as_failure(tmp_path):
    source = make_blank_pdf(tmp_path / "paper.pdf")
    output = tmp_path / "result.xlsx"
    result = CliRunner().invoke(
        app, ["export", str(source), "--output", str(output), "--data-root", str(tmp_path / "data")]
    )
    assert result.exit_code == 1, result.output
    assert "failed: 1" in result.output
    assert output.is_file()
