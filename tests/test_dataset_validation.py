"""How the VLM's verdicts reach the consolidated table -- and, just as important, how they do not.

Three entry points, each pinned: a contradicted value is set aside before anything else is judged; a
conflict in which exactly the surviving side was confirmed becomes ``vlm_resolved``; a value grounding could
not locate is trusted when the VLM located it. Everything else about the two-lane rules is unchanged, and a
verdict about evidence the dataset never held is silently irrelevant rather than wrongly applied.

A fill -- a value quoted from the VLM's transcription of a table -- has one entry point of its own: a cell
that would otherwise be blank, committed as ``vlm_filled``. It never replaces a value a lane read.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from paperfacts.compare import compare_lanes
from paperfacts.dataset import DatasetPayload, DocumentDataset, consolidate_document, write_dataset
from paperfacts.matching import SampleMatch, SampleMatching
from paperfacts.models import DocumentInput, NormalizedBBox
from paperfacts.records import FieldValue, TargetRecord
from paperfacts.validate import FilledValue, RegionCrop, ValidationReport, ValueValidation, Verdict, value_key
from support.extraction import DEFAULT_EXTRACTOR_KEY, make_lane, make_sample
from support.factories import DOC_ID

DOCUMENT = DocumentInput(document_id=DOC_ID, sha256=DOC_ID, pdf_path=Path("paper.pdf"))
VALIDATION_KEY = "vvvvvvvvvvvv"


def value(name: str, raw: str, unit: str | None = None, *, backend: str = "mineru", **kwargs) -> FieldValue:
    return FieldValue(field=name, value_raw=raw, unit_raw=unit, source_ids=(f"{backend}_p0_b1",), **kwargs)


def paired(a: list[FieldValue], b: list[FieldValue]):
    """Two lanes, one matched sample each, and the report; the caller adds verdicts."""
    lane_a = make_lane(samples=[make_sample("A", a)])
    lane_b = make_lane(backend="paddleocr_vl", samples=[make_sample("A", b)])
    matching = SampleMatching(
        pairs=(SampleMatch(a_id="A", b_id="A", confidence=1.0, method="exact", justification=""),)
    )
    return {"mineru": lane_a, "paddleocr_vl": lane_b}, compare_lanes(lane_a, lane_b, matching)


def verdicts(report, *entries: tuple[str, str, FieldValue, Verdict]) -> ValidationReport:
    """A validation report over ``(backend, owner, value, verdict)`` entries, keyed as the stage keys them."""
    return ValidationReport(
        document_id=DOC_ID,
        extractor_key=report.extractor_key,
        comparison_key=report.comparison_key,
        validation_key=VALIDATION_KEY,
        model="fake-vlm",
        policy="disputed",
        values=tuple(
            ValueValidation(
                key=value_key(backend, owner, field),
                backend=backend,
                owner=owner,
                field=field.field,
                value_raw=field.value_raw,
                unit_raw=field.unit_raw,
                condition=field.condition,
                source_ids=field.source_ids,
                reason="conflict",
                verdict=verdict,
                transcription="what the page says",
            )
            for backend, owner, field, verdict in entries
        ),
    )


def decision(result: DocumentDataset, field: str):
    return next(row for row in result.quality_rows if row["field"] == field and row["sample_id"] == "A")


CROP = RegionCrop(
    page=0,
    bbox=NormalizedBBox(x1=0.1, y1=0.4, x2=0.9, y2=0.6),
    dpi=72,
    width_px=100,
    height_px=20,
    source_ids=("mineru_p0_b1",),
    image_sha256="0" * 64,
    path="p000_x_72dpi.png",
)


def fill(
    name: str, raw: str, unit: str | None = None, *, backend: str = "mineru", owner: str = "sample:A"
) -> FilledValue:
    """A fill as the stage stores it: quoted from the transcription, cited to the crop."""
    field = FieldValue(field=name, value_raw=raw, unit_raw=unit, source_ids=("vlm:" + CROP.path,))
    return FilledValue(
        key=value_key(backend, owner, field),
        backend=backend,
        owner=owner,
        field=name,
        value_raw=raw,
        unit_raw=unit,
        source_id="vlm:" + CROP.path,
        crop=CROP,
        transcription="Sample | Thickness\nA | " + raw,
    )


def filled(report, *fills: FilledValue, entries: tuple = ()) -> ValidationReport:
    return verdicts(report, *entries).model_copy(update={"fills": fills})


# ---- A conflict the page settles -------------------------------------------------------------------------------


def test_a_conflict_with_one_side_confirmed_and_the_other_contradicted_is_resolved():
    a, b = value("thickness", "250", "nm"), value("thickness", "2500", "nm", backend="paddleocr_vl")
    lanes, report = paired([a], [b])
    assert [c.status for c in report.comparisons] == ["conflict"]
    validation = verdicts(
        report, ("mineru", "sample:A", a, "confirmed"), ("paddleocr_vl", "sample:A", b, "contradicted")
    )

    result = consolidate_document(DOCUMENT, lanes, report, validation)

    row = decision(result, "thickness")
    assert row["decision"] == "vlm_resolved"
    assert row["value"] == 250
    assert row["lanes"] == "mineru"
    assert row["vlm"] == "mineru: confirmed; paddleocr_vl: contradicted"
    assert "视觉核验裁决" in row["detail"]
    assert result.paper_row["thickness"] == 250


def test_a_conflict_where_both_sides_were_confirmed_stays_a_conflict():
    # The page really does say both numbers (two readings the parsers each caught one of): no verdict can
    # pick between them, and the two-lane rule stands.
    a, b = value("thickness", "250", "nm"), value("thickness", "2500", "nm", backend="paddleocr_vl")
    lanes, report = paired([a], [b])
    validation = verdicts(report, ("mineru", "sample:A", a, "confirmed"), ("paddleocr_vl", "sample:A", b, "confirmed"))

    row = decision(consolidate_document(DOCUMENT, lanes, report, validation), "thickness")

    assert row["decision"] == "conflict" and row["value"] is None


def test_a_conflict_whose_survivor_was_never_checked_stays_a_conflict():
    a, b = value("thickness", "250", "nm"), value("thickness", "2500", "nm", backend="paddleocr_vl")
    lanes, report = paired([a], [b])
    validation = verdicts(report, ("paddleocr_vl", "sample:A", b, "contradicted"))

    row = decision(consolidate_document(DOCUMENT, lanes, report, validation), "thickness")

    # One side denied, the other unread: setting the denied side aside leaves an unverified single value
    # inside a conflict, which is not the shape that resolves.
    assert row["decision"] == "conflict" and row["value"] is None
    assert row["vlm"] == "paddleocr_vl: contradicted"


def test_a_conflict_with_no_verdicts_is_exactly_what_it_was_before():
    a, b = value("thickness", "250", "nm"), value("thickness", "2500", "nm", backend="paddleocr_vl")
    lanes, report = paired([a], [b])

    without = consolidate_document(DOCUMENT, lanes, report)
    with_empty = consolidate_document(DOCUMENT, lanes, report, verdicts(report))

    assert decision(without, "thickness")["decision"] == "conflict"
    assert decision(with_empty, "thickness")["decision"] == "conflict"
    assert decision(with_empty, "thickness")["vlm"] == ""


# ---- Contradictions come first -----------------------------------------------------------------------------------


def test_an_agreement_both_lanes_share_but_the_page_denies_is_refused():
    # The failure this stage exists for: both parsers read the same wrong number.
    a, b = value("thickness", "250", "nm"), value("thickness", "250", "nm", backend="paddleocr_vl")
    lanes, report = paired([a], [b])
    assert [c.status for c in report.comparisons] == ["agree"]
    validation = verdicts(
        report, ("mineru", "sample:A", a, "contradicted"), ("paddleocr_vl", "sample:A", b, "contradicted")
    )

    row = decision(consolidate_document(DOCUMENT, lanes, report, validation), "thickness")

    assert row["decision"] == "vlm_contradicted"
    assert row["value"] is None
    assert row["vlm"] == "mineru: contradicted; paddleocr_vl: contradicted"


def test_a_single_source_value_the_page_denies_is_refused():
    a = value("thickness", "250", "nm")
    lanes, report = paired([a], [])
    validation = verdicts(report, ("mineru", "sample:A", a, "contradicted"))

    row = decision(consolidate_document(DOCUMENT, lanes, report, validation), "thickness")

    assert row["decision"] == "vlm_contradicted" and row["value"] is None


def test_an_agreement_with_one_side_denied_falls_back_to_the_confirmed_side_alone():
    a, b = value("thickness", "250", "nm"), value("thickness", "250", "nm", backend="paddleocr_vl")
    lanes, report = paired([a], [b])
    validation = verdicts(
        report, ("mineru", "sample:A", a, "confirmed"), ("paddleocr_vl", "sample:A", b, "contradicted")
    )

    row = decision(consolidate_document(DOCUMENT, lanes, report, validation), "thickness")

    # Still committed -- the survivor is trusted evidence -- but no longer as a two-lane agreement.
    assert row["decision"] == "single_source"
    assert row["value"] == 250
    assert row["lanes"] == "mineru"


# ---- The appeal against a grounding false negative ------------------------------------------------------------


def test_an_ungrounded_value_the_page_confirms_is_trusted():
    a = value("thickness", "250", "nm", grounded=False)
    lanes, report = paired([a], [])
    assert decision(consolidate_document(DOCUMENT, lanes, report), "thickness")["decision"] == "ungrounded"
    validation = verdicts(report, ("mineru", "sample:A", a, "confirmed"))

    row = decision(consolidate_document(DOCUMENT, lanes, report, validation), "thickness")

    assert row["decision"] == "single_source"
    assert row["value"] == 250
    assert "视觉核验确认" in row["detail"]
    assert "原文定位失败但视觉核验确认" in row["detail"]


def test_an_ungrounded_value_the_page_could_not_read_stays_untrusted():
    a = value("thickness", "250", "nm", grounded=False)
    lanes, report = paired([a], [])
    validation = verdicts(report, ("mineru", "sample:A", a, "illegible"))

    row = decision(consolidate_document(DOCUMENT, lanes, report, validation), "thickness")

    assert row["decision"] == "ungrounded" and row["value"] is None
    assert row["vlm"] == "mineru: illegible"


def test_a_value_without_any_citation_is_not_rescued_by_a_verdict():
    # No citation means no region was ever rendered; a "confirmed" for it cannot exist honestly, and if
    # one were forged the dataset would still refuse the value.
    a = value("thickness", "250", "nm").model_copy(update={"source_ids": (), "grounded": False})
    lanes, report = paired([a], [])
    validation = verdicts(report, ("mineru", "sample:A", a, "confirmed"))

    row = decision(consolidate_document(DOCUMENT, lanes, report, validation), "thickness")

    assert row["decision"] == "ungrounded" and row["value"] is None


# ---- Target (paper-level) values and the key contract ---------------------------------------------------------


def test_target_values_are_keyed_under_the_target_owner():
    a = TargetRecord(fields=(value("component", "SnO2:Ta"),))
    b = TargetRecord(fields=(value("component", "SnO2:Ta", backend="paddleocr_vl"),))
    lane_a = make_lane(target=a)
    lane_b = make_lane(backend="paddleocr_vl", target=b)
    report = compare_lanes(lane_a, lane_b, SampleMatching())
    validation = verdicts(
        report,
        ("mineru", "target", a.fields[0], "contradicted"),
        ("paddleocr_vl", "target", b.fields[0], "contradicted"),
    )

    result = consolidate_document(DOCUMENT, {"mineru": lane_a, "paddleocr_vl": lane_b}, report, validation)

    row = next(row for row in result.quality_rows if row["field"] == "component")
    assert row["decision"] == "vlm_contradicted"


def test_a_verdict_keyed_to_the_wrong_owner_does_not_apply():
    a = value("thickness", "250", "nm")
    lanes, report = paired([a], [])
    validation = verdicts(report, ("mineru", "sample:B", a, "contradicted"))

    row = decision(consolidate_document(DOCUMENT, lanes, report, validation), "thickness")

    assert row["decision"] == "single_source" and row["value"] == 250
    assert row["vlm"] == ""


# ---- Keys travel with the table ---------------------------------------------------------------------------------


def test_the_dataset_records_the_validation_key_it_was_built_with():
    a = value("thickness", "250", "nm")
    lanes, report = paired([a], [])

    without = consolidate_document(DOCUMENT, lanes, report)
    with_verdicts = consolidate_document(DOCUMENT, lanes, report, verdicts(report))

    assert without.validation_key == ""
    assert with_verdicts.validation_key == VALIDATION_KEY
    assert DocumentDataset.from_payload(with_verdicts.to_payload()).validation_key == VALIDATION_KEY
    assert (
        DatasetPayload.model_validate_json(with_verdicts.to_payload().model_dump_json()).validation_key
        == VALIDATION_KEY
    )


def test_a_validation_built_under_other_keys_is_refused():
    a = value("thickness", "250", "nm")
    lanes, report = paired([a], [])
    stale = verdicts(report).model_copy(update={"extractor_key": "ffffffffffff"})

    with pytest.raises(ValueError, match="different keys"):
        consolidate_document(DOCUMENT, lanes, report, stale)
    assert report.extractor_key == DEFAULT_EXTRACTOR_KEY


def test_the_verdict_column_and_the_validation_key_reach_the_workbook(tmp_path: Path):
    a, b = value("thickness", "250", "nm"), value("thickness", "2500", "nm", backend="paddleocr_vl")
    lanes, report = paired([a], [b])
    validation = verdicts(
        report, ("mineru", "sample:A", a, "confirmed"), ("paddleocr_vl", "sample:A", b, "contradicted")
    )
    result = consolidate_document(DOCUMENT, lanes, report, validation)
    output = tmp_path / "out.xlsx"

    write_dataset([result], output)

    workbook = load_workbook(output)
    quality = workbook["数据质量"]
    header = [cell.value for cell in quality[1]]
    assert "视觉核验" in header
    vlm_column = header.index("视觉核验")
    decision_column = header.index("最终决策")
    rows = {row[decision_column].value: row[vlm_column].value for row in quality.iter_rows(min_row=2)}
    assert rows["vlm_resolved"] == "mineru: confirmed; paddleocr_vl: contradicted"
    runs = workbook["运行记录"]
    run_header = [cell.value for cell in runs[1]]
    assert runs[2][run_header.index("视觉核验版本")].value == VALIDATION_KEY


# ---- Fills: only into a cell that would otherwise be blank -------------------------------------------------------


def test_a_blank_cell_is_filled_from_the_table_transcription():
    a = value("sheet_resistance", "12.5", "Ω/sq")
    lanes, report = paired([a], [])
    assert decision(consolidate_document(DOCUMENT, lanes, report), "thickness")["decision"] == "missing"
    validation = filled(report, fill("thickness", "250", "nm"))

    result = consolidate_document(DOCUMENT, lanes, report, validation)

    row = decision(result, "thickness")
    assert row["decision"] == "vlm_filled"
    assert row["value"] == 250
    assert row["lanes"] == "mineru"
    assert row["source_ids"] == "vlm:" + CROP.path
    assert row["vlm"] == "mineru: filled"
    assert "表格转写补全" in row["detail"]
    assert result.paper_row["thickness"] == 250
    # The cell the lane did read is untouched.
    assert decision(result, "sheet_resistance")["decision"] == "single_source"


def test_a_fill_never_replaces_a_value_a_lane_read():
    a = value("thickness", "250", "nm")
    lanes, report = paired([a], [])
    validation = filled(report, fill("thickness", "999", "nm"))

    row = decision(consolidate_document(DOCUMENT, lanes, report, validation), "thickness")

    assert row["decision"] == "single_source" and row["value"] == 250
    assert "filled" not in row["vlm"]


def test_a_fill_does_not_settle_a_conflict():
    # Two readings that disagree are a disagreement to review; a third number from the table transcription
    # is not a tiebreaker, and the cell stays empty with the conflict named.
    a, b = value("thickness", "250", "nm"), value("thickness", "2500", "nm", backend="paddleocr_vl")
    lanes, report = paired([a], [b])
    validation = filled(report, fill("thickness", "250", "nm"))

    row = decision(consolidate_document(DOCUMENT, lanes, report, validation), "thickness")

    assert row["decision"] == "conflict" and row["value"] is None


def test_a_cell_whose_every_reading_was_denied_takes_the_fill():
    # Both parsers misread the number; the VLM denied both and its table transcription carries the real
    # one: the fill is the only evidence left standing.
    a, b = value("thickness", "250", "nm"), value("thickness", "250", "nm", backend="paddleocr_vl")
    lanes, report = paired([a], [b])
    validation = filled(
        report,
        fill("thickness", "2500", "nm"),
        entries=(("mineru", "sample:A", a, "contradicted"), ("paddleocr_vl", "sample:A", b, "contradicted")),
    )

    row = decision(consolidate_document(DOCUMENT, lanes, report, validation), "thickness")

    assert row["decision"] == "vlm_filled" and row["value"] == 2500
    assert row["vlm"] == "mineru: contradicted; paddleocr_vl: contradicted; mineru: filled"


def test_the_first_lane_in_backend_order_supplies_the_fill_when_both_have_one():
    a = value("sheet_resistance", "12.5", "Ω/sq")
    lanes, report = paired([a], [value("sheet_resistance", "12.5", "Ω/sq", backend="paddleocr_vl")])
    validation = filled(report, fill("thickness", "300", "nm", backend="paddleocr_vl"), fill("thickness", "250", "nm"))

    row = decision(consolidate_document(DOCUMENT, lanes, report, validation), "thickness")

    assert row["decision"] == "vlm_filled" and row["value"] == 250 and row["lanes"] == "mineru"


def test_a_fill_keyed_to_a_sample_the_scope_does_not_own_does_not_apply():
    a = value("sheet_resistance", "12.5", "Ω/sq")
    lanes, report = paired([a], [])
    validation = filled(report, fill("thickness", "250", "nm", owner="sample:B"))

    row = decision(consolidate_document(DOCUMENT, lanes, report, validation), "thickness")

    assert row["decision"] == "missing" and row["value"] is None


def test_a_fill_that_is_not_a_scalar_is_refused_like_any_other_value():
    a = value("sheet_resistance", "12.5", "Ω/sq")
    lanes, report = paired([a], [])
    validation = filled(report, fill("thickness", "thick", "nm"))

    row = decision(consolidate_document(DOCUMENT, lanes, report, validation), "thickness")

    assert row["decision"] == "non_scalar" and row["value"] is None


def test_a_filled_cell_reaches_the_workbook_with_its_decision(tmp_path: Path):
    a = value("sheet_resistance", "12.5", "Ω/sq")
    lanes, report = paired([a], [])
    result = consolidate_document(DOCUMENT, lanes, report, filled(report, fill("thickness", "250", "nm")))
    output = tmp_path / "out.xlsx"

    write_dataset([result], output)

    quality = load_workbook(output)["数据质量"]
    header = [cell.value for cell in quality[1]]
    decisions = {row[header.index("最终决策")].value for row in quality.iter_rows(min_row=2)}
    assert "vlm_filled" in decisions
