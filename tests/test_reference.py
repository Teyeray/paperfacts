"""The ``reference`` kind (round 2 spec §5.7): a field of one entity type naming a sample of another -- the coating
a wear test ran on.

The value is a sample id copied from the referenced entity's list. It is grounded when it resolves to one of the
lane's samples of that entity, and that resolution must survive ``workflow.read_lane``, which re-grounds every
value on every read: the tests of it go through ``read_lane``. Two lanes agree when the referenced entity's matching
pairs the samples they name; the cell is the id of the dataset row those samples became.
"""

from __future__ import annotations

import importlib.util
import json
import re
import shutil
import sys
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from paperfacts.columns import field_columns
from paperfacts.compare import compare_lanes
from paperfacts.config import DEFAULT_REPO_ROOT, Settings
from paperfacts.dataset import consolidate_document
from paperfacts.errors import ConfigError
from paperfacts.grounding import ground_lane
from paperfacts.keys import ComparisonOptions, profile_extraction_fingerprint
from paperfacts.kinds import KindContext, rules_for
from paperfacts.matching import SampleMatch, SampleMatching
from paperfacts.models import BACKENDS, DocumentGeometry, DocumentInput
from paperfacts.normalize import normalize_lane
from paperfacts.profile_loader import parse_profile
from paperfacts.prompts import field_user_prompt
from paperfacts.records import FieldValue, LaneExtraction, SampleRecord
from paperfacts.storage import DataLayout
from paperfacts.workbook import write_dataset
from paperfacts.workflow import load_run_profile, read_lane, run_document
from support.extraction import make_artifact
from support.factories import DOC_ID, RawOutputFactory, make_block, paddle_page_entry
from support.llm import FakeLlmClient
from support.profiles import make_entity_profile, make_reference_profile, reference_profile_data

_score_spec = importlib.util.spec_from_file_location(
    "paperfacts_score_reference", DEFAULT_REPO_ROOT / "eval" / "score.py"
)
assert _score_spec is not None and _score_spec.loader is not None
score = importlib.util.module_from_spec(_score_spec)
sys.modules[_score_spec.name] = score
_score_spec.loader.exec_module(score)

PROFILE = make_reference_profile()
COATING, WEAR = PROFILE.entities
SPEC = PROFILE.by_name["tested_coating"]
OPTIONS = ComparisonOptions(profile=PROFILE, ambiguous_match_confidence=0.6)
KEY = "0123456789ab"


# ---- The profile ------------------------------------------------------------------------------------------------


def _refused(change: dict[str, Any], *, entities: bool = True) -> str:
    data = reference_profile_data()
    data["fields"][-1].update(change)
    if not entities:
        data = {key: value for key, value in data.items() if key != "entities"}
        data["groups"] = [{k: v for k, v in group.items() if k != "entity"} for group in data["groups"]]
    with pytest.raises(ConfigError) as caught:
        parse_profile(data, Path("profiles/demo.json"))
    return str(caught.value)


def test_a_reference_field_describes_one_entity_and_names_another():
    assert (SPEC.kind, SPEC.entity, SPEC.references, SPEC.level) == ("reference", "wear_test", "coating", "sample")
    assert all(spec.references is None for spec in PROFILE.fields if spec is not SPEC)


@pytest.mark.parametrize(
    ("change", "message"),
    [
        ({"references": None}, "references names the entity type of a reference field"),
        ({"kind": "text"}, "references names the entity type of a reference field"),
        ({"references": "catalyst"}, "references must name one of the declared entities (coating, wear_test)"),
        ({"references": "wear_test"}, "must name another entity type than the field's own, 'wear_test'"),
        ({"group": "precursor"}, "a reference field belongs to a sample group of a declared entity type"),
        ({"cardinality": "many"}, "cardinality 'many' needs a text or composition field, not a 'reference' one"),
        ({"canonical_unit": "nm"}, "canonical_unit is only meaningful"),
        ({"categories": ["A", "B"]}, "categories is only meaningful for a text field"),
        (
            {"condition_rule": "the load", "condition_hint": "the load", "missing_condition_note_zh": "未注明载荷"},
            "condition_rule is only meaningful",
        ),
    ],
)
def test_the_loader_refuses_a_malformed_reference(change, message):
    assert message in _refused(change)


def test_a_reference_needs_entity_types():
    assert "a reference field belongs to a sample group of a declared entity type" in _refused({}, entities=False)


# ---- The question -----------------------------------------------------------------------------------------------


def test_only_a_reference_question_shows_the_referenced_list_and_its_note():
    question = field_user_prompt(
        SPEC, "- id: W1", "EXCERPTS", "x", WEAR.prompt.sample_list_heading, (COATING, "- id: C1")
    )

    assert 'Copy the id of the referenced coating exactly from the list "Coatings" below.' in question
    assert question.index("Wear tests this paper reports:\n- id: W1\n\n") < question.index(
        "Coatings this paper reports:\n- id: C1\n\nExcerpts"
    )
    # Any other field of the same profile is asked exactly as before the reference existed.
    other = make_entity_profile().by_name["wear_mode"]
    assert field_user_prompt(PROFILE.by_name["wear_mode"], "- id: W1", "E", "x", "Wear tests") == field_user_prompt(
        other, "- id: W1", "E", "x", "Wear tests"
    )
    assert "Copy the id" not in field_user_prompt(PROFILE.by_name["wear_mode"], "- id: W1", "E", "x", "Wear tests")


# ---- Grounding survives read_lane --------------------------------------------------------------------------------


def _reference(value_raw: str, source: str, *, grounded: bool = True) -> FieldValue:
    return FieldValue(field="tested_coating", value_raw=value_raw, source_ids=(source,), grounded=grounded)


def _lane(backend: str, coatings: tuple[str, ...], tests: dict[str, FieldValue | None]) -> LaneExtraction:
    return LaneExtraction(
        document_id=DOC_ID,
        backend=backend,  # type: ignore[arg-type]
        extractor_key=KEY,
        model="m",
        profile_fingerprint=profile_extraction_fingerprint(PROFILE),
        samples=(
            *(SampleRecord(sample_id=sample_id, entity="coating") for sample_id in coatings),
            *(
                SampleRecord(sample_id=sample_id, entity="wear_test", fields=(value,) if value else ())
                for sample_id, value in tests.items()
            ),
        ),
    )


# MinerU's blocks write "C9" and never "C1"; PaddleOCR-VL's never write "coat-1". A text check would therefore ground
# exactly the wrong references: only resolution among the lane's coatings gets them right.
MINERU_TEXT = "Coating C9 was mentioned in passing; the wear tests W1 and W2 ran at 300 °C."
PADDLE_TEXT = "The wear runs were done at 300 °C."
# Stored with the opposite verdicts, so a read that kept or text-checked them would be caught.
STORED = {
    "mineru": _lane(
        "mineru",
        ("C1",),
        {"W1": _reference("C1", "mineru_p0_b0", grounded=False), "W2": _reference("C9", "mineru_p0_b0")},
    ),
    "paddleocr_vl": _lane(
        "paddleocr_vl",
        ("coat-1",),
        {"run-1": _reference("coat 1", "paddleocr_vl_p0_b0", grounded=False), "run-2": None},
    ),
}
MATCHINGS = {
    "coating": SampleMatching(
        pairs=(SampleMatch(a_id="C1", b_id="coat-1", confidence=0.9, justification="j", method="llm"),)
    ),
    "wear_test": SampleMatching(
        pairs=(
            SampleMatch(a_id="W1", b_id="run-1", confidence=0.9, justification="j", method="llm"),
            SampleMatch(a_id="W2", b_id="run-2", confidence=0.9, justification="j", method="llm"),
        )
    ),
}


def _read(tmp_path: Path) -> dict[str, LaneExtraction]:
    layout = DataLayout(tmp_path)
    texts = {"mineru": MINERU_TEXT, "paddleocr_vl": PADDLE_TEXT}
    lanes = {}
    for backend, lane in STORED.items():
        lane.write(layout.extraction_path(DOC_ID, backend, KEY))  # type: ignore[arg-type]
        block = make_block(backend=backend, content=texts[backend])  # type: ignore[arg-type]
        make_artifact((block,), backend=backend).write(layout.artifact_path(DOC_ID, backend))  # type: ignore[arg-type]
        read = read_lane(layout, DOC_ID, backend, KEY, PROFILE)  # type: ignore[arg-type]
        assert read is not None
        lanes[backend] = read
    return lanes


def test_read_lane_grounds_a_reference_by_resolving_it_among_the_lanes_samples(tmp_path: Path):
    lanes = _read(tmp_path)

    mineru, paddle = lanes["mineru"], lanes["paddleocr_vl"]
    resolving = mineru.sample("W1", "wear_test").get("tested_coating")  # type: ignore[union-attr]
    unlisted = mineru.sample("W2", "wear_test").get("tested_coating")  # type: ignore[union-attr]
    # "coat 1" is the listed "coat-1" by sample_key, the key every sample is identified by.
    spelled = paddle.sample("run-1", "wear_test").get("tested_coating")  # type: ignore[union-attr]
    assert (resolving.grounded, resolving.ref_id, resolving.bound) == (True, "C1", None)  # type: ignore[union-attr]
    assert (unlisted.grounded, unlisted.ref_id) == (False, None)  # type: ignore[union-attr]
    assert (spelled.grounded, spelled.ref_id) == (True, "coat-1")  # type: ignore[union-attr]
    # Grounding and reading resolve alike: a reference is grounded exactly when it names a listed sample.
    for lane in lanes.values():
        for sample in lane.samples:
            for value in sample.fields:
                assert value.grounded == (value.ref_id is not None)
    # ref_id is read, never stored: the lane file carries only the model's own words.
    assert "ref_id" not in STORED["mineru"].model_dump_json()


def test_the_dataset_commits_a_resolving_reference_as_its_row_id_and_refuses_an_unlisted_one(tmp_path: Path):
    lanes = _read(tmp_path)
    report = compare_lanes(lanes["mineru"], lanes["paddleocr_vl"], MATCHINGS, OPTIONS)
    document = DocumentInput(document_id=DOC_ID, sha256=DOC_ID, pdf_path=Path("paper.pdf"))

    dataset = consolidate_document(document, lanes, report, OPTIONS)

    rows = {(row["entity"], row["sample_id"]): row for row in dataset.sample_rows}
    # The coating row is "C1 | coat-1", and the wear test names exactly that row.
    assert ("coating", "C1 | coat-1") in rows
    assert rows[("wear_test", "W1 | run-1")]["tested_coating"] == "C1 | coat-1"
    quality = {(row["sample_id"], row["field"]): row for row in dataset.quality_rows}
    assert quality[("W1 | run-1", "tested_coating")]["decision"] == "agree"
    assert rows[("wear_test", "W2 | run-2")]["tested_coating"] is None
    assert quality[("W2 | run-2", "tested_coating")]["decision"] == "ungrounded"


# ---- Comparison through the referenced entity's matching --------------------------------------------------------


def _pair(a: str, b: str, coatings_a: tuple[str, ...], coatings_b: tuple[str, ...], coating: SampleMatching):
    lane_a = normalize_lane(_lane("mineru", coatings_a, {"W1": _reference(a, "b")}), PROFILE)
    lane_b = normalize_lane(_lane("paddleocr_vl", coatings_b, {"run-1": _reference(b, "b")}), PROFILE)
    report = compare_lanes(lane_a, lane_b, {"coating": coating, "wear_test": MATCHINGS["wear_test"]}, OPTIONS)
    return next((c.status, c.detail) for c in report.comparisons if c.field == "tested_coating")


PAIRED = SampleMatching(
    pairs=(
        SampleMatch(a_id="C1", b_id="coat-1", confidence=0.9, justification="j", method="llm"),
        SampleMatch(a_id="C2", b_id="coat-2", confidence=0.9, justification="j", method="llm"),
    )
)


def test_two_lanes_agree_exactly_when_the_referenced_matching_pairs_their_samples():
    coatings_a, coatings_b = ("C1", "C2"), ("coat-1", "coat-2")

    assert _pair("C1", "coat-1", coatings_a, coatings_b, PAIRED)[0] == "agree"
    # Both resolve, to coatings the matching pairs with others: two different coatings.
    assert _pair("C1", "coat-2", coatings_a, coatings_b, PAIRED)[0] == "conflict"
    # One names no listed coating.
    assert _pair("C1", "coat-9", coatings_a, coatings_b, PAIRED)[0] == "ambiguous"
    # One names a coating its matching left unpaired: whether it is the other's is not known.
    unpaired = SampleMatching(
        pairs=(SampleMatch(a_id="C1", b_id="coat-1", confidence=0.9, justification="j", method="llm"),),
        unmatched_a=("C2",),
        unmatched_b=("coat-2",),
    )
    assert _pair("C2", "coat-2", coatings_a, coatings_b, unpaired)[0] == "ambiguous"


def test_the_kind_row_needs_its_context_and_ignores_every_other():
    rules = rules_for(SPEC)
    value = FieldValue(field="tested_coating", value_raw="C1", ref_id="C1")

    assert rules.cell(value, SPEC, PROFILE.units) == (None, "引用的coating样品 'C1' 没有数据行")
    row_ids = KindContext(row_ids={("mineru", "coating", "C1"): "C1 | coat-1"}, backend="mineru")
    assert rules.cell(value, SPEC, PROFILE.units, row_ids) == ("C1 | coat-1", None)
    assert rules.distance(value, value) is None


# ---- End to end, with a second run that reads both lanes back through read_lane ---------------------------------

TITLE = "Wear of sol-gel coatings"
PARAGRAPHS = (
    "Two coatings were deposited from an ethanol solvent; the coating thickness is 100 nm and 200 nm.",
    "The first coating was tested at 300 °C against a steel ball in sliding wear.",
)
INVENTORY = {
    ("mineru", "coatings"): ("C1", "C2"),
    ("mineru", "wear tests"): ("W1",),
    ("paddleocr_vl", "coatings"): ("coat-1", "coat-2"),
    ("paddleocr_vl", "wear tests"): ("run-1",),
}
_SOURCE = re.compile(r"<!-- source: (\S+) -->")
_LISTED_ID = re.compile(r"^- id: (.+?) \| label:", re.MULTILINE)
_FIELD = re.compile(r"\AField to extract:\n- `([^`]+)`")
_PLURAL = re.compile(r"list the (.+?) the paper reports")


def _respond(system: str, user: str) -> str:
    lane = "paddleocr_vl" if "paddleocr_vl_" in user else "mineru"
    cited = _SOURCE.findall(user)[:1]
    if user.startswith("Paper excerpts"):
        plural = _PLURAL.search(system)
        assert plural is not None
        samples = [{"sample_id": sid, "source_ids": cited} for sid in INVENTORY[(lane, plural.group(1))]]
        return json.dumps({"samples": samples, "no_samples_in_scope": False})
    if user.startswith("List A (parser:"):
        list_a, list_b = user.split("\n\nList B (parser:", 1)
        pairs = [
            {"a": a, "b": b, "confidence": 0.9, "justification": "same position"}
            for a, b in zip(_LISTED_ID.findall(list_a), _LISTED_ID.findall(list_b), strict=True)
        ]
        return json.dumps({"pairs": pairs})
    field = _FIELD.match(user)
    assert field is not None, user[:200]
    if field.group(1) != "tested_coating":
        return json.dumps({"values": []})
    # The first listed wear test ran on the first listed coating, whose id is only in the second list.
    tests, coatings = (_LISTED_ID.findall(part) for part in user.split("Coatings this paper reports:", 1))
    value = {"sample_id": tests[0], "value_raw": coatings[0], "source_ids": cited}
    return json.dumps({"values": [value]})


def _seed_parses(document: DocumentInput, layout: DataLayout, geometry: DocumentGeometry) -> None:
    blocks = (TITLE, *PARAGRAPHS)
    for backend in BACKENDS:
        raw_dir = layout.raw_dir(document.document_id, backend)
        factory = RawOutputFactory(raw_dir.parent, document, geometry)
        if backend == "mineru":
            content = [
                {"type": "text", "page_idx": 0, "bbox": [100, 80 + 150 * i, 900, 200 + 150 * i], "text": text}
                | ({"text_level": 1} if i == 0 else {})
                for i, text in enumerate(blocks)
            ]
            factory.mineru(content, dir_name=raw_dir.name)
        else:
            parsing = [
                {
                    "block_id": i,
                    "block_order": i,
                    "block_label": "doc_title" if i == 0 else "text",
                    "block_bbox": [200, 150 + 350 * i, 1450, 450 + 350 * i],
                    "block_content": text,
                }
                for i, text in enumerate(blocks)
            ]
            page = paddle_page_entry(0, {"res": {"parsing_res_list": parsing}}, (1654, 2339))
            factory.paddle([page], dir_name=raw_dir.name)


def test_a_reference_runs_end_to_end_and_still_fills_its_cell_when_the_lanes_are_read_back(
    tmp_path: Path, document: DocumentInput, geometry: DocumentGeometry, monkeypatch
):
    profile_file = tmp_path / "profiles" / "demo.json"
    profile_file.parent.mkdir()
    profile_file.write_text(json.dumps(reference_profile_data()), encoding="utf-8")
    settings = Settings(
        data_root=tmp_path / "data",
        repo_root=tmp_path,
        llm_api_key="sk-test",
        llm_model="fake-model",
        profile=str(profile_file),
    )
    _seed_parses(document, DataLayout(settings.data_root), geometry)
    profile = load_run_profile(settings)
    client = FakeLlmClient(_respond)
    monkeypatch.setattr("paperfacts.workflow.build_llm_client", lambda _settings: client)

    first = run_document(document, settings, profile)
    # The second run serves both lanes from their files, through read_lane, which re-grounds them; without its
    # stored comparison and table, both are derived again from what read_lane returned.
    for derived in ("comparisons", "datasets"):
        shutil.rmtree(DataLayout(settings.data_root).doc_dir(document.document_id) / derived, ignore_errors=True)
    second = run_document(document, settings, profile)

    questions = [call.user for call in client.calls if call.user.startswith("Field to extract:\n- `tested_coating`")]
    assert len(questions) == 2 and all("Coatings this paper reports:\n- id: " in q for q in questions)
    for result in (first, second):
        paddle = result.lanes["paddleocr_vl"].sample("run-1", "wear_test").get("tested_coating")  # type: ignore[union-attr]
        assert (paddle.grounded, paddle.ref_id) == (True, "coat-1")  # type: ignore[union-attr]
        rows = {(row["entity"], row["sample_id"]): row for row in result.dataset.sample_rows}
        assert rows[("wear_test", "W1 | run-1")]["tested_coating"] == "C1 | coat-1"
        comparison = next(c for c in result.report.comparisons if c.field == "tested_coating")
        assert comparison.status == "agree"


# ---- Surfaces: the column, the workbook, the score ---------------------------------------------------------------


def test_the_column_and_the_workbook_name_the_referenced_entity(tmp_path: Path):
    column = next(column for column in field_columns(PROFILE) if column.name == "tested_coating")
    assert (column.kind, column.entity, column.references) == ("reference", "wear_test", "coating")
    document = DocumentInput(document_id=DOC_ID, sha256=DOC_ID, pdf_path=Path("paper.pdf"))
    lanes = {backend: ground_lane(lane, {}, profile=PROFILE) for backend, lane in STORED.items()}
    dataset = consolidate_document(
        document, lanes, compare_lanes(lanes["mineru"], lanes["paddleocr_vl"], MATCHINGS, OPTIONS), OPTIONS
    )
    output = tmp_path / "out.xlsx"

    write_dataset([dataset], output, PROFILE)

    book = load_workbook(output)
    fields = [[cell.value for cell in row] for row in book["字段说明"].iter_rows()]
    row = next(row for row in fields if row[0] == "tested_coating")
    assert "涂层样品ID" in row
    tests = [[cell.value for cell in row] for row in book["磨损测试数据"].iter_rows()]
    column_index = tests[0].index("tested_coating")
    assert "C1 | coat-1" in [row[column_index] for row in tests[1:]]


def test_the_score_resolves_a_reference_to_the_row_aligned_to_the_named_gold_sample():
    specs = PROFILE.by_name
    gold = {
        "doc_id": "d",
        "samples": [
            {"id": "c1", "entity": "coating", "match": {"label": "^C1"}, "fields": {}},
            {"id": "c2", "entity": "coating", "match": {"label": "^C2"}, "fields": {}},
            {
                "id": "t1",
                "entity": "wear_test",
                "match": {"label": "^W"},
                "fields": {"tested_coating": [{"value": "c1"}]},
            },
        ],
    }
    rows = [
        {"entity": "coating", "sample_id": "C1 | coat-1"},
        {"entity": "coating", "sample_id": "C2 | coat-2"},
        {"entity": "wear_test", "sample_id": "W1 | run-1", "tested_coating": "C1 | coat-1"},
    ]

    cells = score.score_document(specs, gold, {"sample_rows": rows, "quality_rows": []})
    wrong = score.score_document(
        specs, gold, {"sample_rows": [*rows[:2], rows[2] | {"tested_coating": "C2 | coat-2"}], "quality_rows": []}
    )

    assert [(c.sample, c.outcome) for c in cells if c.field == "tested_coating"] == [("t1", "correct")]
    assert [(c.sample, c.outcome) for c in wrong if c.field == "tested_coating"] == [("t1", "wrong")]


def test_the_score_keeps_two_entities_gold_samples_of_one_id_apart():
    # A wear test named "1" must not shadow the coating "1" its reference names.
    gold = {
        "doc_id": "d",
        "samples": [
            {"id": "1", "entity": "coating", "match": {"label": "^C1"}, "fields": {}},
            {
                "id": "1",
                "entity": "wear_test",
                "match": {"label": "^W"},
                "fields": {"tested_coating": [{"value": "1"}]},
            },
        ],
    }
    rows = [
        {"entity": "coating", "sample_id": "C1 | coat-1"},
        {"entity": "wear_test", "sample_id": "W1 | run-1", "tested_coating": "C1 | coat-1"},
    ]

    cells = score.score_document(PROFILE.by_name, gold, {"sample_rows": rows, "quality_rows": []})

    assert [c.outcome for c in cells if c.field == "tested_coating"] == ["correct"]


def test_the_prompts_command_prints_a_reference_question(tmp_path: Path, monkeypatch):
    from typer.testing import CliRunner

    from paperfacts.cli import app

    profile_file = tmp_path / "profiles" / "demo.json"
    profile_file.parent.mkdir()
    profile_file.write_text(json.dumps(reference_profile_data()), encoding="utf-8")
    monkeypatch.setenv("PAPERFACTS_EXTRACTION_MODE", "passage")

    printed = CliRunner().invoke(app, ["prompts", "--profile", str(profile_file), "--field", "tested_coating"])

    assert printed.exit_code == 0, printed.output
    assert 'Copy the id of the referenced coating exactly from the list "Coatings" below.' in printed.output
    assert "Coatings this paper reports:\n<referenced sample list>" in printed.output
