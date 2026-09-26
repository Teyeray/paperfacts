"""The Excel export: what a consolidated dataset looks like once it is written to a workbook."""

from pathlib import Path

from openpyxl import load_workbook

from paperfacts.matching import SampleMatch, SampleMatching
from paperfacts.records import PaperRecord
from paperfacts.workbook import write_dataset
from support.extraction import make_lane, make_sample
from support.factories import DOC_ID
from support.profiles import make_profile
from test_dataset import dataset, paired, value


def test_excel_reopens_with_numeric_fields_text_ids_and_no_pdf_formulas(tmp_path, tco_profile):
    a = make_lane(
        samples=[
            make_sample(
                "001",
                [value("thickness", "0.3", "μm"), value("resistivity", "1e-4", "Ω cm")],
                label='=HYPERLINK("https://example.invalid")',
                conditions={"temperature": "001"},
            )
        ],
        paper=PaperRecord(fields=(value("component", "=1+1"),)),
    )
    b = make_lane(backend="paddleocr_vl", samples=[make_sample("001", a.samples[0].fields)], paper=a.paper)
    matching = SampleMatching(
        pairs=(SampleMatch(a_id="001", b_id="001", confidence=1.0, method="exact", justification="same"),)
    )
    result = dataset(a, b, matching, filename="=paper.pdf")
    output = tmp_path / "dataset.xlsx"
    write_dataset(
        [result, result],
        output,
        tco_profile,
        failures=[{"document_id": "b" * 64, "filename": "failed.pdf", "error": "parser failed"}],
    )
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["论文数据", "样品数据", "字段说明", "数据质量", "图中读数", "运行记录"]
    sheet = workbook["论文数据"]
    assert sheet.max_row == 2
    assert sheet.freeze_panes == "D2"
    assert len(sheet.tables) == 1
    columns = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, columns["thickness"]).value == 300
    assert sheet.cell(2, columns["thickness"]).data_type == "n"
    assert sheet.cell(2, columns["样品ID"]).value == "001"
    assert sheet.cell(2, columns["样品ID"]).data_type == "s"
    assert sheet.cell(2, columns["component"]).value == "=1+1"
    assert sheet.cell(2, columns["component"]).data_type == "s"
    assert sheet.cell(2, columns["文件名"]).data_type == "s"
    assert sheet.cell(2, columns["density"]).value is None
    assert "E+00" in sheet.cell(2, columns["resistivity"]).number_format
    assert workbook["运行记录"].max_row == 3
    assert workbook["运行记录"].cell(3, 3).value == "failed"
    assert not list(tmp_path.glob("*.tmp"))
    assert all(cell.data_type != "f" for page in workbook for row in page for cell in row)


def test_an_empty_export_still_records_failures(tmp_path, tco_profile):
    output = tmp_path / "empty.xlsx"
    write_dataset([], output, tco_profile, failures=[{"filename": "bad.pdf", "error": "unreadable"}])
    workbook = load_workbook(output)
    assert workbook["论文数据"].max_row == 1
    assert workbook["运行记录"].cell(2, 3).value == "failed"


def test_the_series_mark_reaches_the_quality_sheet(tmp_path, tco_profile):
    result = paired(
        [value("thickness", "300", "nm", series=True)],
        [value("thickness", "300", "nm", backend="paddleocr_vl", series=True)],
    )
    output = tmp_path / "dataset.xlsx"
    write_dataset([result], output, tco_profile)

    sheet = load_workbook(output)["数据质量"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    rows = {sheet.cell(row, columns["字段"]).value: row for row in range(2, sheet.max_row + 1)}
    assert sheet.cell(rows["thickness"], columns["系列级"]).value is True
    assert sheet.cell(rows["resistivity"], columns["系列级"]).value is False


def test_the_lane_column_reaches_the_quality_sheet(tmp_path, tco_profile):
    result = paired([value("thickness", "300", "nm")], [value("thickness", "300", "nm", backend="paddleocr_vl")])
    output = tmp_path / "dataset.xlsx"
    write_dataset([result], output, tco_profile)

    sheet = load_workbook(output)["数据质量"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    rows = {sheet.cell(row, columns["字段"]).value: row for row in range(2, sheet.max_row + 1)}
    assert sheet.cell(rows["thickness"], columns["证据来源通道"]).value == "mineru; paddleocr_vl"


def test_figure_rows_fill_only_their_own_sheet(tmp_path: Path, tco_profile):
    output = tmp_path / "dataset.xlsx"
    result = dataset(make_lane(samples=[make_sample("A", [value("thickness", "100", "nm")])]))
    row = {"document_id": DOC_ID, "figure": "Fig. 3", "value": None, "value_raw": "25 10^2 ohm/sq", "precision": "±20%"}

    write_dataset([result], output, tco_profile, figure_rows=[row])

    workbook = load_workbook(output)
    sheet = workbook["图中读数"]
    header = [cell.value for cell in sheet[1]]
    cells = {header[i]: cell.value for i, cell in enumerate(sheet[2])}
    assert cells["图"] == "Fig. 3" and cells["读数（近似值）"] is None and cells["精度"] == "±20%"
    assert workbook["样品数据"].max_row == 2  # the sample sheet is what it would have been without the row


def field_levels(output: Path) -> dict[str, str]:
    sheet = load_workbook(output)["字段说明"]
    header = [cell.value for cell in sheet[1]]
    name, level = header.index("字段"), header.index("层级")
    return {row[name].value: row[level].value for row in sheet.iter_rows(min_row=2)}


def test_the_field_sheet_names_each_level_in_the_profiles_own_words(tmp_path: Path, tco_profile):
    write_dataset([], tmp_path / "tco.xlsx", tco_profile)

    levels = field_levels(tmp_path / "tco.xlsx")

    # TCO's copy is what the sheet said before the level label came from the profile.
    assert levels["component"] == "靶材（论文级）"
    assert levels["thickness"] == "样品级"


def test_another_profile_labels_its_levels_its_own_way(tmp_path: Path):
    profile = make_profile({"ui": {"paper_level_label_zh": "前驱体（论文级）", "entity_label_zh": "涂层"}})
    write_dataset([], tmp_path / "demo.xlsx", profile)

    levels = field_levels(tmp_path / "demo.xlsx")

    assert levels == {"precursor_purity": "前驱体（论文级）", "coating_thickness": "涂层级", "solvent": "涂层级"}
