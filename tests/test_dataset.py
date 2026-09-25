"""An ML row must describe one actual sample and contain only defensible scalar values."""

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from paperfacts.compare import compare_lanes
from paperfacts.dataset import (
    DatasetPayload,
    DocumentDataset,
    consolidate_document,
    write_dataset,
    write_dataset_json,
)
from paperfacts.fields import FIELD_SPECS
from paperfacts.matching import SampleMatch, SampleMatching
from paperfacts.models import DocumentInput
from paperfacts.records import FieldValue, TargetRecord
from support.extraction import make_lane, make_sample
from support.factories import DOC_ID


def value(name, raw, unit=None, *, condition=None, backend="mineru", **kwargs):
    return FieldValue(
        field=name, value_raw=raw, unit_raw=unit, condition=condition, source_ids=(f"{backend}_p0_b1",), **kwargs
    )


def dataset(a, b=None, matching=None, *, filename="paper.pdf"):
    b = b or make_lane(backend="paddleocr_vl")
    matching = matching or SampleMatching(
        unmatched_a=tuple(s.sample_id for s in a.samples), unmatched_b=tuple(s.sample_id for s in b.samples)
    )
    document = DocumentInput(document_id=DOC_ID, sha256=DOC_ID, pdf_path=Path(filename))
    report = compare_lanes(a, b, matching)
    return consolidate_document(document, {a.backend: a, b.backend: b}, report)


def paired(a, b, *, confidence=1.0, filename="paper.pdf"):
    return dataset(
        make_lane(samples=[make_sample("A", a)]),
        make_lane(backend="paddleocr_vl", samples=[make_sample("A", b)]),
        SampleMatching(
            pairs=(SampleMatch(a_id="A", b_id="A", confidence=confidence, method="llm", justification="test"),)
        ),
        filename=filename,
    )


def decision(result, field, sample_id=None):
    return next(
        row
        for row in result.quality_rows
        if row["field"] == field and (sample_id is None or row["sample_id"] == sample_id)
    )


def test_paper_row_selects_one_complete_sample_without_cross_sample_fill():
    result = dataset(
        make_lane(
            samples=[
                make_sample(
                    "A",
                    [value("thickness", "300", "nm"), value("transmittance", "85", "%", condition="550 nm")],
                    label="as deposited",
                    conditions={"temperature": "300 K"},
                ),
                make_sample("B", [value("resistivity", "1e-4", "Ω cm")], label="annealed"),
            ]
        )
    )

    assert result.paper_row["sample_id"] == "mineru:A"
    assert result.paper_row["thickness"] == 300
    assert result.paper_row["transmittance"] == 85
    assert result.paper_row["resistivity"] is None
    assert "temperature=300 K" in result.paper_row["conditions"]
    assert len(result.sample_rows) == 2
    assert {spec.name for spec in FIELD_SPECS} <= result.paper_row.keys()


def test_unit_conversion_and_duplicate_sources_yield_one_numeric_value():
    a = value("thickness", "0.3", "μm")
    result = paired(
        [a, a.model_copy(update={"source_ids": ("mineru_p1_b1",)})],
        [value("thickness", "300", "nm", backend="paddleocr_vl")],
    )

    assert result.paper_row["thickness"] == 300
    assert decision(result, "thickness")["decision"] == "agree"
    assert "mineru_p1_b1" in decision(result, "thickness")["source_ids"]
    assert len([row for row in result.quality_rows if row["field"] == "thickness"]) == 1


def test_agree_chooses_highest_repeat_agreement_without_averaging():
    result = paired(
        [value("thickness", "300", "nm", agreement=0.5)],
        [value("thickness", "305", "nm", agreement=1.0, backend="paddleocr_vl")],
    )
    assert result.paper_row["thickness"] == 305


def test_a_comparison_agreement_on_a_value_that_failed_grounding_cannot_label_another_value_agreed():
    # acsnano: both lanes said 230 °C, but lane A's 230 failed grounding; its lone 150 must not be "agree".
    a = [
        value("annealing_temperature", "230", "°C", grounded=False),
        value("annealing_temperature", "150", "°C", condition="second anneal"),
    ]
    b = [value("annealing_temperature", "230", "°C", backend="paddleocr_vl")]

    result = paired(a, b)

    assert decision(result, "annealing_temperature")["decision"] != "agree"
    assert result.paper_row["annealing_temperature"] is None


def test_conflict_and_low_confidence_matching_remain_empty():
    result = paired([value("thickness", "300", "nm")], [value("thickness", "900", "nm", backend="paddleocr_vl")])
    assert result.paper_row["thickness"] is None
    assert decision(result, "thickness")["decision"] == "conflict"
    result = paired(
        [value("thickness", "300", "nm")], [value("thickness", "300", "nm", backend="paddleocr_vl")], confidence=0.2
    )
    assert result.paper_row["thickness"] is None
    assert decision(result, "thickness")["decision"] == "ambiguous"


@pytest.mark.parametrize(
    "raw", ["10-20", "<10", ">=10", "10 (20)", "10 × 20", "10, 20", "1e-4 to 2e-4", "10 nm at 300 K"]
)
def test_lossy_numeric_interpretations_are_never_exported(raw):
    result = paired([value("thickness", raw, "nm")], [value("thickness", raw, "nm", backend="paddleocr_vl")])
    assert result.paper_row["thickness"] is None
    assert decision(result, "thickness")["decision"] == "non_scalar"


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("~300", 300),
        ("300 ± 2", 300),
        ("1e-4 ± 2e-5", 1e-4),
        ("300 ± 2e-1", 300),
        # The uncertainty in parentheses after the unit, as GM1's PaddleOCR-VL lane quotes it.
        ("300 nm (± 5 nm)", 300),
        ("300 nm (  $ \\pm $ 5 nm)", 300),
    ],
)
def test_approximation_and_uncertainty_keep_documented_center(raw, expected):
    result = paired([value("thickness", raw, "nm")], [value("thickness", raw, "nm", backend="paddleocr_vl")])
    assert result.paper_row["thickness"] == expected
    assert "中心值" in decision(result, "thickness")["detail"]


def test_a_lower_bound_beside_a_scalar_is_set_aside_rather_than_refusing_the_cell():
    # Bauden O2-100sccm: both lanes quote 80.6 % (380-780 nm) and ">80 %" (500-2500 nm).
    def lane(backend):
        return [
            value("transmittance", "80.6", "%", condition="380-780 nm", backend=backend),
            value("transmittance", ">80", "%", condition="500-2500 nm", backend=backend),
        ]

    result = paired(lane("mineru"), lane("paddleocr_vl"))

    row = decision(result, "transmittance")
    assert (result.paper_row["transmittance"], row["decision"]) == (80.6, "agree")
    assert "已排除不是唯一标量的候选" in row["detail"] and ">80" in row["detail"]
    assert row["conditions"] == "380-780 nm"


def test_a_cell_holding_only_bounds_and_ranges_is_still_non_scalar():
    def lane(backend):
        return [
            value("transmittance", ">80", "%", condition="500-2500 nm", backend=backend),
            value("transmittance", "80-85", "%", condition="400-800 nm", backend=backend),
        ]

    result = paired(lane("mineru"), lane("paddleocr_vl"))

    assert decision(result, "transmittance")["decision"] == "non_scalar"


def test_one_value_under_differently_worded_conditions_is_one_measurement():
    # GZO HN450: every candidate is 100 nm, under free-text notes rather than measurement conditions.
    def lane(backend):
        return [
            value("thickness", "100", "nm", condition="measured by TEM cross-section", backend=backend),
            value(
                "thickness",
                "100",
                "nm",
                condition="thickness not reduced after forming gas post-treatment",
                backend=backend,
            ),
        ]

    result = paired(lane("mineru"), lane("paddleocr_vl"))

    row = decision(result, "thickness")
    assert (result.paper_row["thickness"], row["decision"]) == (100, "agree")
    assert "视为同一测量" in row["detail"]
    assert "measured by TEM cross-section" in row["conditions"] and "forming gas" in row["conditions"]


def test_different_values_under_differently_worded_conditions_stay_refused():
    fields = [
        value("thickness", "100", "nm", condition="measured by TEM cross-section"),
        value("thickness", "140", "nm", condition="by profilometry"),
    ]

    result = paired(fields, [])

    assert decision(result, "thickness")["decision"] == "multiple_conditions"


def test_a_comparison_cannot_hide_different_same_condition_values_in_one_lane():
    result = paired(
        [value("thickness", "300", "nm"), value("thickness", "400", "nm")],
        [value("thickness", "300", "nm", backend="paddleocr_vl")],
    )
    assert result.paper_row["thickness"] is None
    assert decision(result, "thickness")["decision"] == "multiple_values"


def test_multiple_conditions_are_not_collapsed_even_with_identical_numbers():
    # Neither wavelength is in transmittance's condition_preference, so nothing picks one.
    fields = [value("transmittance", "85", "%", condition=condition) for condition in ("450 nm", "600 nm")]
    result = paired(fields, [v.model_copy(update={"source_ids": ("paddleocr_vl_p0_b1",)}) for v in fields])
    assert result.paper_row["transmittance"] is None
    assert decision(result, "transmittance")["decision"] == "multiple_conditions"


def test_differently_worded_conditions_across_lanes_still_agree():
    """The lanes paraphrase one condition; only a lane's own evidence may signal several conditions."""
    result = paired(
        [value("transmittance", "85", "%", condition="ITO monolayer thickness")],
        [value("transmittance", "85", "%", condition="ITO monolayer film thickness", backend="paddleocr_vl")],
    )
    assert result.paper_row["transmittance"] == 85
    assert decision(result, "transmittance")["decision"] == "agree"
    row = decision(result, "transmittance")
    assert "ITO monolayer thickness" in row["conditions"] and "ITO monolayer film thickness" in row["conditions"]


def test_one_lane_with_two_conditions_is_still_refused():
    result = paired(
        [
            value("transmittance", "85", "%", condition="450 nm"),
            value("transmittance", "85", "%", condition="600 nm"),
        ],
        [value("transmittance", "85", "%", condition="450 nm", backend="paddleocr_vl")],
    )
    assert result.paper_row["transmittance"] is None
    assert decision(result, "transmittance")["decision"] == "multiple_conditions"


def _cited(field, source):
    return field.model_copy(update={"source_ids": (source,)})


def test_the_condition_stated_in_the_same_block_as_the_rest_of_the_row_fills_the_cell():
    # GM1: "5.74e-4 Ω·cm ... 83.5 % (400-1800 nm)" in one abstract sentence, other ranges elsewhere.
    def lane(backend):
        return [
            _cited(value("resistivity", "5.74e-4", "Ω·cm", backend=backend), f"{backend}_p0_b9"),
            _cited(value("transmittance", "83.5", "%", condition="400-1800 nm"), f"{backend}_p0_b9"),
            _cited(value("transmittance", "81.6", "%", condition="400-800 nm"), f"{backend}_p3_b2"),
        ]

    result = paired(lane("mineru"), lane("paddleocr_vl"))

    row = decision(result, "transmittance")
    assert (result.paper_row["transmittance"], row["decision"]) == (83.5, "agree")
    assert row["conditions"] == "400-1800 nm"
    # Traceable to the sentence it came from, not to the blocks of the values set aside.
    assert row["source_ids"] == "mineru_p0_b9; paddleocr_vl_p0_b9"


def test_two_lanes_tying_different_conditions_to_the_row_are_not_committed_as_agreement():
    # The review's scenario: each lane picks one condition by its own row block, but not the same one, and
    # the two measurements used to be committed as one "agree" with conditions "450 nm; 600 nm".
    def lane(backend, row_condition, other_condition):
        values = {"450 nm": "85", "600 nm": "85.2"}
        return [
            _cited(value("resistivity", "5.74e-4", "Ω·cm", backend=backend), f"{backend}_p0_b9"),
            _cited(value("transmittance", values[row_condition], "%", condition=row_condition), f"{backend}_p0_b9"),
            _cited(value("transmittance", values[other_condition], "%", condition=other_condition), f"{backend}_p3_b2"),
        ]

    result = paired(lane("mineru", "450 nm", "600 nm"), lane("paddleocr_vl", "600 nm", "450 nm"))

    row = decision(result, "transmittance")
    assert row["decision"] == "multiple_conditions"
    assert result.paper_row["transmittance"] is None


def test_the_lanes_may_word_the_chosen_condition_differently():
    def lane(backend, row_condition):
        return [
            _cited(value("resistivity", "5.74e-4", "Ω·cm", backend=backend), f"{backend}_p0_b9"),
            _cited(value("transmittance", "85", "%", condition=row_condition), f"{backend}_p0_b9"),
            _cited(value("transmittance", "80", "%", condition="at 600 nm"), f"{backend}_p3_b2"),
        ]

    result = paired(lane("mineru", "at 450 nm"), lane("paddleocr_vl", "450 nm wavelength"))

    assert (result.paper_row["transmittance"], decision(result, "transmittance")["decision"]) == (85.0, "agree")


def test_several_conditions_sharing_the_rows_block_stay_refused():
    fields = [
        _cited(value("resistivity", "5.74e-4", "Ω·cm"), "mineru_p0_b9"),
        _cited(value("transmittance", "83.5", "%", condition="400-1800 nm"), "mineru_p0_b9"),
        _cited(value("transmittance", "81.6", "%", condition="450-700 nm"), "mineru_p0_b9"),
    ]

    result = paired(fields, [])

    assert decision(result, "transmittance")["decision"] == "multiple_conditions"


def test_without_a_row_sharing_condition_the_fields_preference_picks_the_cell():
    # Zhao: 91.9 % averaged over 400-800 nm and 92.2 % at 550 nm, both lanes, no sentence shared with the row.
    def lane(backend):
        return [
            _cited(value("transmittance", "92.2", "%", condition="at 550 nm"), f"{backend}_p4_b2"),
            _cited(value("transmittance", "91.9", "%", condition="average from 400 to 800 nm"), f"{backend}_p4_b3"),
        ]

    result = paired(lane("mineru"), lane("paddleocr_vl"))

    row = decision(result, "transmittance")
    assert (result.paper_row["transmittance"], row["decision"]) == (91.9, "agree")
    assert "优先条件 400-800" in row["detail"]
    assert row["source_ids"] == "mineru_p4_b3; paddleocr_vl_p4_b3"


def test_a_preference_matching_two_conditions_in_one_lane_moves_on_to_the_next():
    # Both name 400 and 800 and neither is a peak; the tie is not settled by the first entry, so 550 decides.
    fields = [
        value("transmittance", "91.9", "%", condition="average 400-800 nm, as deposited"),
        value("transmittance", "90.5", "%", condition="average 400-800 nm, after bending"),
        value("transmittance", "92.2", "%", condition="550 nm"),
    ]

    result = paired(fields, [])

    assert result.paper_row["transmittance"] == 92.2


def test_inside_one_preference_entry_an_average_beats_a_peak():
    fields = [
        value("transmittance", "91.9", "%", condition="average 400-800 nm"),
        value("transmittance", "95.0", "%", condition="peak 400-800 nm"),
        value("transmittance", "92.2", "%", condition="550 nm"),
    ]

    result = paired(fields, [value("transmittance", "91.9", "%", condition="avg. 400-800 nm", backend="paddleocr_vl")])

    row = decision(result, "transmittance")
    assert (result.paper_row["transmittance"], row["decision"]) == (91.9, "agree")
    assert "平均值优先于峰值" in row["detail"]


def test_zhaos_average_over_400_to_1100_nm_is_preferred_over_the_other_ranges():
    # Zhao ICO-30nm annealed: averaged over 400-1100 nm, over 800-1100 nm, and a peak value.
    def lane(backend, average):
        return [
            value("transmittance", average, "%", condition="average 400-1100 nm", backend=backend),
            value("transmittance", "96.6", "%", condition="average 800-1100 nm", backend=backend),
            value("transmittance", "98.1", "%", condition="maximum transmittance", backend=backend),
        ]

    result = paired(lane("mineru", "92.1"), lane("paddleocr_vl", "92.1"))

    row = decision(result, "transmittance")
    assert (result.paper_row["transmittance"], row["decision"]) == (92.1, "agree")
    assert "优先条件 400-1100" in row["detail"]


def test_two_peaks_inside_one_preference_entry_are_still_refused():
    fields = [
        value("transmittance", "95.0", "%", condition="peak 400-800 nm"),
        value("transmittance", "97.5", "%", condition="max 400-800 nm after anneal"),
    ]

    result = paired(fields, [])

    assert decision(result, "transmittance")["decision"] == "multiple_conditions"


def test_the_other_lanes_value_for_a_condition_set_aside_cannot_vouch_for_the_chosen_one():
    # Lane B only quotes 400-800 nm, from another block. The comparison agrees on that condition, but it is
    # not the one chosen for the cell, so the cell is one lane's word, not an agreement.
    a = [
        _cited(value("resistivity", "5.74e-4", "Ω·cm"), "mineru_p0_b9"),
        _cited(value("transmittance", "83.5", "%", condition="400-1800 nm"), "mineru_p0_b9"),
        _cited(value("transmittance", "81.6", "%", condition="400-800 nm"), "mineru_p3_b2"),
    ]
    b = [_cited(value("transmittance", "81.6", "%", condition="400-800 nm"), "paddleocr_vl_p3_b2")]

    result = paired(a, b)

    row = decision(result, "transmittance")
    assert (result.paper_row["transmittance"], row["decision"]) == (83.5, "single_source")
    assert row["source_ids"] == "mineru_p0_b9"


def test_one_mode_quoted_two_ways_is_one_answer_not_a_refusal():
    """A closed category set is judged on the category, so the paper's phrasing cannot manufacture a
    multiple_values refusal out of a single co-sputtering run."""
    result = paired(
        [value("mode", "DC and RF")],
        [value("mode", "DC and RF magnetron co-sputtering", backend="paddleocr_vl")],
    )

    assert decision(result, "mode")["decision"] == "agree"
    assert result.paper_row["mode"] in {"DC and RF", "DC and RF magnetron co-sputtering"}


def test_two_genuinely_different_modes_are_still_refused():
    result = paired([value("mode", "DC")], [value("mode", "RF", backend="paddleocr_vl")])

    assert result.paper_row["mode"] is None
    assert decision(result, "mode")["decision"] == "conflict"


def test_one_lane_quoting_two_spellings_of_one_mode_is_not_multiple_values():
    result = paired(
        [value("mode", "DC and RF"), value("mode", "DC and RF co-sputtering")],
        [value("mode", "DC and RF", backend="paddleocr_vl")],
    )

    assert decision(result, "mode")["decision"] != "multiple_values"
    assert result.paper_row["mode"] is not None


def test_one_lane_quoting_two_different_modes_stays_refused():
    result = paired(
        [value("mode", "DC"), value("mode", "RF")],
        [value("mode", "DC", backend="paddleocr_vl")],
    )

    assert result.paper_row["mode"] is None
    assert decision(result, "mode")["decision"] in {"conflict", "multiple_values"}


def test_a_target_size_written_as_a_number_word_fills_the_cell():
    # metals: "a four-inch ITO target", quoted as "four" with the unit "inch" by both lanes.
    result = dataset(
        make_lane(target=TargetRecord(fields=(value("inch", "four", "inch"),))),
        make_lane(
            backend="paddleocr_vl",
            target=TargetRecord(fields=(value("inch", "four-inch", "inch", backend="paddleocr_vl"),)),
        ),
    )

    assert result.paper_row["inch"] == 4
    assert decision(result, "inch")["decision"] == "agree"


def test_different_target_compositions_cannot_be_picked_or_joined():
    result = dataset(make_lane(target=TargetRecord(fields=(value("component", "SnO2"), value("component", "ZnO")))))
    assert result.paper_row["component"] is None
    assert decision(result, "component")["decision"] == "multiple_values"


def test_unmatched_samples_with_the_same_id_stay_separate_per_backend():
    a = make_lane(samples=[make_sample("001", [value("thickness", "300", "nm")])])
    b = make_lane(
        backend="paddleocr_vl",
        samples=[make_sample("001", [value("resistivity", "1e-4", "Ω cm", backend="paddleocr_vl")])],
    )
    result = dataset(a, b)
    assert len(result.sample_rows) == 2
    assert {row["sample_id"] for row in result.sample_rows} == {"mineru:001", "paddleocr_vl:001"}
    assert sum(row["thickness"] is not None for row in result.sample_rows) == 1
    assert sum(row["resistivity"] is not None for row in result.sample_rows) == 1
    assert result.paper_row["resistivity"] is None


def test_a_failed_match_does_not_turn_into_trusted_single_source_rows():
    a = make_lane(samples=[make_sample("A", [value("thickness", "300", "nm")])])
    result = dataset(a, matching=SampleMatching(unmatched_a=("A",), failed=True, failure="model failed"))
    assert result.paper_row["thickness"] is None


@pytest.mark.parametrize("updates", [{"grounded": False}, {"source_ids": ()}])
def test_values_without_grounded_citations_stay_empty(updates):
    field = value("thickness", "300", "nm").model_copy(update=updates)
    result = dataset(make_lane(samples=[make_sample("A", [field])]))
    assert result.paper_row["thickness"] is None
    assert decision(result, "thickness")["decision"] == "ungrounded"


def test_agreement_with_one_ungrounded_side_uses_only_trusted_evidence():
    result = paired(
        [value("thickness", "300", "nm", grounded=False)], [value("thickness", "305", "nm", backend="paddleocr_vl")]
    )
    assert result.paper_row["thickness"] == 305
    assert decision(result, "thickness")["decision"] == "single_source"


def test_paper_selection_prefers_two_lane_agreement_then_stable_sample_id():
    a = make_lane(
        samples=[
            make_sample("A", [value("thickness", "300", "nm")]),
            make_sample("B", [value("thickness", "400", "nm")]),
        ]
    )
    b = make_lane(
        backend="paddleocr_vl", samples=[make_sample("B", [value("thickness", "400", "nm", backend="paddleocr_vl")])]
    )
    matching = SampleMatching(
        pairs=(SampleMatch(a_id="B", b_id="B", confidence=1.0, method="exact", justification="same"),),
        unmatched_a=("A",),
    )
    assert dataset(a, b, matching).paper_row["sample_id"] == "B"


def test_excel_reopens_with_numeric_fields_text_ids_and_no_pdf_formulas(tmp_path):
    a = make_lane(
        samples=[
            make_sample(
                "001",
                [value("thickness", "0.3", "μm"), value("resistivity", "1e-4", "Ω cm")],
                label='=HYPERLINK("https://example.invalid")',
                conditions={"temperature": "001"},
            )
        ],
        target=TargetRecord(fields=(value("component", "=1+1"),)),
    )
    b = make_lane(backend="paddleocr_vl", samples=[make_sample("001", a.samples[0].fields)], target=a.target)
    matching = SampleMatching(
        pairs=(SampleMatch(a_id="001", b_id="001", confidence=1.0, method="exact", justification="same"),)
    )
    result = dataset(a, b, matching, filename="=paper.pdf")
    output = tmp_path / "dataset.xlsx"
    write_dataset(
        [result, result],
        output,
        failures=[{"document_id": "b" * 64, "filename": "failed.pdf", "error": "parser failed"}],
    )
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["论文数据", "样品数据", "字段说明", "数据质量", "图中读数", "运行记录"]
    sheet = workbook["论文数据"]
    assert sheet.max_row == 2
    assert sheet.freeze_panes == "D2"
    assert len(sheet.tables) == 1
    columns = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, columns["thickness"]).value == 300
    assert sheet.cell(2, columns["thickness"]).data_type == "n"
    assert sheet.cell(2, columns["样品ID"]).value == "001"
    assert sheet.cell(2, columns["样品ID"]).data_type == "s"
    assert sheet.cell(2, columns["component"]).value == "=1+1"
    assert sheet.cell(2, columns["component"]).data_type == "s"
    assert sheet.cell(2, columns["文件名"]).data_type == "s"
    assert sheet.cell(2, columns["density"]).value is None
    assert "E+00" in sheet.cell(2, columns["resistivity"]).number_format
    assert workbook["运行记录"].max_row == 3
    assert workbook["运行记录"].cell(3, 3).value == "failed"
    assert not list(tmp_path.glob("*.tmp"))
    assert all(cell.data_type != "f" for page in workbook for row in page for cell in row)


def test_an_empty_export_still_records_failures(tmp_path):
    output = tmp_path / "empty.xlsx"
    write_dataset([], output, failures=[{"filename": "bad.pdf", "error": "unreadable"}])
    workbook = load_workbook(output)
    assert workbook["论文数据"].max_row == 1
    assert workbook["运行记录"].cell(2, 3).value == "failed"


def test_mismatched_document_ids_are_rejected():
    a, b = make_lane(), make_lane(backend="paddleocr_vl")
    report = compare_lanes(a, b, SampleMatching())
    document = DocumentInput(document_id="b" * 64, sha256="b" * 64, pdf_path=Path("other.pdf"))
    with pytest.raises(ValueError, match="same PDF"):
        consolidate_document(document, {a.backend: a, b.backend: b}, report)


def test_the_json_view_survives_a_round_trip(tmp_path):
    # Rows are MappingProxyType, which json refuses; the web UI reads this file, so the copy must be real.
    result = paired([value("thickness", "300", "nm")], [value("thickness", "300", "nm")])

    path = tmp_path / "dataset.json"
    write_dataset_json(result, path)
    loaded = json.loads(path.read_text(encoding="utf-8"))

    assert loaded["document_id"] == result.document_id
    assert loaded["filename"] == "paper.pdf"
    assert loaded["extractor_key"] == result.extractor_key
    assert loaded["comparison_key"] == result.comparison_key
    assert [field["name"] for field in loaded["fields"]] == [spec.name for spec in FIELD_SPECS]
    assert loaded["paper_row"] == dict(result.paper_row)
    assert loaded["sample_rows"] == [dict(row) for row in result.sample_rows]
    assert loaded["quality_rows"] == [dict(row) for row in result.quality_rows]
    assert loaded["sample_rows"][0]["thickness"] == 300
    assert not list(tmp_path.glob("*.tmp"))


def test_the_field_list_carries_the_chinese_description_for_the_header_tooltip():
    result = paired([value("transmittance", "85", "%")], [value("transmittance", "85", "%", backend="paddleocr_vl")])
    by_name = {field.name: field for field in result.to_payload().fields}
    assert "透光率" in by_name["transmittance"].description
    assert by_name["transmittance"].scope == "sample"


def test_every_configured_field_has_a_chinese_description():
    # The Excel field sheet and the header tooltips both read description_zh; a field shipped in
    # config.json without one would export blank, which this test turns into a visible failure instead.
    assert [spec.name for spec in FIELD_SPECS if not spec.description_zh] == []


def test_the_field_list_carries_the_chinese_label_for_the_column_header():
    # The browser prints this above the column; a field without one falls back to its id, never to blank.
    by_name = {field.name: field for field in dataset(make_lane()).to_payload().fields}

    assert by_name["transmittance"].label == "透光率"
    assert by_name["thickness"].label == "厚度"


def test_from_payload_reverses_to_payload_exactly():
    # The corpus export rebuilds datasets from their JSON, so the round trip through disk has to be lossless.
    result = paired([value("thickness", "300", "nm")], [value("thickness", "300", "nm")])

    parsed = DatasetPayload.model_validate_json(result.to_payload().model_dump_json())
    restored = DocumentDataset.from_payload(parsed)

    assert restored == result
    assert restored.to_payload() == result.to_payload()


def test_a_dataset_file_in_the_wrong_shape_fails_at_the_boundary():
    # Validation is pydantic's job where the JSON is parsed, so a bad file never travels on as a dict.
    with pytest.raises(ValidationError):
        DatasetPayload.model_validate_json("[1, 2, 3]")


def test_the_json_field_list_carries_the_unit_and_the_scope():
    fields = {field.name: field for field in dataset(make_lane()).to_payload().fields}

    assert fields["thickness"].scope == "sample"
    assert fields["component"].scope == "target"
    assert fields["thickness"].unit == "nm"


def test_the_display_name_is_the_filename_a_dataset_reports():
    """A web upload is stored as source.pdf, so the path can't name the paper: the display name does."""
    document = DocumentInput(
        document_id=DOC_ID, sha256=DOC_ID, pdf_path=Path("source.pdf"), display_name="Sputtered ITO.pdf"
    )
    a = make_lane(samples=[make_sample("A", [value("thickness", "100 nm")])])
    b = make_lane(backend="paddleocr_vl")
    matching = SampleMatching(unmatched_a=("A",))
    result = consolidate_document(document, {a.backend: a, b.backend: b}, compare_lanes(a, b, matching))

    assert result.filename == "Sputtered ITO.pdf"
    assert result.paper_row["filename"] == "Sputtered ITO.pdf"
    assert [row["filename"] for row in result.sample_rows] == ["Sputtered ITO.pdf"]
    assert all(row["filename"] == "Sputtered ITO.pdf" for row in result.quality_rows)


def test_without_a_display_name_the_dataset_still_reports_the_path_name():
    result = dataset(make_lane(samples=[make_sample("A", [value("thickness", "100 nm")])]), filename="paper.pdf")

    assert result.filename == "paper.pdf"
    assert result.paper_row["filename"] == "paper.pdf"


def test_a_scale_factor_in_the_unit_reaches_the_paper_row():
    # Observed on the corpus: a column headed "ρ (×10⁻⁴ Ω·cm)" makes the model transcribe the factor as the
    # unit and leave the value a bare "19.4". The exported cell must be the resistivity, not an empty cell.
    result = paired(
        [value("resistivity", "19.4", "×10^-4 Ω-cm")],
        [value("resistivity", "19.4", "×10^-4 Ω·cm", backend="paddleocr_vl")],
    )

    assert result.paper_row["resistivity"] == pytest.approx(1.94e-3)
    assert decision(result, "resistivity")["decision"] == "agree"


def test_an_approximate_value_spaced_out_by_latex_still_exports():
    result = paired(
        [value("thickness", "∼8 . 4", "nm")],
        [value("thickness", "around 8.4", "nm", backend="paddleocr_vl")],
    )

    assert result.paper_row["thickness"] == pytest.approx(8.4)


def test_a_value_stated_only_for_the_whole_series_is_marked_series_level():
    result = paired(
        [value("thickness", "300", "nm", series=True)],
        [value("thickness", "300", "nm", backend="paddleocr_vl", series=True)],
    )

    assert decision(result, "thickness")["decision"] == "agree"
    assert decision(result, "thickness")["series"] is True


def test_one_sample_specific_side_keeps_the_value_off_the_series_mark():
    result = paired(
        [value("thickness", "300", "nm", series=True)],
        [value("thickness", "300", "nm", backend="paddleocr_vl")],
    )

    assert decision(result, "thickness")["decision"] == "agree"
    assert decision(result, "thickness")["series"] is False


def test_a_single_source_series_value_is_still_marked_series_level():
    result = dataset(make_lane(samples=[make_sample("A", [value("thickness", "300", "nm", series=True)])]))

    assert decision(result, "thickness")["decision"] == "single_source"
    assert decision(result, "thickness")["series"] is True


def test_the_series_mark_reaches_the_quality_sheet(tmp_path):
    result = paired(
        [value("thickness", "300", "nm", series=True)],
        [value("thickness", "300", "nm", backend="paddleocr_vl", series=True)],
    )
    output = tmp_path / "dataset.xlsx"
    write_dataset([result], output)

    sheet = load_workbook(output)["数据质量"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    rows = {sheet.cell(row, columns["字段"]).value: row for row in range(2, sheet.max_row + 1)}
    assert sheet.cell(rows["thickness"], columns["系列级"]).value is True
    assert sheet.cell(rows["resistivity"], columns["系列级"]).value is False


def test_the_quality_row_names_the_lanes_behind_a_committed_value():
    agreed = paired([value("thickness", "300", "nm")], [value("thickness", "300", "nm", backend="paddleocr_vl")])
    assert decision(agreed, "thickness")["decision"] == "agree"
    assert decision(agreed, "thickness")["lanes"] == "mineru; paddleocr_vl"

    only_b = make_sample("A", [value("thickness", "300", "nm", backend="paddleocr_vl")])
    single = dataset(make_lane(), make_lane(backend="paddleocr_vl", samples=[only_b]))
    assert decision(single, "thickness")["decision"] == "single_source"
    assert decision(single, "thickness")["lanes"] == "paddleocr_vl"


def test_a_refused_cell_names_no_lane():
    result = dataset(make_lane(samples=[make_sample("A", [value("thickness", "300", "nm", grounded=False)])]))
    assert decision(result, "thickness")["lanes"] == ""


def test_the_lane_column_reaches_the_quality_sheet(tmp_path):
    result = paired([value("thickness", "300", "nm")], [value("thickness", "300", "nm", backend="paddleocr_vl")])
    output = tmp_path / "dataset.xlsx"
    write_dataset([result], output)

    sheet = load_workbook(output)["数据质量"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    rows = {sheet.cell(row, columns["字段"]).value: row for row in range(2, sheet.max_row + 1)}
    assert sheet.cell(rows["thickness"], columns["证据来源通道"]).value == "mineru; paddleocr_vl"


# ---- Figure readings: their own sheet, never a cell -------------------------------------------------


def test_figure_rows_fill_only_their_own_sheet(tmp_path: Path):
    output = tmp_path / "dataset.xlsx"
    result = dataset(make_lane(samples=[make_sample("A", [value("thickness", "100", "nm")])]))
    row = {"document_id": DOC_ID, "figure": "Fig. 3", "value": None, "value_raw": "25 10^2 ohm/sq", "precision": "±20%"}

    write_dataset([result], output, figure_rows=[row])

    workbook = load_workbook(output)
    sheet = workbook["图中读数"]
    header = [cell.value for cell in sheet[1]]
    cells = {header[i]: cell.value for i, cell in enumerate(sheet[2])}
    assert cells["图"] == "Fig. 3" and cells["读数（近似值）"] is None and cells["精度"] == "±20%"
    assert workbook["样品数据"].max_row == 2  # the sample sheet is what it would have been without the row


def test_the_dataset_module_knows_nothing_of_figure_reading():
    # Chart readings take part in no verdict, so the module whose source names every comparison must not
    # depend on the one that reads charts: tuning the figures stage must never rename a stored comparison.
    import ast

    import paperfacts.dataset

    tree = ast.parse(Path(paperfacts.dataset.__file__).read_text(encoding="utf-8"))
    imported = {node.module for node in ast.walk(tree) if isinstance(node, ast.ImportFrom)}
    assert "paperfacts.figures" not in imported
    assert "figure_rows" not in DatasetPayload.model_fields
