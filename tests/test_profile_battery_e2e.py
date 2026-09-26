"""The ``battery_cathode`` example profile, end to end: a second domain runs through the real pipeline with no
code edit (AC-7), converts its own units and refuses a range where its field says so (AC-8), and shares one
data_root with TCO without either run touching the other's files (AC-9).

Nothing is parsed and no model is called. The two parser outputs are hand-built into the data_root, which the
parse stage takes as a cache hit, and a scripted :class:`FakeLlmClient` answers the inventory, field and matching
questions from the question text, so the answers do not depend on the order the pool asks them in.
"""

from __future__ import annotations

import dataclasses
import json
import re
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from paperfacts.config import Settings
from paperfacts.keys import comparison_key_for, extractor_key_for, figure_key_for
from paperfacts.models import BACKENDS, DocumentGeometry, DocumentInput
from paperfacts.profile import DomainProfile
from paperfacts.profile_loader import load_profile
from paperfacts.storage import DataLayout
from paperfacts.workflow import PipelineResult, load_run_profile, run_document
from support.factories import RawOutputFactory, paddle_page_entry
from support.llm import FakeLlmClient
from support.profiles import SHIPPED_PROFILE_PATH

BATTERY_PROFILE_PATH = SHIPPED_PROFILE_PATH.with_name("battery_cathode.json")

# The paper, one block per entry. Each lane names its two samples differently, so matching has to ask the model.
TITLE = "Calcination temperature of LiNi0.8Co0.1Mn0.1O2 cathodes"
PARAGRAPHS = (
    "Two LiNi0.8Co0.1Mn0.1O2 cathode samples were synthesized by co-precipitation and calcined in oxygen for "
    "12 h, at 800 °C (NCM-800) and at 1173 K (NCM-900).",
    "Coin cells were cycled in the voltage window 2.8-4.3 V. NCM-800 delivers an initial discharge capacity of "
    "0.15 Ah/g at 0.1 C, and NCM-900 delivers 0.17 Ah/g at 0.1 C.",
    "After 100 cycles at 1 C, the capacity retention of NCM-800 is 92.5% and that of NCM-900 is 88.0%.",
)
INVENTORY = {
    "mineru": (("NCM-800", "calcined at 800 °C"), ("NCM-900", "calcined at 1173 K")),
    "paddleocr_vl": (("S1", "800 °C sample"), ("S2", "1173 K sample")),
}
# Per field, one (value_raw, unit_raw, condition) per listed sample, in list order.
ANSWERS: dict[str, tuple[tuple[str, str | None, str | None], ...]] = {
    "calcination_temperature": (("800", "°C", None), ("1173", "K", None)),
    "calcination_time": (("12", "h", None), ("12", "h", None)),
    "initial_discharge_capacity": (("0.15", "Ah/g", "0.1 C"), ("0.17", "Ah/g", "0.1 C")),
    # The whole window, quoted where the field asks for its upper limit: refused under range_policy "reject",
    # where the default policy would have invented a 3.55 V cut-off.
    "upper_cutoff_voltage": (("2.8-4.3", "V", None), ("2.8-4.3", "V", None)),
    "capacity_retention": (("92.5", "%", "after 100 cycles at 1 C"), ("88.0", "%", "after 100 cycles at 1 C")),
}

_SOURCE = re.compile(r"<!-- source: (\S+) -->")
_LISTED_ID = re.compile(r"^- id: (.+?) \| label:", re.MULTILINE)
_FIELD = re.compile(r"\AField to extract:\n- `([^`]+)`")


def _lane(user: str) -> str:
    return "paddleocr_vl" if "<!-- source: paddleocr_vl_" in user else "mineru"


def _citing(excerpts: str, text: str) -> list[str]:
    """The id of the block that contains ``text``, so every answer is grounded in the block it cites."""
    for block in re.split(r"(?=<!-- source: )", excerpts):
        found = _SOURCE.match(block)
        if found and text in block:
            return [found.group(1)]
    return _SOURCE.findall(excerpts)[:1]


def scripted_model(profile: DomainProfile, answers: dict[str, Any]):
    """A deterministic model for ``profile``: the inventory lists each lane's samples, a field question gets
    ``answers[field]`` for the listed samples in order (or nothing), and matching pairs the lists by position."""

    def respond(system: str, user: str) -> str:
        if user.startswith("Paper excerpts"):
            cited = _SOURCE.findall(user)[:1]
            samples = [
                {"sample_id": sid, "label": label, "conditions": {}, "source_ids": cited}
                for sid, label in INVENTORY[_lane(user)]
            ]
            return json.dumps({"samples": samples, profile.prompt.no_samples_key: False})
        if user.startswith("List A (parser:"):
            list_a, list_b = user.split("\n\nList B (parser:", 1)
            pairs = [
                {"a": a, "b": b, "confidence": 0.9, "justification": "same calcination temperature"}
                for a, b in zip(_LISTED_ID.findall(list_a), _LISTED_ID.findall(list_b), strict=True)
            ]
            return json.dumps({"pairs": pairs, "unmatched_a": [], "unmatched_b": []})
        field = _FIELD.match(user)
        if field is None:
            raise AssertionError(f"an unexpected request: {user[:200]!r}")
        listed = _LISTED_ID.findall(user.split("Excerpts (Markdown", 1)[0])
        excerpts = user.split("Excerpts (Markdown with provenance markers):", 1)[1]
        values = [
            {
                "sample_id": sample_id,
                "value_raw": value_raw,
                "unit_raw": unit_raw,
                "condition": condition,
                "source_ids": _citing(excerpts, value_raw),
            }
            for sample_id, (value_raw, unit_raw, condition) in zip(
                listed, answers.get(field.group(1), ()), strict=False
            )
        ]
        return json.dumps({"values": values})

    return respond


def seed_parses(document: DocumentInput, layout: DataLayout, geometry: DocumentGeometry) -> None:
    """Both parsers' native output, written where a finished parse would have left it."""
    blocks = (TITLE, *PARAGRAPHS)
    for backend in BACKENDS:
        raw_dir = layout.raw_dir(document.document_id, backend)
        factory = RawOutputFactory(raw_dir.parent, document, geometry)
        if backend == "mineru":
            content = [
                {"type": "text", "page_idx": 0, "bbox": [100, 80 + 150 * i, 900, 200 + 150 * i], "text": text}
                | ({"text_level": 1} if i == 0 else {})
                for i, text in enumerate(blocks)
            ]
            factory.mineru(content, dir_name=raw_dir.name)
        else:
            parsing = [
                {
                    "block_id": i,
                    "block_order": i,
                    "block_label": "doc_title" if i == 0 else "text",
                    "block_bbox": [200, 150 + 350 * i, 1450, 450 + 350 * i],
                    "block_content": text,
                }
                for i, text in enumerate(blocks)
            ]
            page = paddle_page_entry(0, {"res": {"parsing_res_list": parsing}}, (1654, 2339))
            factory.paddle([page], dir_name=raw_dir.name)


@pytest.fixture
def paper(tmp_path: Path, document: DocumentInput, geometry: DocumentGeometry) -> tuple[DocumentInput, Settings]:
    """The synthetic cathode paper, both parses already in a fresh data_root."""
    settings = Settings(data_root=tmp_path / "data", repo_root=tmp_path, llm_api_key="sk-test")
    seed_parses(document, DataLayout(settings.data_root), geometry)
    return document, settings


def run_under(
    profile_path: Path, document: DocumentInput, settings: Settings, answers: dict[str, Any], monkeypatch
) -> tuple[PipelineResult, DomainProfile, Settings, FakeLlmClient]:
    settings = dataclasses.replace(settings, profile=str(profile_path))
    profile = load_run_profile(settings)
    client = FakeLlmClient(scripted_model(profile, answers))
    monkeypatch.setattr("paperfacts.workflow.build_llm_client", lambda _settings: client)
    return run_document(document, settings, profile), profile, settings, client


# ---- AC-7: the profile loads with no code change and is listed ------------------------------------------------


def test_the_battery_profile_loads_as_an_example_with_no_paper_level_field():
    profile = load_profile(BATTERY_PROFILE_PATH)

    assert profile.name == "battery_cathode" and profile.maturity == "example"
    assert profile.paper_groups == () and profile.paper_fields == ()
    assert profile.figures is None and profile.figure_fields == ()
    # Every field that must carry its condition also says what the dataset notes when it does not.
    assert all(spec.missing_condition_note_zh for spec in profile.fields if spec.condition_rule)
    assert {spec.canonical_unit for spec in profile.fields} >= {"mAh/g", "C", "V", "℃", "min", "nm", "%"}


def test_its_declared_units_convert():
    units = load_profile(BATTERY_PROFILE_PATH).units

    assert units.convert("mAh/g", "Ah/g") == (1000.0, 0.0)
    assert units.convert("mAh/g", "mAh g-1") == (1.0, 0.0)
    assert units.convert("V", "mV") == (0.001, 0.0)
    assert units.convert("℃", "K") == (1.0, -273.15)
    # The C-rate unit is case-sensitive: a lower-case "c" is no rate.
    assert units.convert("C", "C") == (1.0, 0.0) and units.convert("C", "c") is None


def test_the_profiles_command_lists_both_profiles():
    from typer.testing import CliRunner

    from paperfacts.cli import app

    result = CliRunner().invoke(app, ["profiles"])

    assert result.exit_code == 0, result.output
    names = [line.split()[0] for line in result.output.splitlines() if line.strip()]
    assert {"battery_cathode", "tco"} <= set(names)


# ---- AC-7 / AC-8: a battery run through the real pipeline ----------------------------------------------------


def test_a_battery_run_has_the_battery_columns_converts_its_units_and_refuses_a_window(paper, monkeypatch):
    document, settings = paper

    result, profile, _, client = run_under(BATTERY_PROFILE_PATH, document, settings, ANSWERS, monkeypatch)

    # All three kinds of question were asked, each under the profile's own wording.
    asked = {user.split("\n", 1)[0].split(" (")[0] for user in client.users}
    assert asked >= {"Paper excerpts", "Field to extract:", "List A"}
    assert all("lithium-ion battery cathode" in system or "cathode samples" in system for system in client.systems)
    dataset = result.dataset
    assert not dataset.incomplete
    assert set(profile.by_name) <= set(dataset.paper_row)
    # A matched pair's row is named after both lanes' ids, A's first.
    rows = {row["sample_id"].split(" | ")[0]: row for row in dataset.sample_rows}
    assert set(rows) == {"NCM-800", "NCM-900"}
    # AC-8: "0.15 Ah/g" is 150 mAh/g; the declared ℃ extension reads 1173 K as 899.85 ℃; 12 h is 720 min.
    assert rows["NCM-800"]["initial_discharge_capacity"] == pytest.approx(150.0)
    assert rows["NCM-900"]["initial_discharge_capacity"] == pytest.approx(170.0)
    assert rows["NCM-900"]["calcination_temperature"] == pytest.approx(899.85)
    assert rows["NCM-800"]["calcination_time"] == pytest.approx(720.0)
    assert rows["NCM-800"]["capacity_retention"] == pytest.approx(92.5)

    # AC-8: a window quoted for the cut-off is refused, with the reason, rather than read as its midpoint.
    for lane in result.lanes.values():
        for sample in lane.samples:
            (cutoff,) = [value for value in sample.fields if value.field == "upper_cutoff_voltage"]
            assert cutoff.value is None
            assert "refused (range_policy 'reject')" in (cutoff.normalization_note or "")
    cutoffs = [row for row in dataset.quality_rows if row["field"] == "upper_cutoff_voltage"]
    assert len(cutoffs) == 2
    for row in cutoffs:
        assert row["value"] is None and row["decision"] == "non_scalar"
        assert "2.8-4.3 V" in row["detail"]
    assert all(row["upper_cutoff_voltage"] is None for row in dataset.sample_rows)

    assert result.excel_path == DataLayout(settings.data_root).doc_dir(document.document_id) / "exports" / (
        "battery_cathode.xlsx"
    )
    workbook = load_workbook(result.excel_path)
    assert any(
        "首次放电比容量" in str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )


# ---- AC-9: two profiles, one data_root -------------------------------------------------------------------------


def test_battery_and_tco_runs_of_one_document_coexist(paper, monkeypatch):
    document, settings = paper
    layout = DataLayout(settings.data_root)

    battery, battery_profile, battery_settings, _ = run_under(
        BATTERY_PROFILE_PATH, document, settings, ANSWERS, monkeypatch
    )
    battery_workbook = battery.excel_path.read_bytes()
    battery_dataset = battery.dataset_json_path.read_bytes()

    tco, tco_profile, tco_settings, tco_client = run_under(SHIPPED_PROFILE_PATH, document, settings, {}, monkeypatch)
    # Nothing was cached for TCO: its questions really were asked, under its own system prompts.
    assert tco_client.call_count > 0
    assert all("cathode" not in system for system in tco_client.systems)

    keys = {
        name: (extractor_key_for(s, p, "fake-model"), comparison_key_for(s, p), figure_key_for(s, p))
        for name, s, p in (("battery", battery_settings, battery_profile), ("tco", tco_settings, tco_profile))
    }
    assert keys["battery"][0] != keys["tco"][0]
    assert keys["battery"][1] != keys["tco"][1]
    assert keys["battery"][2] != keys["tco"][2]
    assert battery.dataset.extractor_key != tco.dataset.extractor_key
    # Separate facts, datasets and workbooks, and the battery ones untouched by the TCO run.
    for backend in BACKENDS:
        for extractor in (battery.dataset.extractor_key, tco.dataset.extractor_key):
            assert layout.extraction_path(document.document_id, backend, extractor).is_file()
    assert tco.excel_path == battery.excel_path.with_name("tco.xlsx")
    assert battery.excel_path.read_bytes() == battery_workbook
    assert battery.dataset_json_path.read_bytes() == battery_dataset
    assert tco.dataset_json_path is not None and tco.dataset_json_path != battery.dataset_json_path
    assert set(tco_profile.by_name) <= set(tco.dataset.paper_row)
    assert not set(tco_profile.by_name) & set(battery.dataset.paper_row)

    # The battery run is still whole under its own keys: running it again asks the model nothing.
    again, _, _, client = run_under(BATTERY_PROFILE_PATH, document, settings, ANSWERS, monkeypatch)
    assert client.call_count == 0
    assert again.dataset.sample_rows == battery.dataset.sample_rows


# ---- a list field (cardinality: many) end to end ----------------------------------------------------------------

PRECURSORS = {
    "name": "precursors",
    "group": "synthesis",
    "kind": "composition",
    "cardinality": "many",
    "description": "Each metal salt or lithium source used to prepare the cathode, as named.",
    "keywords": ["precursors"],
    "label": "前驱体",
}
PRECURSOR_PARAGRAPH = "The precursors were NiSO4, CoSO4 and MnSO4, lithiated with LiOH."
# Per lane, the precursors it reads for every sample: the PaddleOCR-VL lane misses CoSO4 and MnSO4.
PRECURSOR_ANSWERS = {"mineru": ("NiSO4", "CoSO4", "LiOH"), "paddleocr_vl": ("LiOH", "NiSO4")}


def test_a_list_field_runs_end_to_end_as_the_union_of_both_lanes(
    tmp_path: Path, document: DocumentInput, geometry: DocumentGeometry, monkeypatch
):
    data = json.loads(BATTERY_PROFILE_PATH.read_text(encoding="utf-8"))
    data["fields"].append(PRECURSORS)
    profile_dir = tmp_path / "profiles"
    profile_dir.mkdir()
    list_profile = profile_dir / "battery_cathode.json"
    list_profile.write_text(json.dumps(data), encoding="utf-8")
    monkeypatch.setattr(f"{__name__}.PARAGRAPHS", (*PARAGRAPHS, PRECURSOR_PARAGRAPH))
    settings = Settings(data_root=tmp_path / "data", repo_root=tmp_path, llm_api_key="sk-test")
    seed_parses(document, DataLayout(settings.data_root), geometry)
    settings = dataclasses.replace(settings, profile=str(list_profile))
    profile = load_run_profile(settings)
    base = scripted_model(profile, ANSWERS)

    def respond(system: str, user: str) -> str:
        field = _FIELD.match(user)
        if field is None or field.group(1) != "precursors":
            return base(system, user)
        # The field line asks for one entry per value.
        assert "Several values may hold at once: report each as its own entry." in user
        excerpts = user.split("Excerpts (Markdown with provenance markers):", 1)[1]
        listed = _LISTED_ID.findall(user.split("Excerpts (Markdown", 1)[0])
        values = [
            {"sample_id": sample_id, "value_raw": raw, "source_ids": _citing(excerpts, raw)}
            for sample_id in listed
            for raw in PRECURSOR_ANSWERS[_lane(user)]
        ]
        return json.dumps({"values": values})

    client = FakeLlmClient(respond)
    monkeypatch.setattr("paperfacts.workflow.build_llm_client", lambda _settings: client)

    result = run_document(document, settings, profile)

    # A set: the two lanes' lists in another order agree element by element, and a list never conflicts.
    precursors = [c for c in result.report.comparisons if c.field == "precursors"]
    assert {c.status for c in precursors} == {"agree", "missing"}
    assert sorted(c.a.value_raw for c in precursors if c.status == "missing" and c.a) == ["CoSO4", "CoSO4"]
    rows = {row["sample_id"].split(" | ")[0]: row for row in result.dataset.sample_rows}
    for row in rows.values():
        assert row["precursors"] == ["NiSO4", "CoSO4", "LiOH"]
    quality = [row for row in result.dataset.quality_rows if row["field"] == "precursors"]
    assert {row["decision"] for row in quality} == {"single_source"}
    assert all(
        "CoSO4（mineru）" in row["detail"] and "LiOH（mineru, paddleocr_vl）" in row["detail"] for row in quality
    )
    samples = list(load_workbook(result.excel_path)["样品数据"].values)
    column = samples[0].index("precursors")
    assert {line[column] for line in samples[1:]} == {"NiSO4; CoSO4; LiOH"}
