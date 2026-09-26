"""The workbook under entity types (round 2 spec §5.6): one data sheet per entity, and a 实体 column on 字段说明 and
数据质量 when there are several. A profile without entity types keeps today's workbook sheet by sheet: the TCO
workbook is pinned against a snapshot written before entity types reached the workbook.

Regenerate the snapshot only for an intended change to the TCO workbook::

    PYTHONPATH=src:tests uv run python tests/test_workbook_entities.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from paperfacts.dataset import DocumentDataset
from paperfacts.profile import DomainProfile
from paperfacts.workbook import write_dataset
from support.profiles import make_entity_profile, shipped_profile

SNAPSHOT = Path(__file__).parent / "fixtures" / "workbook" / "tco_workbook.json"


def workbook_shape(path: Path) -> dict[str, Any]:
    """Everything a reader of the workbook sees, sheet by sheet: cells with their type and number format, column
    widths, frozen panes, the filter and the tables."""
    book = load_workbook(path)
    return {
        sheet.title: {
            "cells": [[[cell.value, cell.data_type, cell.number_format] for cell in row] for row in sheet.iter_rows()],
            "widths": {key: dim.width for key, dim in sorted(sheet.column_dimensions.items())},
            "freeze": sheet.freeze_panes,
            "filter": sheet.auto_filter.ref,
            "tables": dict(sorted(sheet.tables.items())),
        }
        for sheet in book.worksheets
    }


def tco_dataset(profile: DomainProfile) -> DocumentDataset:
    """A TCO document with a paper row, two samples, their quality rows and every kind of TCO cell."""
    metadata = {"document_id": "0123456789abcdef", "filename": "paper.pdf"}
    values = {
        "component": "In2O3:SnO2=90:10",
        "resistance": None,
        "density": 7.1,
        "sputtering_power": 120.0,
        "mode": "RF",
        "sheet_resistance": 12.5,
        "resistivity": 3.2e-4,
        "transmittance": 88.0,
        "thickness": 150.0,
    }
    samples = tuple(
        {
            **metadata,
            "sample_id": sample_id,
            "sample_label": f"{sample_id} label",
            "conditions": "substrate_temperature=300 °C",
            "available_fields": 8,
            "agree_fields": 6,
            **{spec.name: values.get(spec.name) for spec in profile.fields},
        }
        for sample_id in ("S1", "S2 | S2'")
    )
    quality = tuple(
        {
            **metadata,
            "sample_id": row["sample_id"],
            "field": name,
            "decision": "agree",
            "value": row[name],
            "unit": profile.by_name[name].canonical_unit,
            "conditions": "",
            "source_ids": "mineru_p0_b1",
            "lanes": "mineru; paddleocr_vl",
            "series": False,
            "detail": "=starts with an equals sign",
        }
        for row in samples
        for name in ("sheet_resistance", "mode")
    )
    return DocumentDataset(
        "0123456789abcdef",
        "paper.pdf",
        samples[0],
        samples,
        quality,
        extractor_key="e" * 12,
        comparison_key="c" * 12,
    )


def write_tco(path: Path) -> Path:
    profile = shipped_profile()
    write_dataset([tco_dataset(profile)], path, profile, failures=[{"document_id": "x", "error": "boom"}])
    return path


def test_the_tco_workbook_is_unchanged_sheet_by_sheet(tmp_path):
    shape = workbook_shape(write_tco(tmp_path / "tco.xlsx"))
    expected = json.loads(SNAPSHOT.read_text(encoding="utf-8"))

    assert list(shape) == list(expected) == ["论文数据", "样品数据", "字段说明", "数据质量", "图中读数", "运行记录"]
    for title in expected:
        assert json.loads(json.dumps(shape[title], ensure_ascii=False)) == expected[title], title


def two_entity_dataset(profile: DomainProfile) -> DocumentDataset:
    """Two coatings and one wear test, as ``consolidate_document`` writes them for a two-entity profile."""
    metadata = {"document_id": "0123456789abcdef", "filename": "paper.pdf"}
    paper = {"precursor_purity": 99.9}

    def row(entity: str, sample_id: str, **values: Any) -> dict[str, Any]:
        return {
            **metadata,
            "entity": entity,
            "sample_id": sample_id,
            "sample_label": "",
            "conditions": "",
            "available_fields": 2,
            "agree_fields": 2,
            **paper,
            **values,
        }

    coatings = (row("coating", "S1", coating_thickness=100.0, solvent="water"), row("coating", "S2"))
    test = row("wear_test", "T1", test_temperature=300.0, wear_mode="sliding")
    quality = (
        {**metadata, "sample_id": "paper", "field": "precursor_purity", "decision": "agree", "value": 99.9},
        {**metadata, "entity": "wear_test", "sample_id": "T1", "field": "wear_mode", "decision": "agree"},
    )
    return DocumentDataset("0123456789abcdef", "paper.pdf", coatings[0], (*coatings, test), quality)


def test_each_entity_has_its_own_data_sheet_with_its_own_fields(tmp_path):
    profile = make_entity_profile({"entities.1.label_zh": "磨损/测试[1]"})
    path = tmp_path / "demo.xlsx"
    write_dataset([two_entity_dataset(profile)], path, profile)
    book = load_workbook(path)

    # The title loses what Excel refuses in one.
    assert book.sheetnames == ["论文数据", "涂层数据", "磨损测试1数据", "字段说明", "数据质量", "图中读数", "运行记录"]
    coatings = list(book["涂层数据"].values)
    assert coatings[0][-3:] == ("precursor_purity", "coating_thickness", "solvent")
    assert [line[2] for line in coatings[1:]] == ["S1", "S2"]
    tests = list(book["磨损测试1数据"].values)
    assert tests[0][-3:] == ("precursor_purity", "test_temperature", "wear_mode")
    assert tests[1][2:3] == ("T1",) and tests[1][-2:] == (300.0, "sliding")
    assert book["磨损测试1数据"].freeze_panes == "D2" and set(book["磨损测试1数据"].tables) == {"Samples_wear_test"}
    # The paper sheet has the paper-level fields and the primary entity's.
    assert next(iter(book["论文数据"].values))[-3:] == ("precursor_purity", "coating_thickness", "solvent")


def test_the_field_and_quality_sheets_name_each_row_s_entity(tmp_path):
    profile = make_entity_profile()
    path = tmp_path / "demo.xlsx"
    write_dataset([two_entity_dataset(profile)], path, profile)
    book = load_workbook(path)

    fields = list(book["字段说明"].values)
    assert fields[0][:4] == ("字段", "中文名", "层级", "实体")
    entity = {line[0]: line[3] for line in fields[1:]}
    assert entity == {
        "precursor_purity": None,
        "coating_thickness": "涂层",
        "solvent": "涂层",
        "test_temperature": "磨损测试",
        "wear_mode": "磨损测试",
    }
    quality = list(book["数据质量"].values)
    assert quality[0][:4] == ("文档ID", "文件名", "实体", "样品ID")
    assert [line[2] for line in quality[1:]] == [None, "磨损测试"]
    assert book["数据质量"].freeze_panes == "E2"


def test_an_entity_without_a_label_names_its_sheet_by_its_name_and_a_clash_is_numbered(tmp_path):
    profile = make_entity_profile({"entities.0.label_zh": "", "entities.1.label_zh": "论文"})
    path = tmp_path / "demo.xlsx"
    write_dataset([two_entity_dataset(profile)], path, profile)

    assert load_workbook(path).sheetnames[1:3] == ["coating数据", "论文数据 2"]


if __name__ == "__main__":
    import tempfile

    with tempfile.TemporaryDirectory() as directory:
        shape = workbook_shape(write_tco(Path(directory) / "tco.xlsx"))
    SNAPSHOT.write_text(json.dumps(shape, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"wrote {SNAPSHOT}", file=sys.stderr)
