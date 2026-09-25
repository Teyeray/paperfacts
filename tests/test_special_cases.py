"""The former field-name special cases, now profile attributes: any profile's field gets them by declaring them."""

from openpyxl import load_workbook

from paperfacts.compare import FieldComparison
from paperfacts.dataset import DocumentDataset
from paperfacts.decide import decide
from paperfacts.records import FieldValue
from paperfacts.workbook import write_dataset
from support.extraction import make_lane
from support.profiles import make_profile
from test_dataset import dataset

_RULE = {
    "fields.1.condition_hint": "annealing state, e.g. 'as-deposited'",
    "fields.1.condition_rule": "the annealing state",
}


def _single(name, *, condition=None, raw="120", unit="nm"):
    return [("mineru", FieldValue(field=name, value_raw=raw, unit_raw=unit, condition=condition, source_ids=("b",)))]


def _decide(spec, evidence, units):
    """One lane's value, which the comparison saw only on that lane."""
    comparison = FieldComparison(scope="sample:A", field=spec.name, status="missing", a=evidence[0][1])
    return decide(spec, evidence, [comparison], units=units)


def test_a_field_with_a_condition_rule_and_no_note_gets_the_generic_note():
    profile = make_profile({**_RULE, "fields.1.label": "涂层厚度"})
    spec = profile.by_name["coating_thickness"]

    result = _decide(spec, _single("coating_thickness"), profile.units)

    assert result.value == 120
    assert "原文提取结果未注明涂层厚度的测量条件" in result.detail


def test_the_generic_note_falls_back_to_the_field_name_without_a_label():
    profile = make_profile(_RULE)

    result = _decide(profile.by_name["coating_thickness"], _single("coating_thickness"), profile.units)

    assert "原文提取结果未注明coating_thickness的测量条件" in result.detail


def test_a_custom_note_replaces_the_generic_one():
    profile = make_profile({**_RULE, "fields.1.missing_condition_note_zh": "未注明退火状态"})

    result = _decide(profile.by_name["coating_thickness"], _single("coating_thickness"), profile.units)

    assert "未注明退火状态" in result.detail
    assert "测量条件" not in result.detail


def test_no_note_when_the_condition_is_stated_or_the_field_has_no_rule():
    ruled = make_profile(_RULE)
    stated = _decide(
        ruled.by_name["coating_thickness"], _single("coating_thickness", condition="annealed"), ruled.units
    )
    plain = make_profile()
    unruled = _decide(plain.by_name["coating_thickness"], _single("coating_thickness"), plain.units)

    assert "未注明" not in stated.detail
    assert "未注明" not in unruled.detail


def test_tco_transmittance_keeps_its_own_note(tco_profile):
    evidence = _single("transmittance", raw="90", unit="%")
    result = _decide(tco_profile.by_name["transmittance"], evidence, tco_profile.units)

    assert "原文提取结果未注明透光率波长或波段" in result.detail


def test_display_format_scientific_formats_that_fields_cells(tmp_path):
    profile = make_profile({"fields.1.display_format": "scientific"})
    row = {"document_id": "d", "filename": "p.pdf", "sample_id": "S", "precursor_purity": 99.5}
    document = DocumentDataset(
        "d", "p.pdf", {**row, "coating_thickness": 120.0}, ({**row, "coating_thickness": 120.0},), (), "ek", "ck"
    )
    output = tmp_path / "dataset.xlsx"

    write_dataset([document], output, profile)

    for title in ("论文数据", "样品数据"):
        sheet = load_workbook(output)[title]
        columns = {cell.value: cell.column for cell in sheet[1]}
        assert sheet.cell(2, columns["coating_thickness"]).number_format == "0.0000E+00"
        assert sheet.cell(2, columns["precursor_purity"]).number_format == "0.############"


def test_a_paper_row_without_samples_names_the_paper_level_fields_generically():
    # One of the two documented TCO detail strings (tests/fixtures/s6_detail_strings/README.md).
    selection = next(row for row in dataset(make_lane()).quality_rows if row["field"] == "__selection__")

    assert selection["detail"] == "未提取到可匹配样品；论文行仅保留唯一的论文级字段"
