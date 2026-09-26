"""The kinds ``boolean``, ``date`` and ``interval``: what the model is asked, what survives cleaning and voting, how
a quote is read, when two lanes agree, and what a cell and a workbook column hold.

A profile without such a field -- TCO -- must not notice any of it: its prompts, answer shapes and stored files
keep their bytes (the snapshot and payload pins hold the prompts; the identity tests here hold the classes).
"""

from __future__ import annotations

import typing
from pathlib import Path

import pytest
from openpyxl import load_workbook

from paperfacts.columns import FieldColumn
from paperfacts.compare import FieldComparison, compare_values
from paperfacts.decide import decide
from paperfacts.errors import ConfigError
from paperfacts.extract import FieldHarvest, SampleInventory, _response_models, passage_records
from paperfacts.kinds import RULES
from paperfacts.normalize import drop_implausible, normalize_field
from paperfacts.profile import DomainProfile
from paperfacts.prompts import extraction_system_prompt, field_system_prompt, render_field_table
from paperfacts.readers import read_date
from paperfacts.records import (
    NO_CONTEXT,
    ExtractedRecords,
    FieldResponse,
    FieldValue,
    InventoryResponse,
    InventorySample,
    PaperRecord,
    ResponseField,
    ResponsePaper,
    ResponseSample,
    ResponseValue,
    SampleRecord,
    response_models,
    response_to_records,
)
from paperfacts.voting import deduplicate, merge_passes
from paperfacts.workbook import format_cell, write_dataset
from support.profiles import make_profile, profile_data

_NEW_FIELDS = [
    {
        "name": "doped",
        "group": "coating",
        "kind": "boolean",
        "description": "Whether the coating is intentionally doped.",
        "keywords": ["doped", "dopant"],
    },
    {
        "name": "prepared_on",
        "group": "precursor",
        "kind": "date",
        "description": "Date the precursor was prepared.",
        "keywords": ["prepared"],
    },
    {
        "name": "annealing_window",
        "group": "coating",
        "kind": "interval",
        "description": "Temperature window the coating was annealed in.",
        "keywords": ["annealed"],
        "canonical_unit": "℃",
        "rel_tol": 0.01,
        "valid_range": {"min": 0, "max": 1500},
    },
    {
        "name": "coverage",
        "group": "coating",
        "kind": "interval",
        "description": "Surface coverage of the coating.",
        "keywords": ["coverage"],
        "canonical_unit": "%",
    },
]
PROFILE = make_profile({"fields": [*profile_data()["fields"], *_NEW_FIELDS]})
UNITS = PROFILE.units
SPECS = PROFILE.by_name


def _value(field: str, value_raw: str, **update: object) -> FieldValue:
    return FieldValue(field=field, value_raw=value_raw, source_ids=("b1",), **update)  # type: ignore[arg-type]


def _read(field: str, value_raw: str, **update: object) -> FieldValue:
    return normalize_field(_value(field, value_raw, **update), SPECS[field], UNITS, NO_CONTEXT)


# ---- Response models ----------------------------------------------------------------------------------------


def test_without_a_boolean_field_the_answer_shapes_are_the_module_level_classes():
    models = response_models("target", "no_tco_film", holds=False)

    assert models.field is FieldResponse
    assert models.extraction.model_fields["samples"].annotation == list[ResponseSample]
    assert models.extraction.model_fields["paper"].annotation == ResponsePaper | None
    assert ResponseSample.model_fields["fields"].annotation == list[ResponseField]
    assert "holds" not in ResponseField.model_fields and "holds" not in ResponseValue.model_fields


def test_tco_asks_with_the_module_level_field_answer(tco_profile: DomainProfile):
    assert _response_models(tco_profile).field is FieldResponse
    assert _response_models(PROFILE).field is not FieldResponse


def test_with_a_boolean_field_every_container_is_rebuilt_under_its_own_name():
    models = response_models("paper", "no_samples", holds=True)
    extraction = models.extraction
    sample = typing.get_args(extraction.model_fields["samples"].annotation)[0]
    paper = next(arg for arg in typing.get_args(extraction.model_fields["paper"].annotation) if arg is not type(None))
    value = typing.get_args(models.field.model_fields["values"].annotation)[0]

    for rebuilt, base in (
        (sample, ResponseSample),
        (paper, ResponsePaper),
        (models.field, FieldResponse),
        (value, ResponseValue),
    ):
        assert rebuilt is not base and rebuilt.__name__ == base.__name__
    for container in (sample, paper):
        assert "holds" in typing.get_args(container.model_fields["fields"].annotation)[0].model_fields
    assert "holds" in value.model_fields


def test_holds_survives_every_container_of_a_document_mode_answer():
    extraction = response_models("paper", "no_samples", holds=True).extraction
    answer = extraction.model_validate_json(
        """{"paper": {"source_ids": ["b1"], "fields": [
                {"field": "doped", "value_raw": "all films were doped", "source_ids": ["b1"],
                 "applies_to_all_samples": true, "holds": true}]},
            "samples": [{"sample_id": "S1", "source_ids": ["b1"], "fields": [
                {"field": "doped", "value_raw": "undoped", "source_ids": ["b1"], "holds": false},
                {"field": "solvent", "value_raw": "ethanol", "source_ids": ["b1"], "holds": true}]}]}"""
    )

    records = response_to_records(answer, fields=SPECS, known_ids=frozenset({"b1"}))

    (sample,) = records.samples
    assert [(value.field, value.holds) for value in sample.fields] == [
        ("doped", True),  # the series value, fanned out
        ("doped", False),
        ("solvent", None),  # holds is a boolean field's alone
    ]


def test_holds_survives_a_passage_mode_answer_and_a_boolean_without_it_is_dropped():
    field = response_models("paper", "no_samples", holds=True).field
    answer = field.model_validate_json(
        """{"values": [{"sample_id": "S1", "value_raw": "undoped", "source_ids": ["b1"], "holds": false},
                       {"sample_id": "S2", "value_raw": "doped", "source_ids": ["b1"]}]}"""
    )
    inventory = InventoryResponse(samples=[InventorySample(sample_id="S1"), InventorySample(sample_id="S2")])
    harvest = FieldHarvest(spec=SPECS["doped"], values=tuple(answer.values), known_ids=frozenset({"b1"}))

    listed = SampleInventory(inventory, "", {}, frozenset(), PROFILE.primary)
    records = passage_records([listed], [harvest], PROFILE)

    assert [(s.sample_id, [v.holds for v in s.fields]) for s in records.samples] == [("S1", [False]), ("S2", [])]
    assert "doped: boolean field without holds 'doped'" in records.dropped


# ---- Prompts ------------------------------------------------------------------------------------------------


def test_the_holds_key_is_asked_for_only_by_a_profile_with_a_boolean_field(tco_profile: DomainProfile):
    key = '"applies_to_all_samples": <true|false>,\n         "holds": <true|false for a boolean field, else null>}'

    for prompt in (field_system_prompt, extraction_system_prompt):
        assert key in prompt(PROFILE)
        assert '"holds"' not in prompt(tco_profile)
        assert '"applies_to_all_samples": <true|false>}' in prompt(tco_profile)


@pytest.mark.parametrize(
    ("field", "note"),
    [
        ("doped", "set holds to true when they affirm it, false when they deny it."),
        ("prepared_on", " Quote the date exactly as written."),
        ("annealing_window", 'or a one-sided bound (">80 %").'),
    ],
)
def test_the_field_line_carries_the_kind_note(field: str, note: str):
    assert note in render_field_table((SPECS[field],), "another quantity")


# ---- Stored bytes -------------------------------------------------------------------------------------------


def test_the_new_attributes_are_left_out_of_a_file_when_null():
    plain = _value("coating_thickness", "100").model_dump()
    filled = _value("doped", "doped", holds=False, iso_date="2021", bounds=(1.0, None)).model_dump()

    assert not {"holds", "iso_date", "bounds"} & set(plain)
    assert (filled["holds"], filled["iso_date"], filled["bounds"]) == (False, "2021", (1.0, None))


# ---- Voting -------------------------------------------------------------------------------------------------


def _pass(*holds: bool) -> ExtractedRecords:
    fields = tuple(_value("doped", "doped", holds=value) for value in holds)
    return ExtractedRecords(
        paper=None,
        samples=(SampleRecord(sample_id="S1", fields=fields),),
        invalid_source_ids=(),
        dropped=(),
    )


def test_a_true_pass_and_a_false_pass_never_vote_as_one_value():
    merged = merge_passes([_pass(True), _pass(False)], reference_fields=())

    assert merged.samples[0].fields == ()
    assert any("only 1/2 passes" in entry for entry in merged.dropped)


def test_the_majority_of_passes_decides_holds():
    merged = merge_passes([_pass(True), _pass(False), _pass(True)], reference_fields=())

    (value,) = merged.samples[0].fields
    assert (value.holds, value.agreement) == (True, pytest.approx(2 / 3))


def test_one_pass_keeps_a_true_and_a_false_quote_apart():
    assert [value.holds for value in deduplicate(_pass(True, False, True), reference_fields=()).samples[0].fields] == [
        True,
        False,
    ]


# ---- Boolean ------------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("a", "b", "status"),
    [(True, True, "agree"), (False, False, "agree"), (True, False, "conflict"), (True, None, "ambiguous")],
)
def test_two_booleans_agree_when_their_holds_do(a, b, status):
    spec = SPECS["doped"]
    assert (
        compare_values(_value("doped", "doped", holds=a), _value("doped", "Doped", holds=b), spec, NO_CONTEXT)[0]
        == status
    )


def test_a_boolean_cell_is_its_holds():
    field = _value("doped", "undoped", holds=False)
    comparison = FieldComparison(scope="sample:S1", field="doped", status="agree", a=field, b=field)

    decision = decide(
        SPECS["doped"], [("mineru", field), ("paddleocr_vl", field)], [comparison], units=UNITS, ctx=NO_CONTEXT
    )

    assert (decision.status, decision.value) == ("agree", False)


# ---- Date ---------------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "iso"),
    [
        ("2021", "2021"),
        ("2021-03", "2021-03"),
        ("2021-03-12", "2021-03-12"),
        ("2021/3/12", "2021-03-12"),
        ("2021.03.12", "2021-03-12"),
        ("March 2021", "2021-03"),
        ("Mar. 2021", "2021-03"),
        ("Sept 2020", "2020-09"),
        ("12 March 2021", "2021-03-12"),
        ("March 12th, 2021", "2021-03-12"),
        ("2021 May 3", "2021-05-03"),
        ("March 2021.", "2021-03"),
        ("(March 2021)", "2021-03"),
        ("(2021-03-12).", "2021-03-12"),
        ("2021.", "2021"),
    ],
)
def test_a_date_is_read_to_iso_at_the_precision_written(raw: str, iso: str):
    assert read_date(raw) == (iso, None)


@pytest.mark.parametrize(
    ("raw", "why"),
    [
        ("Mar 21", "two-digit year"),
        ("03/12/21", "not year first"),
        ("03/04/2021", "not year first"),
        ("12.03.2021", "not year first"),
        ("2019-2021", "not year first"),
        ("March-May 2021", "a range"),
        ("March 2021 and April 2021", "not the one month"),
        ("1750", "outside 1800-2100"),
        ("2150", "outside 1800-2100"),
        ("2021-13", "no such calendar date"),
        ("Feb 30, 2021", "no such calendar date"),
        ("spring 2021", "not the one month"),
        ("2021.5", "decimal year"),
    ],
)
def test_a_date_that_could_be_read_two_ways_is_refused(raw: str, why: str):
    iso, note = read_date(raw)

    assert iso is None and why in (note or "")


@pytest.mark.parametrize(
    ("a", "b", "status"),
    [
        ("12 March 2021", "2021-03-12", "agree"),
        ("March 2021", "2021-03-12", "ambiguous"),
        ("2021", "2021-03", "ambiguous"),
        ("March 2021", "April 2021", "conflict"),
        ("2021-03-01", "2021-03-12", "conflict"),
        ("Mar 21", "2021-03", "ambiguous"),
    ],
)
def test_two_dates_agree_when_their_iso_dates_do(a: str, b: str, status: str):
    spec = SPECS["prepared_on"]
    assert compare_values(_read("prepared_on", a), _read("prepared_on", b), spec, NO_CONTEXT)[0] == status


def test_a_date_cell_is_its_iso_string():
    assert RULES["date"].cell(_value("prepared_on", "12 March 2021"), SPECS["prepared_on"], UNITS, NO_CONTEXT) == (
        "2021-03-12",
        None,
    )


# ---- Interval -----------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("field", "raw", "unit", "bound", "bounds"),
    [
        ("annealing_window", "450-500", "°C", None, (450.0, 500.0)),
        ("annealing_window", "450 to 500 °C", None, None, (450.0, 500.0)),
        ("annealing_window", "450 °C to 500 °C", "°C", None, (450.0, 500.0)),
        ("coverage", ">80", "%", None, (80.0, None)),
        ("coverage", "≥ 80", "%", None, (80.0, None)),
        ("coverage", "at most 5", "%", None, (None, 5.0)),
        # grounding found "above" right before the quote: exactly what quoting it would have given
        ("coverage", "80", "%", "above", (80.0, None)),
        ("coverage", "80", "%", "<", (None, 80.0)),
        ("coverage", ">80 %", None, None, (80.0, None)),  # the unit written only in the quote
    ],
)
def test_an_interval_reads_two_ends_or_a_bound_in_the_canonical_unit(field, raw, unit, bound, bounds):
    value = _read(field, raw, unit_raw=unit, bound=bound)

    assert value.bounds == pytest.approx(bounds)
    assert value.unit == SPECS[field].canonical_unit


@pytest.mark.parametrize(
    ("raw", "why"),
    [
        ("80", "an interval field needs two ends or a bound"),
        ("500-450", "descending"),
        ("> 450-500", "refused"),
        ("450-500 °C for 2 h", "needs two ends or a bound"),
        ("overall 450", "needs two ends or a bound"),  # "over" only as a word
    ],
)
def test_a_bare_scalar_or_an_unclean_range_is_no_interval(raw: str, why: str):
    value = _read("annealing_window", raw, unit_raw="°C")

    assert value.bounds is None and why in (value.normalization_note or "")


@pytest.mark.parametrize(
    ("a", "b", "status"),
    [
        ((450.0, 500.0), (452.0, 499.0), "agree"),  # rel_tol 0.01 on each end
        ((450.0, 500.0), (450.0, 550.0), "conflict"),
        ((450.0, None), (450.0, None), "agree"),
        ((450.0, None), (450.0, 500.0), "conflict"),  # an open end equals only an open end
        ((450.0, 500.0), None, "ambiguous"),
    ],
)
def test_two_intervals_agree_end_by_end(a, b, status):
    spec = SPECS["annealing_window"]
    side_a = _value("annealing_window", "x", bounds=a, unit="℃")
    side_b = _value("annealing_window", "y", bounds=b, unit="℃" if b else None)

    assert compare_values(side_a, side_b, spec, NO_CONTEXT)[0] == status


def test_the_pairing_distance_sums_the_finite_ends():
    rules = RULES["interval"]
    a = _value("annealing_window", "x", bounds=(450.0, 500.0))

    assert rules.distance(a, _value("annealing_window", "y", bounds=(455.0, 510.0))) == pytest.approx(15.0)
    assert rules.distance(a, _value("annealing_window", "y", bounds=(455.0, None))) is None


def test_an_interval_cell_is_its_two_ends():
    cell = RULES["interval"].cell(
        _value("coverage", "80", unit_raw="%", bound="above"), SPECS["coverage"], UNITS, NO_CONTEXT
    )

    assert cell == ([80.0, None], None)


def test_the_plausible_range_is_checked_on_each_finite_end():
    records = ExtractedRecords(
        paper=PaperRecord(),
        samples=(
            SampleRecord(
                sample_id="S1",
                fields=(
                    _value("annealing_window", "450-500", unit_raw="°C"),
                    _value("annealing_window", "1400-1600", unit_raw="°C"),
                    _value("annealing_window", ">1600", unit_raw="°C"),
                ),
            ),
        ),
        invalid_source_ids=(),
        dropped=(),
    )

    kept = drop_implausible(records, PROFILE)

    assert [value.value_raw for value in kept.samples[0].fields] == ["450-500"]
    assert "annealing_window: '1400-1600' °C is 1600 ℃, outside the plausible range" in kept.dropped[0]


# ---- Loader -------------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("changes", "message"),
    [
        ({"fields.3.canonical_unit": "nm"}, "canonical_unit is only meaningful for a numeric or interval field"),
        ({"fields.4.valid_range": {"min": 1}}, "valid_range is only meaningful for a numeric or interval field"),
        (
            {"fields.5.bare_number": "assume_canonical"},
            "bare_number is only meaningful for a numeric field, not a 'interval'",
        ),
        ({"fields.5.range_policy": "upper"}, "range_policy is only meaningful for a numeric field"),
        ({"fields.3.rel_tol": 0.1}, "rel_tol is only meaningful for a numeric or interval field, not a 'boolean' one"),
        ({"fields.4.abs_tol": 1}, "abs_tol is only meaningful for a numeric or interval field, not a 'date' one"),
        ({"fields.2.rel_tol": 0.1}, "rel_tol is only meaningful for a numeric or interval field, not a 'text' one"),
        (
            {"fields.3.condition_rule": "the dopant", "fields.3.condition_hint": "dopant"},
            "condition_rule is only meaningful for a numeric, composition, text, date or interval field",
        ),
    ],
)
def test_the_loader_refuses_an_attribute_the_kind_cannot_use(changes, message):
    data = {"fields": [*profile_data()["fields"], *_NEW_FIELDS]}
    fields = data["fields"]
    for dotted, value in changes.items():
        _, index, key = dotted.split(".")
        fields[int(index)] = {**fields[int(index)], key: value}

    with pytest.raises(ConfigError, match=message):
        make_profile(data)


def test_an_attribute_written_at_its_default_is_accepted_on_any_kind():
    # tco.json spells every attribute out; at its default an attribute says nothing about the kind.
    data = {"fields": [*profile_data()["fields"], *_NEW_FIELDS]}
    defaults = {"rel_tol": 0.0, "abs_tol": 0, "canonical_unit": None, "bare_number": "reject"}
    defaults |= {"range_policy": "midpoint", "display_format": "plain", "figure_readable": False}
    data["fields"][3] = {**data["fields"][3], **defaults}

    assert make_profile(data).by_name["doped"].kind == "boolean"


def test_an_interval_canonical_unit_is_checked_against_the_units():
    data = {"fields": [*profile_data()["fields"], {**_NEW_FIELDS[2], "canonical_unit": "furlong"}]}

    with pytest.raises(ConfigError, match="furlong"):
        make_profile(data)


# ---- Presentation -------------------------------------------------------------------------------------------


def _column(kind: str) -> FieldColumn:
    return FieldColumn(name="f", scope="sample", kind=kind)  # type: ignore[arg-type]


@pytest.mark.parametrize(
    ("kind", "value", "written"),
    [
        ("boolean", True, True),
        ("boolean", False, False),
        ("date", "2021-03", "2021-03"),
        ("interval", [2.8, 4.3], "2.8–4.3"),
        ("interval", [80.0, None], "≥ 80"),
        ("interval", [None, 5.0], "≤ 5"),
        ("interval", None, None),
    ],
)
def test_a_new_kind_cell_is_written_as_its_column_says(kind: str, value, written):
    assert format_cell(value, _column(kind)) == written


def test_the_workbook_gives_an_interval_two_numeric_columns(tmp_path: Path):
    from paperfacts.dataset import DocumentDataset

    row = {
        "document_id": "d",
        "sample_id": "S1",
        "doped": True,
        "prepared_on": None,
        "annealing_window": [450.0, 500.0],
        "coverage": [80.0, None],
    }
    dataset = DocumentDataset(
        document_id="d",
        filename="d.pdf",
        extractor_key="e",
        comparison_key="c",
        paper_row={"document_id": "d", "sample_id": "paper"},
        sample_rows=(row,),
        quality_rows=({"document_id": "d", "sample_id": "S1", "field": "annealing_window", "value": [450.0, 500.0]},),
    )
    output = tmp_path / "dataset.xlsx"

    write_dataset([dataset], output, PROFILE)

    book = load_workbook(output)
    samples = book["样品数据"]
    header = [cell.value for cell in samples[1]]
    values = dict(zip(header, (cell.value for cell in samples[2]), strict=True))
    assert "annealing_window" not in header
    assert (values["annealing_window 下限"], values["annealing_window 上限"]) == (450, 500)
    assert (values["coverage 下限"], values["coverage 上限"]) == (80, None)
    assert values["doped"] is True
    quality = book["数据质量"]
    assert [cell.value for cell in quality[2]][5] == "450–500"
    fields = {row[0].value: row[3].value for row in book["字段说明"].iter_rows(min_row=2)}
    assert (fields["doped"], fields["prepared_on"], fields["annealing_window"]) == ("是/否", "日期（ISO）", "℃")
    header = [cell.value for cell in book["字段说明"][1]]
    rules = {row[0].value: row[header.index("单值与缺失规则")].value for row in book["字段说明"].iter_rows(min_row=2)}
    assert rules["annealing_window"].startswith("区间：")
    assert "范围" in rules["doped"]


def test_two_quotes_read_as_opposite_booleans_are_two_values():
    from paperfacts.decide import _identity

    base = dict(field="doped", value_raw="Al-doped", unit_raw=None, condition=None, source_ids=["mineru_p0_b1"])
    assert _identity(FieldValue(**base, holds=True)) != _identity(FieldValue(**base, holds=False))


def test_a_power_of_ten_in_both_the_bound_and_the_unit_is_refused():
    value = _read("annealing_window", "> 4.5 × 10^2", unit_raw="×10^2 °C")

    assert value.bounds is None and "scale factor in both value and unit" in (value.normalization_note or "")


@pytest.mark.parametrize("field", ["doped", "prepared_on"])
def test_equal_answers_pair_first_whatever_order_the_lanes_list_them(field: str):
    from paperfacts.compare import _pair_values

    first = {"doped": dict(holds=True), "prepared_on": dict(value_raw="2021-03")}[field]
    second = {"doped": dict(holds=False), "prepared_on": dict(value_raw="2021-05")}[field]

    def side(*answers: dict) -> list[FieldValue]:
        return [
            _read(field, str(answer.get("value_raw", "doped")), **{k: v for k, v in answer.items() if k == "holds"})
            for answer in answers
        ]

    pairs = _pair_values(side(first, second), side(second, first), SPECS[field], NO_CONTEXT)

    assert all(compare_values(a, b, SPECS[field], NO_CONTEXT)[0] == "agree" for a, b in pairs if a and b)
    assert len(pairs) == 2
