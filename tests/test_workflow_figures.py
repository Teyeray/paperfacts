"""The figures stage inside the pipeline: where its input comes from, what it stores, and that nothing it
does -- switched off, failing, or half-failing -- can cost the paper its other stages.

The vision client is :class:`support.vision.FakeVisionClient`; the PDF is a generated blank one, so the
crops are real PNGs of an empty page and the requests never leave the process.
"""

from __future__ import annotations

import dataclasses
import threading
from pathlib import Path

import pytest
from openpyxl import load_workbook

import paperfacts.workflow as workflow
from paperfacts.config import Settings
from paperfacts.errors import Cancelled, ConfigError, LlmError
from paperfacts.figures import FigureReadings
from paperfacts.keys import figure_key_for
from paperfacts.models import Backend, DocumentInput, NormalizedBBox, PageGeometry, ParsedArtifact
from paperfacts.storage import DataLayout
from paperfacts.workflow import read_document_figures, run_document, shown_figures
from support.factories import make_block
from support.vision import NOT_A_CHART, FakeVisionClient, chart_answer
from test_workflow_run import install_fake_pipeline

BOX = NormalizedBBox(x1=0.1, y1=0.1, x2=0.6, y2=0.5)


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(data_root=tmp_path / "data", repo_root=tmp_path, llm_api_key="sk-test", figures_enabled=True)


def store_artifact(document: DocumentInput, settings: Settings, backend: Backend = "mineru") -> None:
    blocks = (
        make_block(
            page=0, order=0, type="figure", content="a.jpg", bbox=BOX, backend=backend, document_id=document.document_id
        ),
        make_block(
            page=0,
            order=1,
            type="caption",
            content="Fig. 2 Sheet resistance of the films versus O2 flow.",
            bbox=BOX,
            backend=backend,
            document_id=document.document_id,
        ),
    )
    artifact = ParsedArtifact(
        document_id=document.document_id,
        backend=backend,
        backend_version="x",
        pages=(PageGeometry(index=0, width_pt=595, height_pt=842), PageGeometry(index=1, width_pt=612, height_pt=792)),
        blocks=blocks,
    )
    artifact.write(DataLayout(settings.data_root).artifact_path(document.document_id, backend))


def figures_file(document: DocumentInput, settings: Settings) -> Path:
    return DataLayout(settings.data_root).figures_path(document.document_id, figure_key_for(settings))


# ---- read_document_figures ---------------------------------------------------------------------------


def test_the_charts_are_cropped_from_the_pdf_read_and_stored(document: DocumentInput, settings: Settings):
    store_artifact(document, settings)
    client = FakeVisionClient(chart_answer())

    readings = read_document_figures(document, settings, client)

    assert [call.image_png[:8] for call in client.calls] == [b"\x89PNG\r\n\x1a\n"]
    assert {r.source_id for r in readings.readings} == {"mineru_p0_b0"}
    assert FigureReadings.read(figures_file(document, settings)) == readings


def test_stored_readings_are_served_without_asking_again(document: DocumentInput, settings: Settings):
    store_artifact(document, settings)
    read_document_figures(document, settings, FakeVisionClient(chart_answer()))
    again = FakeVisionClient(chart_answer())

    read_document_figures(document, settings, again)

    assert again.calls == []


def test_a_stored_file_with_a_failed_request_is_read_again(document: DocumentInput, settings: Settings):
    store_artifact(document, settings)
    first = read_document_figures(document, settings, FakeVisionClient(LlmError("HTTP 504")))
    retry = FakeVisionClient(chart_answer())

    second = read_document_figures(document, settings, retry)

    assert not first.complete
    assert len(retry.calls) == 1 and second.complete and second.readings


def test_force_re_asks_the_model(document: DocumentInput, settings: Settings):
    store_artifact(document, settings)
    read_document_figures(document, settings, FakeVisionClient(chart_answer()))
    forced = FakeVisionClient(chart_answer())

    read_document_figures(document, settings, forced, force=True)

    assert [call.refresh for call in forced.calls] == [True]


def test_paddle_figure_blocks_stand_in_when_there_is_no_mineru_artifact(document: DocumentInput, settings: Settings):
    store_artifact(document, settings, backend="paddleocr_vl")

    readings = read_document_figures(document, settings, FakeVisionClient(chart_answer()))

    assert readings.backend == "paddleocr_vl"
    assert {r.source_id for r in readings.readings} == {"paddleocr_vl_p0_b0"}


def test_without_the_pdf_nothing_can_be_cropped(document: DocumentInput, settings: Settings, tmp_path: Path):
    store_artifact(document, settings)
    moved = document.model_copy(update={"pdf_path": tmp_path / "gone.pdf"})

    with pytest.raises(FileNotFoundError, match="PDF not available"):
        read_document_figures(moved, settings, FakeVisionClient(chart_answer()))


def test_without_a_parse_nothing_can_be_selected(document: DocumentInput, settings: Settings):
    with pytest.raises(FileNotFoundError, match="no parse artifact"):
        read_document_figures(document, settings, FakeVisionClient(chart_answer()))


# ---- The stage inside run_document ----------------------------------------------------------------------


def run(monkeypatch, document: DocumentInput, settings: Settings, client: object, **kwargs):
    install_fake_pipeline(monkeypatch)
    monkeypatch.setattr("paperfacts.workflow.build_vision_client", lambda settings: client)
    marks: list[tuple[str, str, str]] = []
    result = run_document(document, settings, on_stage=lambda s, st, d: marks.append((s, st, d)), **kwargs)
    return [mark for mark in marks if mark[0] == "figures"], result, marks


def test_the_stage_runs_beside_extraction_and_its_rows_reach_the_workbook_not_the_dataset(
    monkeypatch, document: DocumentInput, settings: Settings
):
    store_artifact(document, settings)
    client = FakeVisionClient(chart_answer())

    figures, result, marks = run(monkeypatch, document, settings, client)

    assert figures == [("figures", "running", ""), ("figures", "done", "2 readings from 1 panels")]
    order = [(name, status) for name, status, _ in marks]
    # announced right after parsing, joined after the comparison and before export
    assert order.index(("figures", "running")) == order.index(("parse:paddleocr_vl", "done")) + 1
    assert order.index(("extract:mineru", "running")) == order.index(("figures", "running")) + 1
    assert order.index(("figures", "done")) == order.index(("export", "running")) - 1
    assert result.figures is not None and len(result.figures.rows) == 2
    assert not hasattr(result.dataset, "figure_rows")
    sheet = load_workbook(result.excel_path)["图中读数"]
    assert sheet.max_row == 3
    assert client.closed


def test_the_stage_overlaps_the_extraction_lanes(monkeypatch, document: DocumentInput, settings: Settings):
    store_artifact(document, settings)
    lanes_started = threading.Event()
    overlapped: list[bool] = []

    def responder(user, image):
        overlapped.append(lanes_started.wait(timeout=5))
        return chart_answer()

    install_fake_pipeline(monkeypatch)
    real_extract = workflow.extract_document

    def extract(*args, **kwargs):
        lanes_started.set()
        return real_extract(*args, **kwargs)

    monkeypatch.setattr(workflow, "extract_document", extract)
    monkeypatch.setattr(workflow, "build_vision_client", lambda settings: FakeVisionClient(responder))

    run_document(document, settings)

    assert overlapped == [True]


def test_a_failing_stage_is_marked_failed_and_the_paper_still_finishes(
    monkeypatch, document: DocumentInput, settings: Settings
):
    # No artifact anywhere: the stage cannot select anything and raises inside.
    figures, result, marks = run(monkeypatch, document, settings, FakeVisionClient(chart_answer()))

    assert figures[-1][:2] == ("figures", "failed")
    assert "no parse artifact" in figures[-1][2]
    assert marks[-1][:2] == ("export", "done")
    assert result.figures is None


def test_a_client_that_cannot_be_built_fails_only_the_stage(monkeypatch, document: DocumentInput, settings: Settings):
    store_artifact(document, settings)
    install_fake_pipeline(monkeypatch)

    def no_key(settings: Settings):
        raise ConfigError("no LLM API key")

    monkeypatch.setattr("paperfacts.workflow.build_vision_client", no_key)
    marks: list[tuple[str, str, str]] = []
    run_document(document, settings, on_stage=lambda s, st, d: marks.append((s, st, d)))

    assert ("figures", "failed", "ConfigError: no LLM API key") in marks
    assert marks[-1][:2] == ("export", "done")


def test_a_panel_whose_request_failed_marks_the_stage_failed(monkeypatch, document: DocumentInput, settings: Settings):
    store_artifact(document, settings)

    figures, result, _ = run(monkeypatch, document, settings, FakeVisionClient(LlmError("timeout")))

    assert figures[-1] == ("figures", "failed", "0 readings from 1 panels, 1 requests failed")
    assert result.figures is not None and result.figures.rows == ()


def test_a_refused_chart_is_a_done_stage_with_nothing_read(monkeypatch, document: DocumentInput, settings: Settings):
    store_artifact(document, settings)

    figures, result, _ = run(monkeypatch, document, settings, FakeVisionClient(NOT_A_CHART))

    assert figures[-1] == ("figures", "done", "0 readings from 1 panels, 1 not a chart")
    assert result.figures is not None and result.figures.rows == ()


def test_switched_off_the_stage_asks_nothing_but_stored_readings_are_still_shown(
    monkeypatch, document: DocumentInput, settings: Settings
):
    store_artifact(document, settings)
    read_document_figures(document, settings, FakeVisionClient(chart_answer()))
    off = dataclasses.replace(settings, figures_enabled=False)
    client = FakeVisionClient(chart_answer())

    figures, result, _ = run(monkeypatch, document, off, client)

    assert figures == [("figures", "skipped", "figures.enabled is false; 2 stored readings kept")]
    assert client.calls == []
    assert result.figures is not None and len(result.figures.rows) == 2


def test_force_does_not_reread_the_charts_but_force_figures_does(
    monkeypatch, document: DocumentInput, settings: Settings
):
    store_artifact(document, settings)
    read_document_figures(document, settings, FakeVisionClient(chart_answer()))
    plain = FakeVisionClient(chart_answer())
    forced = FakeVisionClient(chart_answer())

    run(monkeypatch, document, settings, plain, force=True)
    run(monkeypatch, document, settings, forced, force_figures=True)

    assert plain.calls == []
    assert [call.refresh for call in forced.calls] == [True]


# ---- What is shown: stale and orphaned readings -----------------------------------------------------------


def test_readings_under_an_older_key_are_shown_marked_stale(monkeypatch, document: DocumentInput, settings: Settings):
    store_artifact(document, settings)
    read_document_figures(document, settings, FakeVisionClient(chart_answer()))
    moved = dataclasses.replace(settings, figures_enabled=False, figures_dpi=150)

    figures, result, _ = run(monkeypatch, document, moved, FakeVisionClient(chart_answer()))

    view = shown_figures(document.document_id, "paper.pdf", moved)
    assert view is not None and view.stale
    assert all("旧版本读数" in (row["detail"] or "") for row in view.rows)
    assert figures == [("figures", "skipped", "figures.enabled is false; 2 stored readings kept; stale (older key)")]
    assert result.figures is not None and result.figures.stale


def test_readings_citing_blocks_the_parse_no_longer_has_are_noted(document: DocumentInput, settings: Settings):
    store_artifact(document, settings)
    read_document_figures(document, settings, FakeVisionClient(chart_answer()))
    ParsedArtifact.read(DataLayout(settings.data_root).artifact_path(document.document_id, "mineru")).model_copy(
        update={"blocks": ()}
    ).write(DataLayout(settings.data_root).artifact_path(document.document_id, "mineru"))

    view = shown_figures(document.document_id, "paper.pdf", settings)

    assert view is not None and view.orphaned == ("mineru_p0_b0",)
    assert "1 cite figure blocks missing from the current parse" in view.warning()
    assert all("当前解析里已没有这个图块" in (row["detail"] or "") for row in view.rows)


def test_nothing_stored_means_nothing_shown(document: DocumentInput, settings: Settings):
    assert shown_figures(document.document_id, "paper.pdf", settings) is None


def test_a_corrupt_stored_file_is_read_again(document: DocumentInput, settings: Settings):
    store_artifact(document, settings)
    path = figures_file(document, settings)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("{ torn", encoding="utf-8")
    client = FakeVisionClient(chart_answer())

    readings = read_document_figures(document, settings, client)

    assert len(client.calls) == 1 and readings.readings


def test_readings_of_another_parse_are_read_again(document: DocumentInput, settings: Settings):
    # Read from PaddleOCR-VL's boxes first; once a MinerU parse exists, those citations belong to no block
    # of the artifact the readings would be shown with.
    store_artifact(document, settings, backend="paddleocr_vl")
    read_document_figures(document, settings, FakeVisionClient(chart_answer()))
    store_artifact(document, settings, backend="mineru")
    client = FakeVisionClient(chart_answer())

    readings = read_document_figures(document, settings, client)

    assert len(client.calls) == 1
    assert readings.backend == "mineru" and {r.source_id for r in readings.readings} == {"mineru_p0_b0"}


def test_an_unreadable_panel_is_asked_again_next_run_past_the_cache(document: DocumentInput, settings: Settings):
    store_artifact(document, settings)
    first = read_document_figures(document, settings, FakeVisionClient("no idea"))
    retry = FakeVisionClient(chart_answer())

    second = read_document_figures(document, settings, retry)

    assert first.panels[0].status == "unreadable"
    assert [call.refresh for call in retry.calls] == [True]
    assert second.complete and second.readings


def test_a_page_is_rendered_once_for_all_its_panels(monkeypatch, document: DocumentInput, settings: Settings):
    blocks = [
        make_block(page=0, order=i, type="figure", content="x.jpg", bbox=BOX, document_id=document.document_id)
        for i in range(3)
    ]
    caption = "Fig. 1 Sheet resistance."
    blocks.append(make_block(page=0, order=3, type="caption", content=caption, document_id=document.document_id))
    ParsedArtifact(
        document_id=document.document_id,
        backend="mineru",
        backend_version="x",
        pages=(PageGeometry(index=0, width_pt=595, height_pt=842),),
        blocks=tuple(blocks),
    ).write(DataLayout(settings.data_root).artifact_path(document.document_id, "mineru"))
    renders: list[int] = []
    real = workflow.render_page

    def counting(path, page, *, dpi):
        renders.append(page)
        return real(path, page, dpi=dpi)

    monkeypatch.setattr(workflow, "render_page", counting)

    readings = read_document_figures(document, settings, FakeVisionClient(chart_answer()))

    assert len(readings.panels) == 3 and renders == [0]


# ---- stopping the stage when the rest of the paper fails ------------------------------------------------


def test_a_failed_extraction_stops_the_charts_and_waits_for_them_before_returning(
    monkeypatch, document: DocumentInput, settings: Settings
):
    """The web queue frees a document when run_document returns; a figures thread still asking the vision
    model then would pay twice for a reprocess and race it for the readings file."""
    install_fake_pipeline(monkeypatch)
    figures_started = threading.Event()
    finished: list[str] = []

    def stoppable_figures(document, settings, *, force, artifact, stop):
        figures_started.set()
        assert stop.wait(timeout=5.0), "the failed paper never told the figures stage to stop"
        finished.append("figures stopped")
        return "failed", "stopped"

    def failing_extraction(document, settings, *, force, on_stage):
        assert figures_started.wait(timeout=5.0)
        raise LlmError("the endpoint is down")

    monkeypatch.setattr("paperfacts.workflow._read_figures_stage", stoppable_figures)
    monkeypatch.setattr("paperfacts.workflow._extract_and_compare", failing_extraction)

    with pytest.raises(LlmError):
        run_document(document, settings)

    assert finished == ["figures stopped"]  # joined, not left running
    assert not [t for t in threading.enumerate() if t.name.startswith("paperfacts-figures")]


def test_a_stopped_figures_stage_asks_no_further_panel_and_stores_nothing(document: DocumentInput, settings: Settings):
    store_artifact(document, settings)
    stop = threading.Event()

    def stop_after_the_first(user: str, image: bytes):
        stop.set()
        return chart_answer()

    client = FakeVisionClient(stop_after_the_first)
    artifact = workflow._figure_artifact(document, settings)
    # Three panels of one figure, asked one at a time.
    panels = tuple(
        make_block(page=0, order=i, type="figure", content=f"{i}.jpg", bbox=BOX, document_id=document.document_id)
        for i in range(3)
    )
    artifact = artifact.model_copy(update={"blocks": (*panels, artifact.blocks[1].model_copy(update={"order": 3}))})

    with pytest.raises(Cancelled):
        read_document_figures(
            document, dataclasses.replace(settings, llm_concurrency=1), client, artifact=artifact, stop=stop
        )

    assert len(client.calls) == 1
    assert not figures_file(document, settings).exists()
